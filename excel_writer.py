"""
报销汇总表写入模块（由 invoice_pipeline.py 拆分而来，契约 = A 类：仅换 import，逻辑不变）。

append_to_summary_excel：把本轮所有「处理成功」邮件的 AI 结果按 SUMMARY_COLUMNS 列顺序追加行，
自动补表头；发票日期列写入前经 sanitize_excel_date_display 清洗，避免非法日期串。
"""

from __future__ import annotations

import os
from typing import Any, Dict, List, Optional, Tuple

from openpyxl import Workbook, load_workbook

from config import SUMMARY_COLUMNS
from utils import sanitize_excel_date_display


def append_to_summary_excel(
    rows: List[Dict[str, Any]],
    summary_path: str,
) -> Tuple[Optional[str], Optional[str]]:
    """rows 来自本轮所有「处理成功」邮件的 AI 结果；无行则跳过。返回 (错误信息, 实际写入的文件路径)。"""
    if not rows:
        return None, None
    try:
        if os.path.exists(summary_path):
            wb = load_workbook(summary_path)
        else:
            wb = Workbook()
        ws = wb.active
        if ws.cell(row=1, column=1).value in (None, ""):
            for j, (header, _) in enumerate(SUMMARY_COLUMNS, start=1):
                ws.cell(row=1, column=j, value=header)
        row_idx = 2
        while ws.cell(row=row_idx, column=1).value not in (None, ""):
            row_idx += 1
        for item in rows:
            for j, (_, key) in enumerate(SUMMARY_COLUMNS, start=1):
                val = item.get(key, "Not Found")
                if key == "invoice_date":
                    val = sanitize_excel_date_display(val)
                ws.cell(row=row_idx, column=j, value=val)
            row_idx += 1
        wb.save(summary_path)
        return None, summary_path
    except PermissionError:
        return "EXCEL_PERMISSION_DENIED: 请关闭 发票信息汇总表.xlsx 后重试", None
    except Exception as e:
        return f"EXCEL_WRITE_ERROR: {str(e)}", None