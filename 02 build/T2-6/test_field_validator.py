"""T2-6 研发自测：使用人工给定的字段与预期，不从程序输出倒推。"""

from __future__ import annotations

import unittest
import json
import tempfile
from decimal import Decimal
from email.message import EmailMessage
from pathlib import Path
from unittest.mock import patch

from config import RunConfig
from data_model import (
    FileRecord, InvoiceFields, InvoiceRecord, ItineraryRecord, MailRecord,
    MatchedItinerary, SourceMeta, mail_to_dict,
)
from field_validator import validate_records
from itinerary_matcher import match_itineraries
from pipeline import _excel_rows, _itinerary_record, _travel_time_to_minute, run_pipeline
from processor_contracts import DocumentResult
from stage_ports import StagePorts, load_stage_ports


NUMBER = "26327000000775232739"


def fixture(kind: str = "b类", receipt: str = "办公用品") -> MailRecord:
    record = InvoiceRecord(
        record_id="Msg1-01", kind=kind,
        fields=InvoiceFields(invoice_date="2026-10-01", invoice_number=NUMBER,
                             total_amount=Decimal("123.4567"), invoice_type="",
                             receipt_type=receipt, seller="甲公司"),
        source=SourceMeta(["01"]), status="complete",
    )
    return MailRecord(mail_id="Msg1", mail_uid="701", subject="发票邮件",
                      mail_date="2026-10-02 08:00",
                      files=[FileRecord("01", original_name="发票.pdf")],
                      records=[record])


class FieldValidatorTests(unittest.TestCase):
    def test_second_precision_is_normalized_at_itinerary_ingress(self):
        mail = fixture("a类", "客运服务费")
        mail.files.append(FileRecord("02", original_name="行程单.pdf"))
        result = DocumentResult(file_seq="02", role="Itinerary", fields={
            "match_amount": "123.4567", "travel_time": "2026-10-01 09:05:28",
            "origin_dest": "杭州→上海"}, source_file_seq=["02"])
        itinerary = _itinerary_record(mail, result)
        mail.itineraries.append(itinerary)
        record = mail.records[0]
        record.matched_itinerary = MatchedItinerary("02", Decimal("123.4567"))
        record.fields.travel_time = itinerary.travel_time
        record.fields.origin_dest = itinerary.origin_dest
        self.assertEqual(itinerary.travel_time, "2026-10-01 09:05")
        self.assertEqual(validate_records(mail), [])
        self.assertEqual(mail_to_dict(mail)["records"][0]["fields"]["travel_time"],
                         "2026-10-01 09:05")

    def test_invalid_second_precision_is_not_silently_accepted(self):
        for raw in ("2026-02-30 09:05:28", "2026-10-01 25:05:28",
                    "2026-10-01 09:05:99", "2026-10-01 09:05:28x"):
            with self.subTest(raw=raw):
                self.assertEqual(_travel_time_to_minute(raw), raw)

    def test_port_attached_and_complete_b_ignores_travel_fields(self):
        mail = fixture()
        mail.records[0].fields.travel_time = "坏时间"
        mail.records[0].fields.origin_dest = "坏路线"
        self.assertIsNotNone(load_stage_ports().validate_records)
        self.assertEqual(validate_records(mail), [])
        self.assertEqual(mail.records[0].status, "complete")
        self.assertEqual(mail.records[0].fields.travel_time, "")
        self.assertEqual(mail.records[0].fields.origin_dest, "")
        self.assertEqual(_excel_rows([mail])[0]["travel_time"], "")
        self.assertEqual(_excel_rows([mail])[0]["origin_dest"], "")

    def test_a_checks_merged_itinerary_and_keeps_decimal_precision(self):
        mail = fixture("a类", "客运服务费")
        mail.files.append(FileRecord("02", original_name="行程单.pdf"))
        mail.itineraries.append(ItineraryRecord("Msg1-02", Decimal("123.4567"),
                                               "2026-10-01 09:05", "杭州→上海"))
        record = mail.records[0]
        record.matched_itinerary = MatchedItinerary("02", Decimal("123.4567"))
        record.fields.travel_time = "2026-10-01 09:05"
        record.fields.origin_dest = "杭州→上海"
        self.assertEqual(validate_records(mail), [])
        saved = mail_to_dict(mail)
        self.assertEqual(saved["records"][0]["fields"]["total_amount"], "123.4567")
        self.assertEqual(saved["itineraries"][0]["match_amount"], "123.4567")
        self.assertEqual(_excel_rows([mail])[0]["travel_time"], "2026-10-01 09:05")

    def test_required_missing_and_partial_row_remains(self):
        mail = fixture()
        fields = mail.records[0].fields
        fields.invoice_date = "Not Found"
        fields.invoice_number = ""
        fields.seller = "Not Found"
        alerts = validate_records(mail)
        self.assertEqual(len(_excel_rows([mail])), 1)
        self.assertEqual(mail.records[0].status, "人工待核")
        self.assertIn("开票日期缺失", alerts[0]["flags"][0])
        self.assertTrue(any("发票号码缺失" in flag for flag in alerts[0]["flags"]))
        self.assertTrue(any("销售方缺失" in flag for flag in alerts[0]["flags"]))
        self.assertEqual(alerts[0]["file_seq"], "01")
        self.assertEqual(alerts[0]["filename"], "发票.pdf")
        self.assertIn("Msg1-01", mail_to_dict(mail)["error_reason"])

    def test_all_three_excel_keys_missing_excludes_only_excel(self):
        mail = fixture()
        fields = mail.records[0].fields
        fields.invoice_date = fields.invoice_number = fields.total_amount = "Not Found"
        validate_records(mail)
        self.assertEqual(_excel_rows([mail]), [])
        self.assertEqual(len(mail_to_dict(mail)["records"]), 1)

    def test_invalid_date_number_amount_and_optional_type(self):
        cases = (
            ("invoice_date", "2026-02-30", "开票日期不是有效日期"),
            ("invoice_date", "2026/10/01", "开票日期格式无效"),
            ("invoice_number", "12345678", "发票号码须为 20 位数字"),
            ("invoice_number", "１２３４５６７８９０１２３４５６７８９０", "发票号码须为 20 位数字"),
            ("total_amount", "0", "报销金额须大于 0"),
            ("total_amount", "12.34567", "报销金额小数超过 4 位"),
            ("total_amount", "100000000000000", "报销金额超出 decimal(18,4) 范围"),
            ("total_amount", "not-money", "报销金额不是有效十进制金额"),
            ("invoice_type", "甲" * 21, "发票类型超过 20 字符"),
        )
        for field, value, expected in cases:
            with self.subTest(field=field, value=value):
                mail = fixture()
                setattr(mail.records[0].fields, field, value)
                alerts = validate_records(mail)
                self.assertTrue(any(expected in flag for flag in alerts[0]["flags"]))
                self.assertEqual(mail.records[0].status, "人工待核")

    def test_optional_type_and_seller_normalization(self):
        mail = fixture()
        fields = mail.records[0].fields
        fields.invoice_type = "Not Found"
        fields.seller = "  甲，乙,公司  "
        self.assertEqual(validate_records(mail), [])
        self.assertEqual(fields.invoice_type, "")
        self.assertEqual(fields.seller, "甲乙公司")

    def test_text_fields_do_not_accept_numeric_values(self):
        labels = {"invoice_type": "发票类型", "receipt_type": "票面类型", "seller": "销售方"}
        for field, label in labels.items():
            with self.subTest(field=field):
                mail = fixture()
                setattr(mail.records[0].fields, field, 123)
                alerts = validate_records(mail)
                self.assertTrue(any(f"{label}类型无效" in flag for flag in alerts[0]["flags"]))
                self.assertEqual(mail.records[0].status, "人工待核")

    def test_receipt_missing_or_overlong_leaves_kind_unknown(self):
        for receipt in ("Not Found", "甲" * 51, "*交通服务*"):
            with self.subTest(receipt=receipt):
                mail = fixture("b类", receipt)
                alerts = validate_records(mail)
                self.assertEqual(mail.records[0].kind, "")
                self.assertEqual(mail.records[0].status, "人工待核")
                self.assertTrue(any("票面类型" in flag for flag in alerts[0]["flags"]))

    def test_b_rejects_stale_itinerary_link_without_validating_travel_values(self):
        mail = fixture()
        record = mail.records[0]
        record.matched_itinerary = MatchedItinerary("02", "123.4567")
        record.fields.travel_time = "坏时间"
        alerts = validate_records(mail)
        self.assertTrue(any("b 类不应关联行程单" in flag for flag in alerts[0]["flags"]))
        self.assertFalse(any("出行时间" in flag for flag in alerts[0]["flags"]))
        self.assertIsNone(record.matched_itinerary)
        self.assertEqual(record.fields.travel_time, "")

    def test_a_missing_match_and_itinerary_invalid_fields(self):
        mail = fixture("a类", "客运服务费")
        mail.files.append(FileRecord("02", original_name="行程单.pdf"))
        mail.itineraries.append(ItineraryRecord("Msg1-02", "-1", "2026-02-30 09:00",
                                               "杭州-上海"))
        alerts = validate_records(mail)
        self.assertEqual(len(alerts), 2)
        itinerary_flags = alerts[0]["flags"]
        invoice_flags = alerts[1]["flags"]
        self.assertTrue(any("行程单匹配金额须大于 0" in flag for flag in itinerary_flags))
        self.assertTrue(any("出行时间不是有效日期时间" in flag for flag in itinerary_flags))
        self.assertTrue(any("出发地/目的地格式无效" in flag for flag in itinerary_flags))
        self.assertTrue(any("a 类缺少已匹配行程单" in flag for flag in invoice_flags))
        self.assertEqual(mail.records[0].status, "人工待核")

    def test_a_rejects_wrong_merge_and_preserves_prior_status(self):
        mail = fixture("a类", "客运服务费")
        mail.files.append(FileRecord("02", original_name="行程单.pdf"))
        mail.itineraries.append(ItineraryRecord("Msg1-02", "123.4567",
                                               "2026-10-01 09:05", "杭州→上海"))
        record = mail.records[0]
        record.status = "缺凭证"
        record.matched_itinerary = MatchedItinerary("02", "123.4567")
        record.fields.travel_time = "2026-10-01 09:06"
        record.fields.origin_dest = "杭州→北京"
        alerts = validate_records(mail)
        self.assertTrue(any("出行时间与已匹配行程单不一致" in flag
                            for flag in alerts[0]["flags"]))
        self.assertTrue(any("出发地/目的地与已匹配行程单不一致" in flag
                            for flag in alerts[0]["flags"]))
        self.assertEqual(record.status, "缺凭证")

    def test_a_marks_record_when_matched_itinerary_amount_invalid(self):
        mail = fixture("a类", "客运服务费")
        mail.files.append(FileRecord("02", original_name="行程单.pdf"))
        mail.itineraries.append(ItineraryRecord("Msg1-02", "0",
                                               "2026-10-01 09:05", "杭州→上海"))
        record = mail.records[0]
        record.matched_itinerary = MatchedItinerary("02", "0")
        record.fields.travel_time = "2026-10-01 09:05"
        record.fields.origin_dest = "杭州→上海"
        alerts = validate_records(mail)
        self.assertEqual(record.status, "人工待核")
        self.assertTrue(any("行程单匹配金额须大于 0" in flag
                            for flag in alerts[1]["flags"]))

    def test_repeat_validation_does_not_duplicate_saved_reasons(self):
        mail = fixture()
        mail.records[0].fields.invoice_number = "bad"
        validate_records(mail)
        first = mail.error_reason
        validate_records(mail)
        self.assertEqual(mail.error_reason, first)

    def test_controlled_pipeline_writes_record_specific_reason_and_excel_row(self):
        message = EmailMessage()
        message["Subject"] = "电子发票"
        message["Date"] = "Thu, 15 Oct 2026 08:00:00 +0800"
        message.set_content("合成测试")
        message.add_attachment(b"synthetic PDF", maintype="application",
                               subtype="pdf", filename="发票.pdf")

        class FakeImap:
            def login(self, *_args): return "OK", []
            def select(self, *_args): return "OK", [b"1"]
            def uid(self, command, *_args):
                return (("OK", [b"701"]) if command == "SEARCH" else
                        ("OK", [(b"1", message.as_bytes())]))
            def status(self, *_args): return "OK", [b"INBOX (MESSAGES 1 UNSEEN 1)"]
            def logout(self): return "BYE", []

        def process_pdf(file, *_args):
            return DocumentResult(file_seq=file.meta.file_seq, role="Invoice",
                                  fields={"invoice_date": "2026-10-01",
                                          "invoice_number": NUMBER,
                                          "total_amount": "123.4567",
                                          "invoice_type": "",
                                          "receipt_type": "办公用品",
                                          "seller": "Not Found"})

        with tempfile.TemporaryDirectory() as root:
            cfg = RunConfig("imap.example", "user", "unused",
                            output_folder_name="run", output_root=root,
                            dashscope_api_key="synthetic-key")
            ports = StagePorts(arbitrate_records=lambda mails: mails,
                               validate_records=validate_records)
            with patch("pipeline.imaplib.IMAP4_SSL", return_value=FakeImap()), \
                 patch("pipeline.process_pdf", side_effect=process_pdf), \
                 patch("pipeline.load_stage_ports", return_value=ports):
                _, stats = run_pipeline(cfg)
            self.assertIsNone(stats["imap_error"])
            saved = json.loads(Path(cfg.debug_file).read_text(encoding="utf-8"))
            self.assertEqual(saved[0]["records"][0]["status"], "人工待核")
            self.assertIn("T2-6 Msg1-01 销售方缺失", saved[0]["error_reason"])
            self.assertEqual(saved[0]["records"][0]["fields"]["total_amount"], "123.4567")
            self.assertTrue(any("Msg1-01 销售方缺失" in flag
                                for alert in stats["not_found_alerts"]
                                for flag in alert["flags"]))
            from openpyxl import load_workbook
            workbook = load_workbook(cfg.summary_excel_file, read_only=True)
            try:
                rows = list(workbook.active.values)
            finally:
                workbook.close()
            self.assertEqual(len(rows), 2)
            self.assertEqual(rows[1][1], NUMBER)
            self.assertEqual(rows[1][2], 123.46)
            self.assertEqual(rows[1][4], "办公用品")

    def test_controlled_pipeline_a_merges_only_itinerary_travel_fields(self):
        message = EmailMessage()
        message["Subject"] = "电子发票"
        message["Date"] = "Thu, 15 Oct 2026 08:00:00 +0800"
        message.set_content("合成测试")
        for filename in ("发票.pdf", "行程单.pdf"):
            message.add_attachment(filename.encode(), maintype="application",
                                   subtype="pdf", filename=filename)

        class FakeImap:
            def login(self, *_args): return "OK", []
            def select(self, *_args): return "OK", [b"1"]
            def uid(self, command, *_args):
                return (("OK", [b"701"]) if command == "SEARCH" else
                        ("OK", [(b"1", message.as_bytes())]))
            def status(self, *_args): return "OK", [b"INBOX (MESSAGES 1 UNSEEN 1)"]
            def logout(self): return "BYE", []

        def process_pdf(file, *_args):
            if file.meta.file_seq == "01":
                return DocumentResult(file_seq="01", role="Invoice",
                                      fields={"invoice_date": "2026-10-01",
                                              "invoice_number": NUMBER,
                                              "total_amount": "53.93",
                                              "receipt_type": "*运输服务*客运服务费",
                                              "seller": "甲公司"})
            return DocumentResult(file_seq="02", role="Itinerary",
                                  fields={"match_amount": "53.93",
                                          "travel_time": "2026-10-01 09:05",
                                          "origin_dest": "杭州→上海"})

        with tempfile.TemporaryDirectory() as root:
            cfg = RunConfig("imap.example", "user", "unused",
                            output_folder_name="run", output_root=root,
                            dashscope_api_key="synthetic-key")
            ports = StagePorts(match_itineraries=match_itineraries,
                               arbitrate_records=lambda mails: mails,
                               validate_records=validate_records)
            with patch("pipeline.imaplib.IMAP4_SSL", return_value=FakeImap()), \
                 patch("pipeline.process_pdf", side_effect=process_pdf), \
                 patch("pipeline.load_stage_ports", return_value=ports):
                _, stats = run_pipeline(cfg)
            self.assertIsNone(stats["imap_error"])
            saved = json.loads(Path(cfg.debug_file).read_text(encoding="utf-8"))[0]
            record = saved["records"][0]
            self.assertEqual(record["kind"], "a类")
            self.assertEqual(record["status"], "complete")
            self.assertEqual(record["matched_itinerary"],
                             {"file_seq": "02", "match_amount": "53.93"})
            self.assertEqual(record["fields"]["receipt_type"], "客运服务费")
            self.assertEqual(record["fields"]["travel_time"], "2026-10-01 09:05")
            self.assertEqual(record["fields"]["origin_dest"], "杭州→上海")
            from openpyxl import load_workbook
            workbook = load_workbook(cfg.summary_excel_file, read_only=True)
            try:
                rows = list(workbook.active.values)
            finally:
                workbook.close()
            self.assertEqual(rows[1][5], "2026-10-01 09:05")
            self.assertEqual(rows[1][7], "杭州→上海")


if __name__ == "__main__":
    unittest.main()
