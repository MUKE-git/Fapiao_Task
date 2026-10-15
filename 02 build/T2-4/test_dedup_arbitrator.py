"""T2-4 研发自测：受控发票记录与 AI 返回，不调用真实邮箱或模型。"""
from __future__ import annotations

import json
import tempfile
import unittest
from decimal import Decimal
from email.message import EmailMessage
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from config import RunConfig
from data_model import (
    FileRecord, InvoiceFields, InvoiceRecord, MailRecord, MatchedItinerary, SourceMeta,
    mail_to_dict,
)
from dedup_arbitrator import arbitrate_records
from excel_writer import append_to_summary_excel
from pipeline import _excel_rows, run_pipeline
from processor_contracts import DocumentResult
from stage_ports import load_stage_ports


NUMBER = "26327000000775232739"


def fixture(index: int, *, number: str = NUMBER, amount: str = "100.00",
            seller: str = "甲公司", invoice_type: str = "", travel_time: str = "Not Found",
            date: str = "2026-10-01") -> MailRecord:
    mail_id = f"Msg{index}"
    record = InvoiceRecord(
        record_id=f"{mail_id}-01", kind="b类",
        fields=InvoiceFields(invoice_date=date, invoice_number=number,
                             total_amount=Decimal(amount), invoice_type=invoice_type,
                             receipt_type="办公用品", seller=seller, travel_time=travel_time),
        source=SourceMeta(["01"]), status="complete",
    )
    return MailRecord(mail_id=mail_id, mail_uid=str(1000 + index),
                      subject=f"第{index}封发票邮件", mail_date="2026-10-01 10:00",
                      files=[FileRecord("01", file_type="PDF", original_name=f"发票{index}.pdf",
                                        internal_name=f"{mail_id}_01_发票.pdf")],
                      records=[record])


class DedupArbitratorTests(unittest.TestCase):
    def setUp(self):
        self.cfg = RunConfig("host", "user", "password", output_root=tempfile.gettempdir())

    def test_same_required_fields_skip_later_excel_only_across_mails(self):
        first = fixture(1, amount="100.0", invoice_type="普通发票")
        later = fixture(2, amount="100.00", invoice_type="专用发票",
                        travel_time="2026-09-20 08:00")
        arbitrate_records([first, later], self.cfg)
        self.assertEqual(len(_excel_rows([first, later])), 1)
        self.assertEqual(later.excel_skip_record_ids, {"Msg2-01"})
        self.assertEqual(len(mail_to_dict(first)["records"]), 1)
        self.assertEqual(len(mail_to_dict(later)["records"]), 1)
        self.assertEqual(mail_to_dict(later)["records"][0]["source"]["source_file_seq"], ["01"])
        self.assertEqual(later.records[0].status, "complete")
        self.assertIn("Msg2-01", later.error_reason)

    def test_empty_number_never_merges_but_same_invalid_string_does(self):
        empty = [fixture(1, number="Not Found"), fixture(2, number="")]
        arbitrate_records(empty, self.cfg)
        self.assertEqual(len(_excel_rows(empty)), 2)
        bad = [fixture(3, number="ABC123"), fixture(4, number="ABC123")]
        arbitrate_records(bad, self.cfg)
        self.assertEqual(len(_excel_rows(bad)), 1)

    def test_same_mail_duplicate_keeps_both_sources_and_matched_itinerary(self):
        mail = fixture(1)
        mail.records[0].kind = "a类"
        mail.records[0].matched_itinerary = MatchedItinerary("03", Decimal("100.00"))
        mail.records[0].fields.travel_time = "2026-09-20 08:00"
        mail.records[0].fields.origin_dest = "北京→上海"
        mail.records[0].status = "人工待核"
        duplicate = fixture(2).records[0]
        duplicate.record_id = "Msg1-02"
        duplicate.source.source_file_seq = ["02"]
        duplicate.fields.invoice_type = "专用发票"
        mail.files.append(FileRecord("02", file_type="PDF", original_name="同票副本.pdf"))
        mail.records.append(duplicate)

        arbitrate_records([mail], self.cfg)
        rows = _excel_rows([mail])
        saved = mail_to_dict(mail)
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["travel_time"], "2026-09-20 08:00")
        self.assertEqual([r["source"]["source_file_seq"] for r in saved["records"]],
                         [["01"], ["02"]])
        self.assertEqual(mail.records[0].status, "人工待核")
        self.assertEqual(mail.records[1].status, "complete")

    def test_different_numbers_and_missing_excel_keys_follow_existing_filter(self):
        first, second, empty = fixture(1), fixture(2, number="99999999999999999999"), fixture(3, number="")
        empty.records[0].fields.invoice_date = "Not Found"
        empty.records[0].fields.total_amount = "Not Found"
        arbitrate_records([first, second, empty], self.cfg)
        self.assertEqual(len(_excel_rows([first, second, empty])), 2)
        self.assertEqual(sum(len(mail_to_dict(mail)["records"]) for mail in (first, second, empty)), 3)

    def test_conflict_ai_selects_existing_candidate_only_with_source_evidence(self):
        first, second = fixture(1, amount="100.00"), fixture(2, amount="200.00")
        seen = {"calls": 0}

        def model_call(*, model, messages):
            seen["calls"] += 1
            seen["model"] = model
            seen["payload"] = json.loads(messages[1]["content"])
            return {"selected_record_id": "Msg2-01", "evidence": [
                {"field": "total_amount", "quote": "价税合计 200.00 元"}], "reason": "第二张原文清晰"}

        def text_loader(_mail, record, _cfg):
            # 第一候选提取成 100，但自己的原文不支持 100；第二候选有原文证据。
            return ("票面文字模糊，金额无法辨认" if record.record_id == "Msg1-01"
                    else "价税合计 200.00 元")

        arbitrate_records([first, second], self.cfg, text_loader=text_loader,
                          ai_call=model_call)
        self.assertEqual([row["total_amount"] for row in _excel_rows([first, second])], [Decimal("200.00")])
        self.assertEqual(first.records[0].status, "人工待核")
        self.assertEqual(len(mail_to_dict(first)["records"]), 1)
        self.assertEqual(len(mail_to_dict(second)["records"]), 1)
        self.assertEqual(seen["model"], self.cfg.dashscope_model)
        self.assertEqual(seen["calls"], 1)
        self.assertEqual(seen["payload"]["different_fields"], ["total_amount"])
        self.assertEqual(len(seen["payload"]["candidates"]), 2)
        self.assertEqual(seen["payload"]["candidates"][0]["mail_subject"], first.subject)
        self.assertEqual([item["pdf_text"] for item in seen["payload"]["candidates"]],
                         ["票面文字模糊，金额无法辨认", "价税合计 200.00 元"])

    def test_third_value_or_unsupported_quote_keeps_conflicting_rows_for_review(self):
        for reply in (
            {"selected_record_id": "Msg3-01", "evidence": [], "reason": "第三种"},
            {"selected_record_id": "Msg2-01", "evidence": [
                {"field": "total_amount", "quote": "价税合计 200.00 元"}], "reason": "无原文证据"},
        ):
            with self.subTest(reply=reply):
                first, second = fixture(1, amount="100.00"), fixture(2, amount="200.00")
                arbitrate_records([first, second], self.cfg,
                    text_loader=lambda *_: "价税合计 100.00 元",
                    ai_call=lambda **_: reply)
                self.assertEqual(len(_excel_rows([first, second])), 2)
                self.assertEqual([first.records[0].status, second.records[0].status],
                                 ["人工待核", "人工待核"])
                self.assertTrue(first.arbitration_alerts[0]["filename"])

    def test_ai_must_support_every_changed_required_field(self):
        first = fixture(1, amount="100.00", seller="甲公司")
        second = fixture(2, amount="200.00", seller="乙公司")
        texts = {"Msg1-01": "价税合计 200.00 元 销售方乙公司",
                 "Msg2-01": "价税合计 200.00 元 销售方乙公司"}
        call = lambda **_: {"selected_record_id": "Msg2-01", "evidence": [
            {"field": "total_amount", "quote": "价税合计 200.00 元"}], "reason": ""}
        arbitrate_records([first, second], self.cfg,
                          text_loader=lambda _mail, record, _cfg: texts[record.record_id],
                          ai_call=call)
        self.assertEqual(len(_excel_rows([first, second])), 2)
        self.assertEqual(first.records[0].status, "人工待核")

        first = fixture(1, amount="100.00", seller="甲公司")
        second = fixture(2, amount="200.00", seller="乙公司")
        call = lambda **_: {"selected_record_id": "Msg2-01", "evidence": [
            {"field": "total_amount", "quote": "价税合计 200.00 元"},
            {"field": "seller", "quote": "销售方乙公司"}], "reason": ""}
        arbitrate_records([first, second], self.cfg,
                          text_loader=lambda _mail, record, _cfg: texts[record.record_id],
                          ai_call=call)
        self.assertEqual(len(_excel_rows([first, second])), 1)
        self.assertIn("销售方：销售方乙公司", second.error_reason)

    def test_ai_cannot_choose_when_both_pdfs_support_their_own_conflicting_values(self):
        first, second = fixture(1, amount="100.00"), fixture(2, amount="200.00")
        texts = {"Msg1-01": "价税合计 100.00 元",
                 "Msg2-01": "价税合计 200.00 元"}
        reply = {"selected_record_id": "Msg2-01", "evidence": [
            {"field": "total_amount", "quote": "价税合计 200.00 元"}], "reason": "猜第二张"}
        arbitrate_records([first, second], self.cfg,
            text_loader=lambda _mail, record, _cfg: texts[record.record_id],
            ai_call=lambda **_: reply)
        self.assertEqual(len(_excel_rows([first, second])), 2)
        self.assertEqual([first.records[0].status, second.records[0].status],
                         ["人工待核", "人工待核"])

    def test_ai_api_response_json_shape_is_parsed(self):
        first, second = fixture(1, amount="100.00"), fixture(2, amount="200.00")
        reply = {"selected_record_id": "Msg2-01", "evidence": [
            {"field": "total_amount", "quote": "价税合计 200.00 元"}], "reason": "原文"}
        response = SimpleNamespace(status_code=200,
            output=SimpleNamespace(text=json.dumps(reply, ensure_ascii=False)))
        arbitrate_records([first, second], self.cfg,
            text_loader=lambda *_: "价税合计 200.00 元",
            ai_call=lambda **_: response)
        self.assertEqual(len(_excel_rows([first, second])), 1)
        self.assertEqual(second.records[0].record_id, "Msg2-01")

    def test_conflict_reads_each_source_pdf_text_then_calls_ai_once(self):
        first, second = fixture(1, amount="100.00"), fixture(2, amount="200.00")
        first.files.append(FileRecord("02", file_type="PDF", original_name="相关行程单.pdf",
                                      internal_name="Msg1_02_行程单.pdf"))
        with tempfile.TemporaryDirectory() as temp:
            cfg = RunConfig("host", "user", "password", output_root=temp)
            root = Path(cfg.extract_dir)
            root.mkdir(parents=True)
            for mail in (first, second):
                for file in mail.files:
                    (root / file.internal_name).write_bytes(b"controlled-pdf-placeholder")
            opened, payloads = [], []

            class FakePdf:
                def __init__(self, path):
                    opened.append(Path(path).name)
                    name = Path(path).name
                    content = ("相关行程单 车次 G123" if "Msg1_02" in name else
                               "金额字样无法辨认" if "Msg1_01" in name else
                               "价税合计 200.00 元")
                    self.pages = [SimpleNamespace(extract_text=lambda: content)]

                def __enter__(self): return self
                def __exit__(self, *_): return False

            def model_call(*, model, messages):
                payloads.append(json.loads(messages[1]["content"]))
                return {"selected_record_id": "Msg2-01", "evidence": [
                    {"field": "total_amount", "quote": "价税合计 200.00 元"}]}

            with patch("dedup_arbitrator.pdfplumber.open", side_effect=FakePdf):
                arbitrate_records([first, second], cfg, ai_call=model_call)
            self.assertEqual(opened, [first.files[0].internal_name,
                                      second.files[0].internal_name,
                                      first.files[1].internal_name])
            self.assertEqual(len(payloads), 1)
            self.assertEqual([c["record_id"] for c in payloads[0]["candidates"]],
                             ["Msg1-01", "Msg2-01"])
            self.assertEqual(payloads[0]["attachment_groups"][0]["attachments"][1]["pdf_text"],
                             "相关行程单 车次 G123")
            self.assertEqual(len(_excel_rows([first, second])), 1)

    def test_unreadable_related_pdf_keeps_conflict_for_manual_review(self):
        first, second = fixture(1, amount="100.00"), fixture(2, amount="200.00")
        first.files.append(FileRecord("02", file_type="PDF", original_name="扫描行程单.pdf"))
        with patch("dedup_arbitrator._pdf_text_file", return_value=""):
            arbitrate_records([first, second], self.cfg,
                text_loader=lambda _mail, record, _cfg: "金额不明" if record is first.records[0]
                            else "价税合计 200.00 元",
                ai_call=lambda **_: self.fail("相关 PDF 缺文字时不可自动仲裁"))
        self.assertEqual(len(_excel_rows([first, second])), 2)
        self.assertEqual([first.records[0].status, second.records[0].status],
                         ["人工待核", "人工待核"])

    def test_duplicate_inside_conflict_keeps_one_per_different_variant(self):
        mails = [fixture(1, amount="100.00"), fixture(2, amount="100.00"),
                 fixture(3, amount="200.00")]
        arbitrate_records(mails, self.cfg, text_loader=lambda *_: "",
                          ai_call=lambda **_: self.fail("无 PDF 原文时不应调用 AI"))
        self.assertEqual(len(_excel_rows(mails)), 2)
        self.assertEqual(sum(len(mail_to_dict(mail)["records"]) for mail in mails), 3)
        self.assertEqual(mails[1].excel_skip_record_ids, {"Msg2-01"})

    def test_stage_port_and_real_excel_writer_receive_selected_rows(self):
        first, second = fixture(1), fixture(2)
        ports = load_stage_ports(self.cfg)
        self.assertIsNotNone(ports.arbitrate_records)
        ports.arbitrate_records([first, second])
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "summary.xlsx"
            error, written = append_to_summary_excel(_excel_rows([first, second]), str(path))
            self.assertIsNone(error)
            self.assertEqual(written, str(path))
            from openpyxl import load_workbook
            book = load_workbook(path)
            self.assertEqual(book.active.max_row, 2)  # 表头 + 一条代表记录
            self.assertEqual(book.active.cell(2, 2).value, NUMBER)
            book.close()

    def test_run_pipeline_writes_one_excel_row_but_both_json_sources(self):
        def message(filename, content):
            msg = EmailMessage()
            msg["Subject"] = "电子发票"
            msg["Date"] = "Tue, 13 Oct 2026 10:00:00 +0800"
            msg.set_content("受控输入")
            msg.add_attachment(content, maintype="application", subtype="pdf", filename=filename)
            return msg

        messages = {"101": message("甲发票.pdf", b"fake-pdf-a"),
                    "102": message("乙发票.pdf", b"fake-pdf-b")}

        class FakeImap:
            def login(self, *_): return "OK", []
            def select(self, *_): return "OK", []
            def status(self, *_): return "OK", [b"INBOX (MESSAGES 2 UNSEEN 2)"]
            def logout(self): return "BYE", []

            def uid(self, command, *args):
                if command.upper() == "SEARCH":
                    return "OK", [b"101 102"]
                uid = args[0].decode() if isinstance(args[0], bytes) else str(args[0])
                return "OK", [(b"body", messages[uid].as_bytes())]

        def extract(file, *_):
            return DocumentResult(file.meta.file_seq, "Invoice", {
                "invoice_date": "2026-10-13", "invoice_number": NUMBER,
                "total_amount": "100.00", "receipt_type": "办公用品", "seller": "甲公司",
            }, [file.meta.file_seq])

        with tempfile.TemporaryDirectory() as temp:
            cfg = RunConfig("imap.test", "user", "password", dashscope_api_key="test-only",
                            output_root=temp)
            with patch("pipeline.imaplib.IMAP4_SSL", return_value=FakeImap()), \
                 patch("pipeline.process_pdf", side_effect=extract):
                _, stats = run_pipeline(cfg)
            self.assertIsNone(stats["imap_error"])
            self.assertTrue(any("已跳过重复 Excel 行" in str(item["flags"])
                                for item in stats["not_found_alerts"]))
            saved = json.loads(Path(cfg.debug_file).read_text(encoding="utf-8"))
            self.assertEqual([len(mail["records"]) for mail in saved], [1, 1])
            self.assertEqual([mail["records"][0]["record_id"] for mail in saved],
                             ["Msg1-01", "Msg2-01"])
            self.assertEqual([mail["mail_uid"] for mail in saved], ["101", "102"])
            from openpyxl import load_workbook
            book = load_workbook(cfg.summary_excel_file, read_only=True)
            self.assertEqual(book.active.max_row, 2)
            book.close()


if __name__ == "__main__":
    unittest.main()
