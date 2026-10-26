"""T3-9 development checks. PDF and HTTP outputs below are fixed inputs."""
from __future__ import annotations

import tempfile
import unittest
import io
import json
import zipfile
from decimal import Decimal
from email.message import EmailMessage
from pathlib import Path
from unittest.mock import patch
from openpyxl import load_workbook

from config import RunConfig, ensure_output_dirs
from data_model import FileRecord, InvoiceFields, InvoiceRecord, MailRecord, SourceMeta
from dedup_arbitrator import arbitrate_records
from mail_link_downloader import DownloadedFile
from pipeline import _check_source_chain, process_message, run_pipeline
from processor_contracts import DocumentResult
from stage_ports import StagePorts
from voucher_matcher import combine_vouchers
from xml_processor import process_xml

SAMPLE_XML = Path(r"C:\Users\lenovo\Downloads\T3出行电子发票(XML)A.xml")
SAMPLE_OFD = Path(r"C:\Users\lenovo\Downloads\T3出行电子发票(OFD)A.ofd")
NUMBER = "26447000000342589935"


class MainFlowTests(unittest.TestCase):
    def setUp(self):
        temporary = tempfile.TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        self.cfg = RunConfig("imap.test", "user", "password", output_root=temporary.name,
                             dashscope_api_key="fixed-test-only")
        ensure_output_dirs(self.cfg)
        self.ports = StagePorts(document_handlers={"XML": process_xml},
                                combine_vouchers=combine_vouchers)

    @staticmethod
    def message(xml: bytes, pdf: bool = False) -> EmailMessage:
        msg = EmailMessage()
        msg["Subject"] = "发票 A"
        msg["Date"] = "Mon, 27 Apr 2026 09:00:00 +0800"
        msg.set_content("本次发票")
        msg.add_attachment(xml, maintype="application", subtype="xml", filename="A.xml")
        if pdf:
            msg.add_attachment(b"%PDF-controlled", maintype="application", subtype="pdf",
                               filename="A.pdf")
        return msg

    @unittest.skipUnless(SAMPLE_XML.is_file(), "本机无用户 A XML 原件")
    def test_real_xml_only_preserves_missing_voucher(self):
        outcome = process_message(self.message(SAMPLE_XML.read_bytes()),
                                  "Msg1", "101", self.cfg, ports=self.ports)
        self.assertEqual(len(outcome.mail.records), 1)
        record = outcome.mail.records[0]
        self.assertEqual(record.fields.invoice_number, NUMBER)
        self.assertEqual(record.fields.total_amount, Decimal("121.74"))
        self.assertEqual(record.status, "缺凭证")
        self.assertTrue(record.source.field_evidence.get("invoice_number"))
        _check_source_chain([outcome.mail], self.cfg)

    @unittest.skipUnless(SAMPLE_XML.is_file(), "本机无用户 A XML 原件")
    def test_real_xml_and_controlled_pdf_result_merge_once(self):
        fixed = DocumentResult("02", "Invoice", {
            "invoice_number": NUMBER, "invoice_date": "2026-04-27",
            "total_amount": "121.74", "receipt_type": "客运服务费"}, ["02"])
        with patch("pipeline.process_pdf", return_value=fixed):
            outcome = process_message(self.message(SAMPLE_XML.read_bytes(), True),
                                      "Msg1", "101", self.cfg, ports=self.ports)
        self.assertEqual(len(outcome.mail.records), 1)
        record = outcome.mail.records[0]
        self.assertEqual(record.record_id, "Msg1-01")
        self.assertEqual(set(record.source.source_file_seq), {"01", "02"})
        _check_source_chain([outcome.mail], self.cfg)

    @unittest.skipUnless(SAMPLE_XML.is_file(), "本机无用户 A XML 原件")
    def test_attachment_and_controlled_link_share_md5(self):
        xml = SAMPLE_XML.read_bytes()
        linked = DownloadedFile("link.xml", xml, "正文链接 example.test（XML）", "XML")
        with patch("pipeline.download_links", return_value=([linked], [])):
            outcome = process_message(self.message(xml), "Msg1", "101",
                                      self.cfg, ports=self.ports)
        self.assertEqual(len(outcome.mail.files), 1)
        self.assertEqual(len(outcome.mail.dedup_skipped), 1)

    def test_cross_mail_xml_evidence_wins_without_ai(self):
        xml = MailRecord(mail_id="Msg1", mail_uid="101",
                         files=[FileRecord("01", file_type="XML")])
        pdf = MailRecord(mail_id="Msg2", mail_uid="102",
                         files=[FileRecord("01", file_type="PDF")])
        xml.records = [InvoiceRecord(record_id="Msg1-01",
            fields=InvoiceFields(invoice_number=NUMBER, invoice_date="2026-04-27",
                total_amount=Decimal("121.74"), receipt_type="客运服务费", seller="销售方"),
            source=SourceMeta(["01"], field_evidence={"total_amount": [{
                "file_seq": "01", "raw": "121.74", "normalized": "121.74",
                "location": "/EInvoice/TotalTaxIncludedAmount"}]}))]
        pdf.records = [InvoiceRecord(record_id="Msg2-01",
            fields=InvoiceFields(invoice_number=NUMBER, invoice_date="2026-04-27",
                total_amount=Decimal("131.74"), receipt_type="客运服务费", seller="销售方"),
            source=SourceMeta(["01"]))]
        arbitrate_records([pdf, xml])
        self.assertIn("Msg2-01", pdf.excel_skip_record_ids)
        self.assertTrue(pdf.records[0].source.delivery_blocked)
        self.assertFalse(xml.records[0].source.delivery_blocked)

    def _run_fixed_imap(self, msg: EmailMessage):
        class FixedImap:
            def login(self, *_args):
                return "OK", []

            def select(self, *_args):
                return "OK", [b"1"]

            def uid(self, command, *_args):
                if command == "SEARCH":
                    return "OK", [b"101"]
                return "OK", [(b"1 (UID 101)", msg.as_bytes())]

            def status(self, *_args):
                return "OK", [b"INBOX (MESSAGES 1 UNSEEN 1)"]

            def logout(self):
                return "BYE", []

        with patch("pipeline.imaplib.IMAP4_SSL", return_value=FixedImap()):
            return run_pipeline(self.cfg)

    @unittest.skipUnless(SAMPLE_XML.is_file(), "本机无用户 A XML 原件")
    def test_fixed_imap_real_xml_to_excel_json_and_checked_zip(self):
        _, stats = self._run_fixed_imap(self.message(SAMPLE_XML.read_bytes()))
        self.assertEqual(stats["delivery_status"], "ready", stats.get("delivery_error"))
        self.assertEqual(stats["summary_row_count"], 1)
        with zipfile.ZipFile(io.BytesIO(stats["delivery_zip_bytes"])) as archive:
            names = archive.namelist()
            self.assertEqual(set(names), {"Invoice_Task/发票信息汇总表.xlsx",
                                          "Invoice_Task/task_debug.json"})
            rows = list(load_workbook(io.BytesIO(archive.read(
                "Invoice_Task/发票信息汇总表.xlsx")), data_only=True).active.values)
            self.assertEqual(rows[1][1], NUMBER)
            mails = json.loads(archive.read("Invoice_Task/task_debug.json"))
            self.assertEqual(mails[0]["records"][0]["status"], "缺凭证")
            self.assertEqual(mails[0]["records"][0]["source"]["source_file_seq"], ["01"])

    def test_fixed_imap_no_files_still_makes_header_only_zip(self):
        msg = EmailMessage()
        msg["Subject"] = "发票"
        msg["Date"] = "Mon, 27 Apr 2026 09:00:00 +0800"
        msg.set_content("无附件")
        _, stats = self._run_fixed_imap(msg)
        self.assertEqual(stats["delivery_status"], "ready", stats.get("delivery_error"))
        self.assertEqual(stats["summary_row_count"], 0)
        with zipfile.ZipFile(io.BytesIO(stats["delivery_zip_bytes"])) as archive:
            rows = list(load_workbook(io.BytesIO(archive.read(
                "Invoice_Task/发票信息汇总表.xlsx")), data_only=True).active.values)
            self.assertEqual(len(rows), 1)
            self.assertEqual(len(rows[0]), 8)

    def test_exhausted_round_link_bytes_does_not_start_download(self):
        msg = EmailMessage()
        msg["Subject"] = "链接发票"
        msg.set_content("请取 https://example.com/invoice.pdf")
        with patch("pipeline.download_links", side_effect=AssertionError("download started")):
            outcome = process_message(msg, "Msg1", "101", self.cfg,
                                      ports=self.ports, max_links=2, max_link_bytes=0)
        self.assertEqual(outcome.link_attempted_count, 0)
        self.assertEqual(outcome.link_downloaded_bytes, 0)
        self.assertTrue(any(alert.get("code") == "P04" for alert in outcome.alerts))
        self.assertIn("LINK_ROUND_BYTES_LIMIT", outcome.mail.error_reason)

    @unittest.skipUnless(SAMPLE_OFD.is_file(), "本机无用户 A OFD 原件")
    def test_real_a_ofd_keeps_fields_when_conversion_is_unsafe(self):
        msg = EmailMessage()
        msg["Subject"] = "发票 OFD A"
        msg["Date"] = "Mon, 27 Apr 2026 09:00:00 +0800"
        msg.set_content("本次 OFD")
        msg.add_attachment(SAMPLE_OFD.read_bytes(), maintype="application",
                           subtype="ofd", filename="A.ofd")
        _, stats = self._run_fixed_imap(msg)
        self.assertEqual(stats["delivery_status"], "ready", stats.get("delivery_error"))
        self.assertTrue(any(alert.get("code") == "P08"
                            for alert in stats["not_found_alerts"]))
        with zipfile.ZipFile(io.BytesIO(stats["delivery_zip_bytes"])) as archive:
            self.assertFalse(any("用户票据/" in name for name in archive.namelist()))
            mails = json.loads(archive.read("Invoice_Task/task_debug.json"))
            self.assertEqual(mails[0]["records"][0]["fields"]["invoice_number"], NUMBER)
            self.assertEqual(mails[0]["records"][0]["fields"]["total_amount"], "121.74")


if __name__ == "__main__":
    unittest.main()
