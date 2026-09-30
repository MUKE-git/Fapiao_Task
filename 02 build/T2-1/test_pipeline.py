from __future__ import annotations

import tempfile
import json
import os
import zipfile
from io import BytesIO
import unittest
from email.message import EmailMessage
from unittest.mock import patch

from openpyxl import load_workbook

from config import RunConfig, ensure_output_dirs
from data_model import MatchedItinerary, STATUS_MANUAL_REVIEW
from pipeline import _apply_matching, _excel_rows, process_message, run_pipeline
from pdf_processor import process_pdf
from processor_contracts import DocumentResult
from stage_ports import StagePorts


def make_message(*attachments: tuple[str, bytes]) -> EmailMessage:
    msg = EmailMessage()
    msg["Subject"] = "电子发票"
    msg["Date"] = "Sat, 26 Sep 2026 09:30:00 +0800"
    msg.set_content("测试邮件")
    for name, content in attachments:
        msg.add_attachment(content, maintype="application", subtype="octet-stream", filename=name)
    return msg


class PipelineDriverTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.cfg = RunConfig("imap.test", "user", "password", output_root=self.tmp.name,
                             dashscope_api_key="test")
        ensure_output_dirs(self.cfg)

    @staticmethod
    def audit(work, *_args):
        name = work.meta.original_name
        if "行程单" in name:
            data = {"total_amount": 53.93, "travel_time": "2026-09-25 10:00",
                    "origin_dest": "杭州→上海"}
            role = "Itinerary"
        else:
            number = "1" * 20 if "甲" in name else "2" * 20
            data = {"invoice_date": "2026-09-26", "invoice_number": number,
                    "total_amount": 53.93, "invoice_type": "电子发票（普通发票）",
                    "receipt_type": "客运服务" if "甲" in name else "办公用品",
                    "seller": "测试公司"}
            role = "Invoice"
        work.meta.pre_read_type = role
        return DocumentResult(file_seq=work.meta.file_seq, role=role, fields=data,
                              source_file_seq=[work.meta.file_seq])

    def test_multiple_invoices_do_not_share_mail_level_fields(self):
        msg = make_message(("甲发票.pdf", b"one"), ("乙发票.pdf", b"two"),
                           ("结构化数据.xml", b"<invoice/>"))
        with patch("pipeline.process_pdf", side_effect=self.audit):
            outcome = process_message(msg, "Msg1", "123", self.cfg)
        mail, alerts = outcome.mail, outcome.alerts
        self.assertEqual([r.record_id for r in mail.records], ["Msg1-01", "Msg1-02"])
        self.assertEqual([r.fields.invoice_number for r in mail.records], ["1" * 20, "2" * 20])
        self.assertEqual([f.file_type for f in mail.files], ["PDF", "PDF", "XML"])
        self.assertEqual([r.source.source_file_seq for r in mail.records], [["01"], ["02"]])
        self.assertEqual(len(_excel_rows([mail])), 2)
        self.assertTrue(any("XML 文件已登记" in str(a["flags"]) for a in alerts))

    def test_single_itinerary_merges_only_into_transport_invoice(self):
        msg = make_message(("甲发票.pdf", b"invoice"), ("甲行程单.pdf", b"itinerary"))
        with patch("pipeline.process_pdf", side_effect=self.audit):
            outcome = process_message(msg, "Msg2", "124", self.cfg)
        mail = outcome.mail
        self.assertEqual(len(mail.records), 1)
        self.assertEqual(mail.records[0].matched_itinerary, None)

        def fake_match(target, itineraries):
            target.records[0].matched_itinerary = MatchedItinerary(itineraries[0].file_seq, 53.93)
            target.records[0].fields.travel_time = itineraries[0].fields["travel_time"]
            target.records[0].fields.origin_dest = itineraries[0].fields["origin_dest"]
            return []

        _apply_matching([outcome], StagePorts(match_itineraries=fake_match))
        row = mail.records[0]
        self.assertEqual(row.matched_itinerary.file_seq, "02")
        self.assertEqual(row.fields.travel_time, "2026-09-25 10:00")
        self.assertEqual(row.fields.origin_dest, "杭州→上海")
        excel_row = _excel_rows([mail])[0]
        self.assertEqual(excel_row["travel_time"], "2026-09-25 10:00")
        self.assertEqual(excel_row["origin_dest"], "杭州→上海")

    def test_unimplemented_formats_do_not_create_fake_records(self):
        msg = make_message(("数据.xml", b"<invoice/>"), ("票据.ofd", b"ofd"),
                           ("照片.png", b"png"), ("未知.dat", b"other"))
        outcome = process_message(msg, "Msg3", "125", self.cfg)
        mail, alerts = outcome.mail, outcome.alerts
        self.assertEqual([item.file_type for item in mail.files], ["XML", "OFD", "图片", "其他"])
        self.assertEqual(mail.records, [])
        self.assertEqual(len(alerts), 4)

    def test_imap_driver_writes_v2_json_and_business_excel(self):
        msg = make_message(("甲发票.pdf", b"one"), ("乙发票.pdf", b"two"))

        class FakeImap:
            def __init__(self):
                self.fetches = []

            def login(self, *_args):
                return "OK", []

            def select(self, *_args):
                return "OK", [b"1"]

            def uid(self, command, *args):
                if command == "SEARCH":
                    return "OK", [b"123"]
                self.fetches.append(args[-1])
                return "OK", [(b"1 (UID 123)", msg.as_bytes())]

            def status(self, *_args):
                return "OK", [b"INBOX (MESSAGES 1 UNSEEN 1)"]

            def logout(self):
                return "BYE", []

        fake = FakeImap()
        with patch("pipeline.imaplib.IMAP4_SSL", return_value=fake), \
             patch("pipeline.process_pdf", side_effect=self.audit):
            mismatches, stats = run_pipeline(self.cfg)
        self.assertEqual(mismatches, [])
        self.assertIsNone(stats["imap_error"])
        self.assertEqual(stats["invoice_keyword_mails"], 1)
        self.assertEqual(fake.fetches, ["(BODY.PEEK[HEADER.FIELDS (SUBJECT DATE)])", "(BODY.PEEK[])"])
        with open(self.cfg.debug_file, encoding="utf-8") as stream:
            saved = json.load(stream)
        self.assertEqual(len(saved[0]["records"]), 2)
        self.assertEqual(saved[0]["mail_uid"], "123")
        self.assertTrue(os.path.isfile(self.cfg.summary_excel_file))
        files = {item["file_seq"] for item in saved[0]["files"]}
        for row in saved[0]["records"]:
            seq = row["source"]["source_file_seq"][0]
            self.assertIn(seq, files)
            self.assertEqual(row["record_id"], f"{saved[0]['mail_id']}-{seq}")
        book = load_workbook(self.cfg.summary_excel_file, read_only=True)
        try:
            sheet = book.active
            self.assertEqual(sheet.max_row, 3)
            self.assertEqual([sheet.cell(row=i, column=2).value for i in (2, 3)],
                             [row["fields"]["invoice_number"] for row in saved[0]["records"]])
        finally:
            book.close()

    def test_zip_duplicate_is_skipped_with_one_file_sequence(self):
        buffer = BytesIO()
        with zipfile.ZipFile(buffer, "w") as archive:
            archive.writestr("重复发票.pdf", b"same")
            archive.writestr("结构化数据.xml", b"<invoice/>")
        msg = make_message(("甲发票.pdf", b"same"), ("附件.zip", buffer.getvalue()))
        with patch("pipeline.process_pdf", side_effect=self.audit):
            mail = process_message(msg, "Msg4", "126", self.cfg).mail
        self.assertEqual([f.file_seq for f in mail.files], ["01", "02"])
        self.assertEqual([f.file_type for f in mail.files], ["PDF", "XML"])
        self.assertEqual(len(mail.dedup_skipped), 1)

    def test_pdf_adapter_submits_only_one_file_to_old_ai_logic(self):
        msg = make_message(("甲发票.pdf", b"first"), ("乙发票.pdf", b"second"))
        from attachment_processor import collect_mail_files
        from data_model import MailRecord
        mail = MailRecord(mail_id="Msg5", mail_uid="127")
        files = collect_mail_files(msg, mail, self.cfg)
        seen = []

        def fake_ai(task_info, **_kwargs):
            self.assertEqual(len(task_info["files"]), 1)
            name = task_info["files"][0]["attachment_original_name"]
            seen.append(name)
            number = "1" * 20 if "甲" in name else "2" * 20
            return {"invoice_number": number, "total_amount": 10,
                    "file_classifications": [{"role": "Invoice"}]}, "", [], []

        with patch("pdf_processor.local_inspect_pdf", return_value=("Invoice", "发票号码：123")), \
             patch("pdf_processor.call_ai_audit_and_rename", side_effect=fake_ai):
            results = [process_pdf(file, self.cfg, "电子发票", "2026-09-26") for file in files]
        self.assertEqual(seen, ["甲发票.pdf", "乙发票.pdf"])
        self.assertEqual([result.fields["invoice_number"] for result in results],
                         ["1" * 20, "2" * 20])

    def test_t2_ports_run_after_file_processing_in_contract_order(self):
        msg = make_message(("甲发票.pdf", b"invoice"), ("甲行程单.pdf", b"itinerary"))

        class FakeImap:
            def login(self, *_args): return "OK", []
            def select(self, *_args): return "OK", [b"1"]
            def uid(self, command, *_args):
                return ("OK", [b"123"]) if command == "SEARCH" else ("OK", [(b"1", msg.as_bytes())])
            def status(self, *_args): return "OK", [b"INBOX (MESSAGES 1 UNSEEN 1)"]
            def logout(self): return "BYE", []

        called = []

        def match(mail, evidence):
            called.append("T2-3")
            self.assertEqual(len(evidence), 1)
            mail.records[0].matched_itinerary = MatchedItinerary(evidence[0].file_seq, 53.93)
            return []

        def arbitrate(mails):
            called.append("T2-4")
            return mails

        def validate(mail):
            called.append("T2-6")
            self.assertEqual(mail.records[0].matched_itinerary.file_seq, "02")
            return []

        ports = StagePorts(match_itineraries=match, arbitrate_records=arbitrate,
                           validate_records=validate)
        with patch("pipeline.imaplib.IMAP4_SSL", return_value=FakeImap()), \
             patch("pipeline.process_pdf", side_effect=self.audit), \
             patch("pipeline.load_stage_ports", return_value=ports):
            _, stats = run_pipeline(self.cfg)
        self.assertIsNone(stats["imap_error"])
        self.assertEqual(called, ["T2-3", "T2-4", "T2-6"])


if __name__ == "__main__":
    unittest.main()
