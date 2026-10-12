"""十七项修订中 T2-1 增量的研发自测。"""
from __future__ import annotations

import json
import tempfile
import unittest
from decimal import Decimal
from unittest.mock import patch

from openpyxl import load_workbook

from attachment_processor import collect_mail_files
from config import RunConfig, ensure_output_dirs
from data_model import MailRecord, NOT_FOUND, STATUS_MANUAL_REVIEW
from excel_writer import append_to_summary_excel
from pdf_processor import process_pdf
from pipeline import run_pipeline
from processor_contracts import DocumentResult
from test_pipeline import make_message


class ConflictIncrementalTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.cfg = RunConfig("imap.test", "user", "password", output_root=self.tmp.name,
                             dashscope_api_key="test")
        ensure_output_dirs(self.cfg)

    def test_excel_filters_only_all_missing_per_invoice(self):
        msg = make_message(*[(f"发票{i}.pdf", f"pdf{i}".encode()) for i in range(4)])

        class FakeImap:
            def login(self, *_): return "OK", []
            def select(self, *_): return "OK", [b"1"]
            def uid(self, command, *_):
                return ("OK", [b"123"]) if command == "SEARCH" else ("OK", [(b"1", msg.as_bytes())])
            def status(self, *_): return "OK", [b"INBOX (MESSAGES 1 UNSEEN 1)"]
            def logout(self): return "BYE", []

        values = [
            (NOT_FOUND, NOT_FOUND, NOT_FOUND),
            ("2026-10-04", NOT_FOUND, NOT_FOUND),
            ("2026-10-04", "123", NOT_FOUND),
            ("invalid-date", "456", Decimal("12.3456")),
        ]

        def extract(file, *_):
            date, number, amount = values[int(file.meta.original_name[2])]
            return DocumentResult(file.meta.file_seq, "Invoice", {
                "invoice_date": date, "invoice_number": number,
                "total_amount": amount, "receipt_type": "办公用品"},
                [file.meta.file_seq])

        with patch("pipeline.imaplib.IMAP4_SSL", return_value=FakeImap()), \
             patch("pipeline.process_pdf", side_effect=extract):
            _, stats = run_pipeline(self.cfg)
        self.assertIsNone(stats["imap_error"])
        with open(self.cfg.debug_file, encoding="utf-8") as stream:
            saved = json.load(stream)
        self.assertEqual(len(saved[0]["records"]), 4)
        self.assertEqual(saved[0]["records"][0]["status"], STATUS_MANUAL_REVIEW)
        self.assertEqual(saved[0]["records"][3]["fields"]["total_amount"], "12.3456")
        book = load_workbook(self.cfg.summary_excel_file, read_only=True)
        try:
            self.assertEqual(book.active.max_row, 4)
            self.assertEqual(book.active.cell(4, 3).value, 12.35)
        finally:
            book.close()
        flags = str(stats["not_found_alerts"])
        self.assertIn("均缺失，未写入 Excel", flags)
        self.assertIn("已写入 Excel", flags)
        self.assertNotIn("Msg1-04 缺少", flags)

    def test_pdf_classification_must_identify_the_one_input(self):
        msg = make_message(("发票.pdf", b"pdf"))
        mail = MailRecord(mail_id="Msg1", mail_uid="123")
        file = collect_mail_files(msg, mail, self.cfg)[0]

        def check(classifications, expected_error):
            def fake_ai(task_info, **_):
                name = task_info["files"][0]["local_path"].split("\\")[-1]
                result = classifications(name)
                return {"invoice_number": "123", "file_classifications": result}, "", [], []
            with patch("pdf_processor.local_inspect_pdf", return_value=("Unknown", "")), \
                 patch("pdf_processor.call_ai_audit_and_rename", side_effect=fake_ai):
                result = process_pdf(file, self.cfg, "电子发票", "2026-10-04")
            self.assertEqual(bool(result.error_reason), expected_error)
            self.assertEqual(result.role, "Invoice")
            self.assertEqual(bool(result.manual_unknown), expected_error)

        check(lambda name: [{"file": name, "role": "Invoice"}], False)
        check(lambda name: [], True)
        check(lambda name: [{"file": name, "role": "Invoice"}] * 2, True)
        check(lambda name: [{"file": "other.pdf", "role": "Invoice"}], True)
        check(lambda name: [{"file": name, "role": "Other"}], True)

    def test_partial_missing_date_does_not_overwrite_on_next_run(self):
        first = {"invoice_date": NOT_FOUND, "invoice_number": "123", "total_amount": Decimal("1.2")}
        second = {"invoice_date": "2026-10-04", "invoice_number": "456", "total_amount": Decimal("2.3")}
        self.assertIsNone(append_to_summary_excel([first], self.cfg.summary_excel_file)[0])
        self.assertIsNone(append_to_summary_excel([second], self.cfg.summary_excel_file)[0])
        book = load_workbook(self.cfg.summary_excel_file, read_only=True)
        try:
            self.assertEqual(book.active.max_row, 3)
            self.assertEqual([book.active.cell(i, 2).value for i in (2, 3)], ["123", "456"])
        finally:
            book.close()


if __name__ == "__main__":
    unittest.main()
