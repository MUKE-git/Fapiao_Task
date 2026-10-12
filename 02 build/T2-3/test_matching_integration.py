"""T2-3 研发集成验证：真实匹配端口，模拟收件及逐文件提取。"""
from __future__ import annotations

from email.message import EmailMessage
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import load_workbook

from config import RunConfig, SUMMARY_COLUMNS, ensure_output_dirs
from data_model import NOT_FOUND, mail_to_dict
from itinerary_matcher import match_itineraries
from pipeline import _apply_matching, _check_source_chain, _excel_rows, process_message, run_pipeline
from processor_contracts import DocumentResult
from stage_ports import StagePorts, load_stage_ports


def message(*names):
    msg = EmailMessage()
    msg["Subject"] = "电子发票研发验证"
    msg["Date"] = "Sun, 04 Oct 2026 10:00:00 +0800"
    msg.set_content("模拟内容")
    for index, name in enumerate(names):
        msg.add_attachment(f"different-content-{index}".encode(),
                           maintype="application", subtype="octet-stream", filename=name)
    return msg


class FakeImap:
    def __init__(self, messages):
        self.messages = messages
        self.queries = []

    def login(self, *_): return "OK", []
    def select(self, *_): return "OK", [str(len(self.messages)).encode()]
    def status(self, *_):
        return "OK", [f"INBOX (MESSAGES {len(self.messages)} UNSEEN {len(self.messages)})".encode()]
    def logout(self): return "BYE", []

    def uid(self, command, *args):
        if command.upper() == "SEARCH":
            return "OK", [" ".join(self.messages).encode()]
        uid, query = args
        uid = uid.decode() if isinstance(uid, bytes) else str(uid)
        self.queries.append(query)
        return "OK", [(b"body", self.messages[uid].as_bytes())]


class MatchingIntegrationTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.cfg = RunConfig("imap.test", "test-user", "test-password",
                             output_root=self.tmp.name, dashscope_api_key="test-only")
        ensure_output_dirs(self.cfg)

    def extraction(self, mapping):
        def extract(file, *_):
            role = "Itinerary" if "行程单" in file.meta.original_name else "Invoice"
            fields = dict(mapping[file.meta.original_name])
            if role == "Invoice":
                fields.setdefault("invoice_date", "2026-10-04")
                fields.setdefault("invoice_number", file.meta.file_seq * 10)
                fields.setdefault("receipt_type", "客运服务")
                fields.setdefault("seller", "发票企业")
            else:
                fields.setdefault("travel_time", "2026-10-03 09:00")
                fields.setdefault("origin_dest", "杭州→上海")
            return DocumentResult(file.meta.file_seq, role, fields, [file.meta.file_seq])
        return extract

    def mailbox(self, messages, mapping):
        fake = FakeImap(messages)
        with patch("pipeline.imaplib.IMAP4_SSL", return_value=fake), \
             patch("pipeline.process_pdf", side_effect=self.extraction(mapping)):
            _, stats = run_pipeline(self.cfg)
        self.assertIsNone(stats["imap_error"])
        saved = json.loads(Path(self.cfg.debug_file).read_text(encoding="utf-8"))
        return stats, saved, fake

    def excel(self):
        book = load_workbook(self.cfg.summary_excel_file, read_only=True, data_only=True)
        try:
            return list(book.active.values), [book.active.cell(i, 3).number_format
                                              for i in range(2, book.active.max_row + 1)]
        finally:
            book.close()

    def test_stage_ports_automatically_load_real_matcher(self):
        self.assertIs(load_stage_ports().match_itineraries, match_itineraries)

    def test_real_match_reaches_arbitration_then_validation(self):
        mapping = {"发票A.pdf": {"total_amount": "100.00"},
                   "行程单A.pdf": {"match_amount": "100.0000"}}
        steps = []

        def matching(mail, evidence):
            steps.append("T2-3")
            return match_itineraries(mail, evidence)

        def arbitrate(mails):
            steps.append("T2-4")
            self.assertEqual(mails[0].records[0].matched_itinerary.file_seq, "02")
            return mails

        def validate(mail):
            steps.append("T2-6")
            record = mail.records[0]
            self.assertEqual(record.fields.origin_dest, "杭州→上海")
            record.status = "人工待核"
            mail.error_reason = "后续字段校验待核"
            return [{"mail_subject": mail.subject, "flags": ["后续字段校验待核"]}]

        ports = StagePorts(match_itineraries=matching, arbitrate_records=arbitrate,
                           validate_records=validate)
        with patch("pipeline.load_stage_ports", return_value=ports):
            stats, saved, _ = self.mailbox({"101": message(*mapping)}, mapping)
        self.assertEqual(steps, ["T2-3", "T2-4", "T2-6"])
        self.assertEqual(saved[0]["records"][0]["status"], "人工待核")
        self.assertIn("后续字段校验待核", str(stats["not_found_alerts"]))
        self.assertEqual(self.excel()[0][1][7], "杭州→上海")

    def test_B03_business_with_actual_port_not_empty_ports(self):
        mapping = {"客运发票C.pdf": {"total_amount": "100.00"},
                   "行程单A.pdf": {"match_amount": "900.00"}}
        with patch("pipeline.process_pdf", side_effect=self.extraction(mapping)):
            outcome = process_message(message(*mapping), "Msg1", "101", self.cfg)
        self.assertIsNone(outcome.mail.records[0].matched_itinerary)
        _apply_matching([outcome], load_stage_ports())
        _check_source_chain([outcome.mail])
        record = outcome.mail.records[0]
        self.assertEqual(record.matched_itinerary.file_seq, "02")
        self.assertEqual(record.matched_itinerary.match_amount, 900)
        self.assertEqual(_excel_rows([outcome.mail])[0]["origin_dest"], "杭州→上海")
        saved = mail_to_dict(outcome.mail)
        self.assertEqual(saved["records"][0]["matched_itinerary"]["match_amount"], "900.00")
        self.assertEqual(saved["records"][0]["source"]["source_file_seq"], ["01"])

    def test_full_pipeline_precision_json_and_eight_column_excel(self):
        mapping = {"发票A.pdf": {"total_amount": "100.00"},
                   "行程单A.pdf": {"match_amount": "100.0049"},
                   "行程单B.pdf": {"match_amount": "100.50"}}
        stats, saved, fake = self.mailbox({"101": message(*mapping)}, mapping)
        self.assertEqual(saved[0]["records"][0]["matched_itinerary"],
                         {"file_seq": "02", "match_amount": "100.0049"})
        self.assertEqual(len(saved[0]["itineraries"]), 2)
        self.assertEqual(set(saved[0]["itineraries"][0]),
                         {"record_id", "match_amount", "travel_time", "origin_dest"})
        self.assertIn("Msg1-03 行程单未匹配", saved[0]["error_reason"])
        rows, formats = self.excel()
        self.assertEqual(rows[0], tuple(name for name, _ in SUMMARY_COLUMNS))
        self.assertEqual(len(rows), 2)
        self.assertEqual((rows[1][2], rows[1][7]), (100, "杭州→上海"))
        self.assertEqual(formats, ["0.00"])
        self.assertNotIn("匹配尚待 T2-3", str(stats["not_found_alerts"]))
        self.assertTrue(all("BODY.PEEK" in query for query in fake.queries))

    def test_conflict_and_independent_match_reach_excel_and_json(self):
        # 无唯一明确尾缀时保留原冲突处理；A/单A 现在应优先关联。
        mapping = {"发票D.pdf": {"total_amount": "100.00"},
                   "发票B.pdf": {"total_amount": "100.00"},
                   "发票C.pdf": {"total_amount": "200.00"},
                   "行程单A.pdf": {"match_amount": "100.00"},
                   "行程单C.pdf": {"match_amount": "200.00"}}
        stats, saved, _ = self.mailbox({"101": message(*mapping)}, mapping)
        records = saved[0]["records"]
        self.assertTrue(all("matched_itinerary" not in item for item in records[:2]))
        self.assertEqual(records[2]["matched_itinerary"]["file_seq"], "05")
        self.assertIn("一对一关系冲突", saved[0]["error_reason"])
        self.assertIn("一对一关系冲突", str(stats["not_found_alerts"]))
        rows, _ = self.excel()
        self.assertEqual(len(rows), 4)
        self.assertEqual([row[7] for row in rows[1:]], [None, None, "杭州→上海"])

    def test_iterative_same_amount_pairs_reach_json_and_excel(self):
        mapping = {"发票A.pdf": {"total_amount": "100.00"},
                   "发票.pdf": {"total_amount": "100.00"},
                   "行程单A.pdf": {"match_amount": "100.0000", "origin_dest": "杭州→上海"},
                   "行程单.pdf": {"match_amount": "100.00", "origin_dest": "上海→苏州"}}
        stats, saved, _ = self.mailbox({"101": message(*mapping)}, mapping)
        records = saved[0]["records"]
        self.assertEqual([r["matched_itinerary"]["file_seq"] for r in records], ["03", "04"])
        self.assertEqual([r["source"]["source_file_seq"] for r in records], [["01"], ["02"]])
        self.assertEqual(len(saved[0]["itineraries"]), 2)
        self.assertEqual(records[0]["matched_itinerary"]["match_amount"], "100.0000")
        self.assertNotIn("候选无法唯一确定", str(stats["not_found_alerts"]))
        self.assertNotIn("行程单未匹配", saved[0]["error_reason"])
        rows, _ = self.excel()
        self.assertEqual(len(rows), 3)
        self.assertEqual(len(rows[0]), 8)
        self.assertEqual([row[7] for row in rows[1:]], ["杭州→上海", "上海→苏州"])

    def test_suffix_priority_and_residual_mismatch_reach_pipeline(self):
        mapping = {"发票A.pdf": {"total_amount": "100.00"},
                   "发票.pdf": {"total_amount": "100.00"},
                   "行程单A.pdf": {"match_amount": "100.00"},
                   "行程单.pdf": {"match_amount": "200.0000", "origin_dest": "上海→苏州"}}
        stats, saved, _ = self.mailbox({"101": message(*mapping)}, mapping)
        records = saved[0]["records"]
        self.assertEqual([r["matched_itinerary"]["file_seq"] for r in records], ["03", "04"])
        self.assertEqual(records[1]["fields"]["total_amount"], "100.00")
        self.assertEqual(records[1]["matched_itinerary"]["match_amount"], "200.0000")
        self.assertNotIn("一对一关系冲突", str(stats["not_found_alerts"]))
        self.assertEqual([row[2] for row in self.excel()[0][1:]], [100, 100])

    def test_residual_missing_invoice_amount_still_gets_travel_fields(self):
        mapping = {"发票.pdf": {"total_amount": NOT_FOUND},
                   "发票A.pdf": {"total_amount": "100.00"},
                   "行程单.pdf": {"match_amount": "200.0000", "origin_dest": "上海→苏州"},
                   "行程单A.pdf": {"match_amount": "100.00"}}
        _, saved, _ = self.mailbox({"101": message(*mapping)}, mapping)
        first = saved[0]["records"][0]
        self.assertEqual(first["matched_itinerary"]["file_seq"], "03")
        self.assertEqual(first["fields"]["total_amount"], NOT_FOUND)
        self.assertEqual(first["fields"]["origin_dest"], "上海→苏州")
        self.assertEqual(first["status"], "人工待核")
        self.assertNotIn("金额与原名尾缀均未命中", saved[0]["error_reason"])
        rows, _ = self.excel()
        self.assertEqual(len(rows), 3)
        self.assertEqual(rows[1][7], "上海→苏州")

    def test_complex_filename_fallback_and_b_invoice_isolation(self):
        mapping = {"发票A.pdf": {"total_amount": NOT_FOUND},
                   "发票B.pdf": {"total_amount": "200.00"},
                   "办公发票C.pdf": {"total_amount": "300.00", "receipt_type": "办公用品"},
                   "行程单A.pdf": {"match_amount": "900.00"},
                   "行程单B.pdf": {"match_amount": "800.00"}}
        _, saved, _ = self.mailbox({"101": message(*mapping)}, mapping)
        records = saved[0]["records"]
        self.assertEqual([records[i]["matched_itinerary"]["file_seq"] for i in (0, 1)], ["04", "05"])
        self.assertNotIn("matched_itinerary", records[2])
        self.assertNotIn("travel_time", records[2]["fields"])
        rows, _ = self.excel()
        self.assertEqual(len(rows), 4)
        self.assertEqual(rows[3][7], None)

    def test_all_missing_invoice_stays_in_json_with_match_and_is_filtered_from_excel(self):
        mapping = {"发票A.pdf": {"invoice_date": NOT_FOUND, "invoice_number": NOT_FOUND,
                                 "total_amount": NOT_FOUND},
                   "行程单A.pdf": {"match_amount": NOT_FOUND}}
        stats, saved, _ = self.mailbox({"101": message(*mapping)}, mapping)
        self.assertEqual(saved[0]["records"][0]["matched_itinerary"]["file_seq"], "02")
        self.assertEqual(saved[0]["records"][0]["matched_itinerary"]["match_amount"], NOT_FOUND)
        self.assertFalse(Path(self.cfg.summary_excel_file).exists())
        self.assertIn("均缺失，未写入 Excel", str(stats["not_found_alerts"]))

    def test_itinerary_only_is_retained_without_invoice_row(self):
        mapping = {"行程单A.pdf": {"match_amount": "100.0000"}}
        _, saved, _ = self.mailbox({"101": message(*mapping)}, mapping)
        self.assertEqual(saved[0]["records"], [])
        self.assertEqual(saved[0]["itineraries"][0]["match_amount"], "100.0000")
        self.assertFalse(Path(self.cfg.summary_excel_file).exists())

    def test_matching_does_not_cross_mail_boundaries(self):
        mapping = {"发票A.pdf": {"total_amount": "100.00"},
                   "行程单A.pdf": {"match_amount": "100.00"}}
        _, saved, _ = self.mailbox({"101": message("发票A.pdf"),
                                  "102": message("行程单A.pdf")}, mapping)
        self.assertNotIn("matched_itinerary", saved[0]["records"][0])
        self.assertEqual(saved[1]["records"], [])


if __name__ == "__main__":
    unittest.main()
