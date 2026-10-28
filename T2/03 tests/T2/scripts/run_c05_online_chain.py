"""C05 approved synthetic email through the real T2 PDF/AI/business chain.

Only IMAP is replaced by the approved RFC822 fixture. No PDF or AI return is
stubbed. The online model is capped at two requests, and credentials are never
written to the report. Expected business values come from approved expected.json.
"""
from __future__ import annotations

import hashlib
import json
import sys
import tempfile
import tomllib
from datetime import datetime
from decimal import Decimal
from email import message_from_bytes
from email.message import EmailMessage
from pathlib import Path
from unittest.mock import patch

import dashscope
from dashscope import Generation
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[3]
sys.path.insert(0, str(ROOT))
import ai_client  # noqa: E402
import pdf_processor  # noqa: E402
from config import RunConfig, SUMMARY_COLUMNS  # noqa: E402
from pipeline import run_pipeline  # noqa: E402

FIXTURE = ROOT / "03 tests" / "owner test" / "固定用例库" / "cases" / "C05" / "variants" / "v1"
BASE = ROOT / "03 tests" / "T2"
RUN = BASE / "result" / f"C05_online_{datetime.now().strftime('%Y%m%d_%H%M%S')}"
RUN.mkdir(parents=True, exist_ok=False)
tempfile.tempdir = str(RUN.resolve())


def sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def check(diffs: list[dict], path: str, got, wanted) -> None:
    if got != wanted:
        diffs.append({"path": path, "expected": wanted, "actual": got})


class FixtureImap:
    def __init__(self, uid: str, raw: bytes):
        self.uid_text, self.raw = uid, raw
        self.events: list[str] = []

    def login(self, *_):
        self.events.append("LOGIN")
        return "OK", []

    def select(self, *_):
        self.events.append("SELECT INBOX")
        return "OK", [b"1"]

    def status(self, *_):
        self.events.append("STATUS INBOX")
        return "OK", [b"INBOX (MESSAGES 1 UNSEEN 1)"]

    def uid(self, command, *args):
        if command == "SEARCH":
            self.events.append("UID SEARCH UNSEEN")
            return "OK", [self.uid_text.encode()]
        if command != "FETCH":
            raise AssertionError(f"Unexpected IMAP command: {command}")
        uid = args[0].decode() if isinstance(args[0], bytes) else str(args[0])
        if uid != self.uid_text:
            raise AssertionError(f"Unexpected UID: {uid}")
        query = str(args[1])
        self.events.append(f"UID FETCH {query}")
        if "HEADER.FIELDS" in query:
            source = message_from_bytes(self.raw)
            header = EmailMessage()
            header["Subject"] = source["Subject"]
            header["Date"] = source["Date"]
            return "OK", [(b"header", header.as_bytes())]
        return "OK", [(b"body", self.raw)]

    def logout(self):
        self.events.append("LOGOUT")
        return "BYE", []


def main() -> int:
    secrets_path = ROOT / ".streamlit" / "secrets.toml"
    secrets = tomllib.loads(secrets_path.read_text(encoding="utf-8-sig"))
    api_key = str(secrets.get("DASHSCOPE_API_KEY") or "").strip()
    model = str(secrets.get("DASHSCOPE_MODEL") or "qwen-plus").strip()
    if not api_key:
        raise RuntimeError("No configured DashScope key; online chain not started")
    dashscope.api_key = api_key
    approved = json.loads((FIXTURE / "expected.json").read_text(encoding="utf-8"))["expected"]
    state = json.loads((FIXTURE / "imap_state.json").read_text(encoding="utf-8"))
    inventory = json.loads((FIXTURE / "input_inventory.json").read_text(encoding="utf-8"))
    raw = (FIXTURE / "00_raw.eml").read_bytes()
    fake = FixtureImap(str(state["uid"]), raw)
    cfg = RunConfig("imap.fixture.test", "fixture", "unused", output_folder_name="output",
                    output_root=str(RUN), dashscope_api_key=api_key, dashscope_model=model)
    calls: list[dict] = []
    documents: list[dict] = []
    original_call = Generation.call
    original_pdf = pdf_processor.process_pdf

    def capped_call(*args, **kwargs):
        if len(calls) >= 2:
            raise RuntimeError("Online request cap reached: 2")
        call = {"model": str(kwargs.get("model", "")), "status": None}
        calls.append(call)
        response = original_call(*args, **kwargs)
        call["status"] = getattr(response, "status_code", None)
        return response

    def observed_pdf(file, *args, **kwargs):
        result = original_pdf(file, *args, **kwargs)
        documents.append({"file_seq": file.meta.file_seq,
                          "original_name": file.meta.original_name,
                          "role": result.role,
                          "fields": {key: str(value) for key, value in result.fields.items()
                                     if key in ("invoice_date", "invoice_number", "total_amount",
                                                "invoice_type", "receipt_type", "seller",
                                                "match_amount", "travel_time", "origin_dest",
                                                "data_source")},
                          "error_reason": result.error_reason,
                          "pre_read_type": file.meta.pre_read_type,
                          "pre_read_invoice_no": file.meta.pre_read_invoice_no,
                          "pre_read_amount": str(file.meta.pre_read_amount)})
        return result

    with patch("pipeline.imaplib.IMAP4_SSL", return_value=fake), \
         patch.object(Generation, "call", side_effect=capped_call), \
         patch("pipeline.process_pdf", side_effect=observed_pdf), \
         patch.object(ai_client, "__file__", str(RUN / "ai_client.py")):
        mismatches, stats = run_pipeline(cfg)

    diffs: list[dict] = []
    check(diffs, "input.eml.sha256", sha(FIXTURE / "00_raw.eml"), inventory[0]["sha256"])
    for item in inventory[1:]:
        source = FIXTURE / item["path"]
        check(diffs, f"input.{item['path']}.sha256", sha(source), item["sha256"])
    check(diffs, "online.requests", len(calls), 2)
    check(diffs, "online.statuses", [call["status"] for call in calls], [200, 200])
    check(diffs, "pdf_processor.files", [item["original_name"] for item in documents],
          ["发票.pdf", "行程单.pdf"])
    check(diffs, "pdf_processor.roles", [item["role"] for item in documents],
          ["Invoice", "Itinerary"])
    for item in documents:
        check(diffs, f"pdf_processor.{item['original_name']}.error_reason", item["error_reason"], "")
        check(diffs, f"pdf_processor.{item['original_name']}.fallback",
              item["fields"].get("data_source") == "Fallback_Local_Regex", False)
    check(diffs, "pipeline.mismatches", mismatches, [])
    for key, wanted in (("imap_error", None), ("run_error", None), ("excel_error", None),
                        ("unseen_total", approved["unseen_total"]),
                        ("invoice_keyword_mails", approved["selected_mails"])):
        check(diffs, f"stats.{key}", stats.get(key), wanted)
    check(diffs, "imap.no_store", any("STORE" in event for event in fake.events), False)
    debug = Path(cfg.debug_file)
    excel = Path(cfg.summary_excel_file)
    data = json.loads(debug.read_text(encoding="utf-8")) if debug.is_file() else []
    check(diffs, "JSON.mail_count", len(data), 1)
    if len(data) == 1:
        mail = data[0]
        check(diffs, "JSON.mail_id", mail.get("mail_id"), "Msg1")
        check(diffs, "JSON.mail_uid", mail.get("mail_uid"), str(state["uid"]))
        check(diffs, "JSON.subject", mail.get("subject"), state["subject"])
        check(diffs, "JSON.mail_date", mail.get("mail_date"), "2026-09-29 09:30")
        check(diffs, "JSON.file_count", len(mail.get("files", [])), approved["file_count"])
        check(diffs, "JSON.record_count", len(mail.get("records", [])), approved["record_count"])
        check(diffs, "JSON.itinerary_count", len(mail.get("itineraries", [])), approved["itinerary_count"])
        check(diffs, "JSON.dedup_skipped", mail.get("dedup_skipped"), [])
        check(diffs, "JSON.error_reason", mail.get("error_reason"), "")
        invoice_expected = approved["excel_field_checks"]["发票.pdf"]
        for index, name in enumerate(("发票.pdf", "行程单.pdf"), 1):
            if index > len(mail.get("files", [])):
                break
            meta = mail["files"][index - 1]
            seq = f"{index:02d}"
            check(diffs, f"JSON.files[{index - 1}].file_seq", meta.get("file_seq"), seq)
            check(diffs, f"JSON.files[{index - 1}].file_type", meta.get("file_type"), "PDF")
            check(diffs, f"JSON.files[{index - 1}].original_name", meta.get("original_name"), name)
            check(diffs, f"JSON.files[{index - 1}].exported_name", meta.get("exported_name"), "")
            internal = meta.get("internal_name", "")
            check(diffs, f"JSON.files[{index - 1}].internal_prefix",
                  internal.startswith(f"Msg1_{seq}_") and len(internal) <= 120, True)
            physical = Path(cfg.extract_dir) / internal
            check(diffs, f"file[{index}].sha256", sha(physical) if physical.is_file() else "missing",
                  inventory[index]["sha256"])
        if mail.get("itineraries"):
            trip = mail["itineraries"][0]
            check(diffs, "JSON.itineraries[0]", trip,
                  {"record_id": "Msg1-02", "match_amount": invoice_expected["total_amount"],
                   "travel_time": invoice_expected["travel_time"],
                   "origin_dest": invoice_expected["origin_dest"]})
        if mail.get("records"):
            record = mail["records"][0]
            fields = {key: invoice_expected[key] for key in
                      ("invoice_date", "invoice_number", "total_amount", "invoice_type",
                       "receipt_type", "seller", "travel_time", "origin_dest")}
            check(diffs, "JSON.records[0]", record,
                  {"record_id": "Msg1-01", "kind": "a类", "fields": fields,
                   "source": {"source_file_seq": ["01"]},
                   "matched_itinerary": {"file_seq": "02",
                                         "match_amount": invoice_expected["total_amount"]},
                   "status": "complete"})
    if not excel.is_file():
        diffs.append({"path": "Excel.file", "expected": "exists", "actual": "missing"})
    else:
        book = load_workbook(excel, data_only=True)
        sheet = book.active
        check(diffs, "Excel.headers", [sheet.cell(1, col).value for col in range(1, 9)],
              [label for label, _ in SUMMARY_COLUMNS])
        check(diffs, "Excel.row_count", sheet.max_row - 1, approved["excel_rows"])
        expected_row = approved["excel_field_checks"]["发票.pdf"]
        for col, key in enumerate(("invoice_date", "invoice_number", "total_amount",
                                   "invoice_type", "receipt_type", "travel_time",
                                   "seller", "origin_dest"), 1):
            cell = sheet.cell(2, col)
            wanted = expected_row[key]
            if key == "total_amount":
                check(diffs, "Excel.amount", Decimal(str(cell.value)), Decimal(wanted))
                check(diffs, "Excel.amount_format", cell.number_format, "0.00")
                check(diffs, "Excel.amount_type", cell.data_type, "n")
            else:
                check(diffs, f"Excel.{key}", cell.value, wanted or None)
        book.close()
    source_names = ("pipeline.py", "pdf_processor.py", "ai_client.py", "attachment_processor.py",
                    "itinerary_matcher.py", "dedup_arbitrator.py", "field_validator.py",
                    "data_model.py", "excel_writer.py", "stage_ports.py",
                    "03 tests/T2/scripts/run_c05_online_chain.py")
    report = {"case": "C05/v1", "scope": "approved synthetic EML and PDFs, actual PDF/AI/T2 modules; IMAP fixture",
              "model": model, "online_request_cap": 2, "online_calls": calls,
              "source_sha256": {path: sha(ROOT / path) for path in source_names},
              "input_sha256": {item["path"]: item["sha256"] for item in inventory},
              "output_sha256": {name: sha(path) for name, path in
                                (("task_debug.json", debug), ("summary.xlsx", excel))
                                if path.is_file()},
              "imap_events": fake.events, "pdf_documents": documents,
              "stats": {key: value for key, value in stats.items()
                        if key not in ("dashscope_api_key",)},
              "passed": not diffs, "differences": diffs}
    result = RUN / "report.json"
    result.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"report": str(result), "passed": report["passed"],
                      "online_statuses": [call["status"] for call in calls],
                      "differences": diffs}, ensure_ascii=False))
    return 0 if report["passed"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
