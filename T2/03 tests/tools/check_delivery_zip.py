"""Read-only consistency check for one downloaded delivery ZIP.

This checks structure and links between outputs. It cannot establish whether
the extracted business values match the original PDF or an approved answer.
"""
from __future__ import annotations

import argparse
import io
import json
import re
import sys
import zipfile
from datetime import date, datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path

from openpyxl import load_workbook

HEADERS = ("发票日期", "发票号码", "报销金额", "发票类型", "票面类型",
           "出行时间", "销售方", "出发地/目的地")
FIELD_KEYS = ("invoice_date", "invoice_number", "total_amount", "invoice_type",
              "receipt_type", "seller")
MISSING = {"", "Not Found", "None"}


def amount(value: object) -> Decimal | None:
    try:
        parsed = Decimal(str(value))
        return parsed if parsed.is_finite() else None
    except (InvalidOperation, TypeError, ValueError):
        return None


def name_amount(value: object) -> str:
    parsed = amount(value)
    if (parsed is None or parsed <= 0 or parsed.as_tuple().exponent < -4
            or parsed.adjusted() + 1 > 14):
        return "金额缺失"
    return format(parsed, "f") + "元"


def name_date(value: object, *, travel: bool) -> str:
    raw = str(value or "")
    try:
        if travel:
            if not re.fullmatch(r"\d{4}-\d{2}-\d{2} \d{2}:\d{2}(?::\d{2})?", raw):
                return "日期缺失"
            datetime.strptime(raw, "%Y-%m-%d %H:%M:%S" if len(raw) == 19 else "%Y-%m-%d %H:%M")
            return raw[:10]
        if not re.fullmatch(r"\d{4}-\d{2}-\d{2}", raw):
            return "日期缺失"
        date.fromisoformat(raw)
        return raw
    except ValueError:
        return "日期缺失"


def compare_name(name: str, expected_date: str, expected_amount: str,
                 where: str, errors: list[str]) -> None:
    stem = re.sub(r"\(\d+\)$", "", name[:-4]) if name.lower().endswith(".pdf") else ""
    expected_tail = f"+{expected_date}+{expected_amount}"
    if not stem.endswith(expected_tail):
        errors.append(f"{where}: 导出名日期/金额应以 {expected_tail} 结尾，实际 {name}")


def check_json(mails: object, pdf_paths: set[str], root: str) -> tuple[list[str], list[str], list[dict]]:
    errors: list[str] = []
    warnings: list[str] = []
    records: list[dict] = []
    if not isinstance(mails, list):
        return ["task_debug.json 顶层必须是邮件数组"], warnings, records
    mail_ids: set[str] = set()
    uids: set[str] = set()
    exported: set[str] = set()
    for index, mail in enumerate(mails, 1):
        at = f"JSON[{index}]"
        if not isinstance(mail, dict):
            errors.append(f"{at}: 邮件条目不是对象")
            continue
        mail_id = str(mail.get("mail_id") or "")
        uid = str(mail.get("mail_uid") or "")
        if mail_id != f"Msg{index}" or mail_id in mail_ids:
            errors.append(f"{at}.mail_id: 应为 Msg{index} 且本轮唯一，实际 {mail_id}")
        if not uid.isdigit() or uid in uids:
            errors.append(f"{at}.mail_uid: 缺失、非数字或重复")
        mail_ids.add(mail_id)
        uids.add(uid)
        if mail.get("error_reason"):
            warnings.append(f"{at}: error_reason 非空，需查看该邮件的处理原因")
        files = mail.get("files")
        itineraries = mail.get("itineraries")
        invoice_records = mail.get("records")
        if not all(isinstance(value, list) for value in (files, itineraries, invoice_records)):
            errors.append(f"{at}: files、itineraries、records 必须是数组")
            continue
        by_seq: dict[str, dict] = {}
        for number, file in enumerate(files, 1):
            place = f"{at}.files[{number}]"
            if not isinstance(file, dict):
                errors.append(f"{place}: 文件条目不是对象")
                continue
            seq = str(file.get("file_seq") or "")
            if seq != f"{number:02d}" or seq in by_seq:
                errors.append(f"{place}.file_seq: 不连续或重复")
            by_seq[seq] = file
            for key in ("file_type", "original_name", "internal_name", "exported_name",
                        "pre_read_type", "pre_read_invoice_no", "pre_read_amount"):
                if key not in file:
                    errors.append(f"{place}.{key}: 缺失")
            name = str(file.get("exported_name") or "")
            if name:
                if file.get("file_type") not in ("PDF", "OFD", "图片"):
                    errors.append(f"{place}.exported_name: 非票据来源却标为已交付")
                if Path(name).name != name or "\\" in name or not name.lower().endswith(".pdf"):
                    errors.append(f"{place}.exported_name: 不是安全的 PDF 文件名")
                pdf_path = f"{root}/用户票据/{name}"
                if pdf_path not in pdf_paths:
                    errors.append(f"{place}.exported_name: ZIP 中找不到 {name}")
                exported.add(pdf_path)
        trip_by_seq: dict[str, dict] = {}
        for number, trip in enumerate(itineraries, 1):
            place = f"{at}.itineraries[{number}]"
            if not isinstance(trip, dict):
                errors.append(f"{place}: 行程单条目不是对象")
                continue
            record_id = str(trip.get("record_id") or "")
            seq = record_id[len(mail_id) + 1:] if record_id.startswith(mail_id + "-") else ""
            if not seq or seq not in by_seq or seq in trip_by_seq:
                errors.append(f"{place}.record_id: 无法唯一反查本邮件文件")
            trip_by_seq[seq] = trip
            for key in ("match_amount", "travel_time", "origin_dest"):
                if key not in trip:
                    errors.append(f"{place}.{key}: 缺失")
        matched: set[str] = set()
        record_ids: set[str] = set()
        for number, record in enumerate(invoice_records, 1):
            place = f"{at}.records[{number}]"
            if not isinstance(record, dict):
                errors.append(f"{place}: 发票记录不是对象")
                continue
            record_id = str(record.get("record_id") or "")
            if record_id in record_ids:
                errors.append(f"{place}.record_id: 重复")
            record_ids.add(record_id)
            source = (record.get("source") or {}).get("source_file_seq") or []
            if (not isinstance(source, list) or not source
                    or any(not isinstance(seq, str) for seq in source)
                    or len(source) != len(set(source))
                    or any(str(seq) not in by_seq for seq in source)
                    or record_id != f"{mail_id}-{source[0] if source else ''}"):
                errors.append(f"{place}.source: 来源序号不能唯一反查文件")
            fields = record.get("fields")
            if not isinstance(fields, dict):
                errors.append(f"{place}.fields: 不是对象")
                continue
            for key in FIELD_KEYS:
                if key not in fields:
                    errors.append(f"{place}.fields.{key}: 缺失")
            status = record.get("status")
            if status not in ("complete", "人工待核"):
                errors.append(f"{place}.status: 未知状态 {status}")
            elif status == "人工待核":
                warnings.append(f"{place}: 人工待核，不能算自动验收通过")
            link = record.get("matched_itinerary")
            trip = None
            if link is not None:
                if not isinstance(link, dict):
                    errors.append(f"{place}.matched_itinerary: 不是对象")
                else:
                    seq = str(link.get("file_seq") or "")
                    trip = trip_by_seq.get(seq)
                    if trip is None or seq in matched:
                        errors.append(f"{place}.matched_itinerary: 行程单缺失或被多票重复关联")
                    matched.add(seq)
                    if trip is not None:
                        left, right = amount(link.get("match_amount")), amount(trip.get("match_amount"))
                        if left != right and not (left is None and right is None):
                            errors.append(f"{place}.matched_itinerary.match_amount: 与行程单记录不一致")
                        for field in ("travel_time", "origin_dest"):
                            if fields.get(field) != trip.get(field):
                                errors.append(f"{place}.fields.{field}: 与所关联行程单不一致")
            if record.get("kind") == "a类" and trip is None:
                warnings.append(f"{place}: a 类发票没有有效行程单关联，需核对状态和原因")
            if record.get("kind") == "b类" and link is not None:
                errors.append(f"{place}: b 类发票不应关联行程单")
            if status == "complete" and any(str(fields.get(key) or "") in MISSING
                                            for key in FIELD_KEYS if key != "invoice_type"):
                warnings.append(f"{place}: complete 记录仍有必需字段缺失")
            records.append({"where": place, "fields": fields})
            if source and str(source[0]) in by_seq:
                filename = str(by_seq[str(source[0])].get("exported_name") or "")
                if filename:
                    date_value = (name_date(trip.get("travel_time"), travel=True) if trip is not None
                                  else name_date(fields.get("invoice_date"), travel=False))
                    compare_name(filename, date_value, name_amount(fields.get("total_amount")),
                                 place, errors)
        for seq, trip in trip_by_seq.items():
            if seq not in by_seq:
                continue
            filename = str(by_seq[seq].get("exported_name") or "")
            if not filename:
                continue
            owner = next((record for record in invoice_records
                          if isinstance(record, dict) and isinstance(record.get("matched_itinerary"), dict)
                          and str(record["matched_itinerary"].get("file_seq")) == seq), None)
            chosen_amount = ((owner.get("fields") or {}).get("total_amount") if owner
                             else trip.get("match_amount"))
            compare_name(filename, name_date(trip.get("travel_time"), travel=True),
                         name_amount(chosen_amount), f"{at}.itineraries[{seq}]", errors)
    for pdf_path in sorted(pdf_paths - exported):
        errors.append(f"ZIP: PDF 未被任何 files[].exported_name 引用：{pdf_path}")
    return errors, warnings, records


def check_excel(content: bytes, records: list[dict]) -> tuple[list[str], list[str], int]:
    errors: list[str] = []
    warnings: list[str] = []
    try:
        book = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        sheet = book.active
        rows = sheet.iter_rows(values_only=True)
        headers = tuple(next(rows))
        if headers != HEADERS:
            errors.append(f"Excel 表头/列顺序不符：{headers}")
        actual_rows = list(rows)
        book.close()
    except Exception as exc:
        return [f"Excel 无法读取：{type(exc).__name__}: {exc}"], warnings, 0
    by_number: dict[str, list[dict]] = {}
    for record in records:
        number = str(record["fields"].get("invoice_number") or "")
        if number not in MISSING:
            by_number.setdefault(number, []).append(record)
    if all(len(items) == 1 for items in by_number.values()):
        eligible = [record for record in records if not all(
            str(record["fields"].get(key) or "") in MISSING
            for key in ("invoice_date", "invoice_number", "total_amount"))]
        if len(actual_rows) != len(eligible):
            errors.append(f"Excel 业务行数 {len(actual_rows)} 与 JSON 可入表发票数 {len(eligible)} 不一致")
    else:
        warnings.append("存在重复票号；JSON 不记录全部去重写表决策，行数需对照已审核预期")
    compared = 0
    for row_no, row in enumerate(actual_rows, 2):
        if len(row) != len(HEADERS):
            errors.append(f"Excel 第 {row_no} 行列数不为 8")
            continue
        number = str(row[1] or "")
        candidates = by_number.get(number, [])
        if len(candidates) != 1:
            warnings.append(f"Excel 第 {row_no} 行票号无法唯一定位 JSON 记录；需依据已审核预期核对")
            continue
        compared += 1
        fields = candidates[0]["fields"]
        expected = (fields.get("invoice_date"), fields.get("invoice_number"),
                    fields.get("total_amount"), fields.get("invoice_type"),
                    fields.get("receipt_type"), fields.get("travel_time"),
                    fields.get("seller"), fields.get("origin_dest"))
        for col, (got, want) in enumerate(zip(row, expected), 1):
            if col == 3:
                parsed = amount(want)
                if parsed is not None and isinstance(got, (int, float)) and not isinstance(got, bool):
                    if Decimal(str(got)) != parsed.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP):
                        errors.append(f"Excel 第 {row_no} 行报销金额与 JSON 不一致")
                    continue
            if (got or "") != (want or ""):
                errors.append(f"Excel 第 {row_no} 行第 {col} 列与 {candidates[0]['where']} 不一致")
    return errors, warnings, compared


def check_bundle(path: Path) -> tuple[list[str], list[str], dict]:
    errors: list[str] = []
    warnings: list[str] = []
    counts: dict = {}
    try:
        with zipfile.ZipFile(path) as archive:
            names = archive.namelist()
            if len(names) != len(set(names)):
                errors.append("ZIP 内有重名条目")
            damaged = archive.testzip()
            if damaged:
                errors.append(f"ZIP 内容校验失败：{damaged}")
            json_names = [name for name in names if name.count("/") == 1 and name.endswith("/task_debug.json")]
            if len(json_names) != 1:
                return ["ZIP 中应恰有一份顶层 task_debug.json"], warnings, counts
            root = json_names[0].split("/", 1)[0]
            excel_name = f"{root}/发票信息汇总表.xlsx"
            if excel_name not in names:
                errors.append("ZIP 缺少发票信息汇总表.xlsx")
            pdf_paths = {name for name in names if name.startswith(f"{root}/用户票据/")
                         and name.lower().endswith(".pdf")}
            allowed = {json_names[0], excel_name} | pdf_paths
            for name in names:
                if name not in allowed and not name.endswith("/"):
                    errors.append(f"ZIP 含未约定文件：{name}")
            for pdf_path in pdf_paths:
                if b"%PDF-" not in archive.read(pdf_path)[:1024]:
                    errors.append(f"ZIP 中不是可识别 PDF：{pdf_path}")
            mails = json.loads(archive.read(json_names[0]).decode("utf-8"))
            json_errors, json_warnings, records = check_json(mails, pdf_paths, root)
            errors.extend(json_errors)
            warnings.extend(json_warnings)
            counts = {"mails": len(mails) if isinstance(mails, list) else 0,
                      "invoice_records": len(records), "delivered_pdfs": len(pdf_paths)}
            if excel_name in names:
                excel_errors, excel_warnings, compared = check_excel(archive.read(excel_name), records)
                errors.extend(excel_errors)
                warnings.extend(excel_warnings)
                counts["excel_rows_compared"] = compared
    except (OSError, zipfile.BadZipFile, UnicodeError, json.JSONDecodeError) as exc:
        errors.append(f"无法检查交付包：{type(exc).__name__}: {exc}")
    return errors, warnings, counts


def main() -> int:
    parser = argparse.ArgumentParser(description="只读检查交付 ZIP 的结构、关联和 JSON/Excel 一致性")
    parser.add_argument("zip_path", type=Path, help="本地下载的 ZIP 文件路径")
    args = parser.parse_args()
    errors, warnings, counts = check_bundle(args.zip_path)
    print("检查范围：", counts)
    for item in errors:
        print("错误：", item)
    for item in warnings:
        print("待核：", item)
    print(f"结果：{len(errors)} 项结构/关联错误，{len(warnings)} 项待核。"
          "此结果不证明字段与原始票面相符。")
    return 1 if errors else 0


if __name__ == "__main__":
    sys.exit(main())
