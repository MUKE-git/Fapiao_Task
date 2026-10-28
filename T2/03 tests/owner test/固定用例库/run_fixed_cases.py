"""Run all frozen mandatory fixtures through the offline main flow.

PDF extraction is deliberately supplied by each variant's extract.json. The
result proves dispatch, attachment handling, record assembly and T2-3 matching,
not the production PDF/AI recognizer. Raw PDF and MIME integrity is checked too.
"""
from __future__ import annotations

import argparse
import email
import hashlib
import json
import os
import shutil
import sys
import tempfile
import types
from contextlib import ExitStack
from decimal import Decimal, InvalidOperation
from email import policy as email_policy
from email.message import EmailMessage
from pathlib import Path
from unittest.mock import patch
from types import SimpleNamespace

import pdfplumber
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parent
PROJECT = Path(os.environ.get("FIXED_CASE_PROJECT_ROOT", str(ROOT.parents[2]))).resolve()
sys.path.insert(0, str(PROJECT))

# This run never calls the AI. Stub its SDK so the offline test cannot make a
# network request and can also run without the optional production package.
offline_dashscope = types.ModuleType("dashscope")
class OfflineGeneration:
    @staticmethod
    def call(*_args, **_kwargs):
        raise AssertionError("AI call is forbidden in the controlled offline run")
offline_dashscope.Generation = OfflineGeneration
offline_dashscope.api_key = ""
sys.modules["dashscope"] = offline_dashscope

from config import RunConfig, SUMMARY_COLUMNS  # noqa: E402
from pipeline import run_pipeline  # noqa: E402
from processor_contracts import DocumentResult  # noqa: E402
from attachment_processor import local_inspect_pdf  # noqa: E402
import ai_client  # noqa: E402


class FakeImap:
    def __init__(self, unseen_uids: list[str], raw_by_uid: dict[str, bytes], messages: int):
        self.unseen_uids = unseen_uids
        self.raw_by_uid = raw_by_uid
        self.messages = messages
        self.events: list[str] = []
        self.header_fetched_uids: list[str] = []
        self.full_fetched_uids: list[str] = []

    def login(self, *_args):
        self.events.append("LOGIN")
        return "OK", []

    def select(self, *_args):
        self.events.append("SELECT INBOX")
        return "OK", [str(self.messages).encode("ascii")]

    def uid(self, command, *_args):
        if command == "SEARCH":
            self.events.append("UID SEARCH UNSEEN")
            return "OK", [" ".join(self.unseen_uids).encode("ascii")]
        if command != "FETCH":
            raise AssertionError(f"unexpected IMAP command {command}")
        uid = _args[0].decode("ascii") if isinstance(_args[0], bytes) else str(_args[0])
        if uid not in self.raw_by_uid:
            raise AssertionError(f"unexpected UID FETCH {uid}")
        raw = self.raw_by_uid[uid]
        query = _args[1]
        self.events.append(f"UID FETCH {query}")
        if "HEADER.FIELDS" in query:
            self.header_fetched_uids.append(uid)
            source = email.message_from_bytes(raw)
            header = EmailMessage()
            header["Subject"] = source["Subject"]
            header["Date"] = source["Date"]
            return "OK", [(b"header", header.as_bytes())]
        self.full_fetched_uids.append(uid)
        return "OK", [(b"body", raw)]

    def status(self, *_args):
        self.events.append("STATUS INBOX")
        return "OK", [f"INBOX (MESSAGES {self.messages} UNSEEN {len(self.unseen_uids)})".encode()]

    def logout(self):
        self.events.append("LOGOUT")
        return "BYE", []


def sha256(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def integrity(folder: Path, inventory: list[dict], extracted: list[dict]) -> list[str]:
    errors = []
    for item in inventory:
        path = folder / item["path"]
        if not path.is_file():
            errors.append(f"missing input {item['path']}")
            continue
        content = path.read_bytes()
        if sha256(content) != item["sha256"] or len(content) != item["bytes"]:
            errors.append(f"hash/size mismatch {item['path']}")
    raw_items = [item for item in inventory if item["path"].lower().endswith(".eml")]
    for raw_item in raw_items:
        message = email.message_from_bytes((folder / raw_item["path"]).read_bytes())
        attached = [part for part in message.walk() if part.get_content_disposition() == "attachment"]
        direct = [item for item in inventory if item["source"] == "MIME"
                  and item.get("mail_uid") == raw_item.get("mail_uid")]
        if len(attached) != len(direct):
            errors.append(f"MIME attachment count mismatch {raw_item['path']}")
        for part, item in zip(attached, direct):
            if sha256(part.get_payload(decode=True) or b"") != item["sha256"]:
                errors.append(f"MIME payload mismatch {item['name']}")
            if part.get_filename() != item["name"]:
                errors.append(f"MIME filename mismatch {item['name']}")
    for item in extracted:
        candidates = [entry for entry in inventory if entry["name"] == item["original_name"]]
        if not candidates:
            errors.append(f"missing extracted attachment {item['original_name']}")
            continue
        attachment = folder / candidates[0]["path"]
        if item["role"] == "Unknown":
            try:
                with pdfplumber.open(attachment) as pdf:
                    _ = pdf.pages[0].extract_text()
                errors.append("intentionally corrupt PDF unexpectedly readable")
            except Exception:
                pass
            continue
        try:
            with pdfplumber.open(attachment) as pdf:
                text = "\n".join(page.extract_text() or "" for page in pdf.pages)
            if not text.strip():
                errors.append(f"PDF text empty {item['original_name']}")
            for value in item["fields"].values():
                if value not in (None, "Not Found") and str(value) not in text:
                    errors.append(f"PDF missing field value {item['original_name']}: {value}")
        except Exception as exc:
            errors.append(f"PDF open failed {item['original_name']}: {type(exc).__name__}")
    return errors


def inspect_json(data: list[dict], expected: dict) -> dict:
    files = [item for mail in data for item in mail.get("files", [])]
    records = [item for mail in data for item in mail.get("records", [])]
    itineraries = [item for mail in data for item in mail.get("itineraries", [])]
    duplicates = [item for mail in data for item in mail.get("dedup_skipped", [])]
    match = {}
    fields = {}
    itinerary_names = []
    records_by_uid = {}
    for mail in data:
        by_seq = {item["file_seq"]: item["original_name"] for item in mail.get("files", [])}
        uid = str(mail.get("mail_uid", ""))
        records_by_uid[uid] = []
        for record in mail.get("records", []):
            seq = (record.get("source") or {}).get("source_file_seq", [""])[0]
            name = by_seq.get(seq, f"unresolved:{seq}")
            fields[name] = record["fields"]
            records_by_uid[uid].append(name)
            target = record.get("matched_itinerary")
            if target:
                match[name] = by_seq.get(target["file_seq"], f"unresolved:{target['file_seq']}")
        itinerary_names.extend(by_seq.get(item["record_id"].split("-")[-1], "")
                               for item in mail.get("itineraries", []))
    return {"file_count": len(files), "record_count": len(records),
            "itinerary_count": len(itineraries),
            "duplicate_skipped": len(duplicates),
            "matched_by_invoice": match, "fields_by_invoice": fields,
            "all_itinerary_names": itinerary_names,
            "records_by_uid": records_by_uid}


def check_route_transfer(expected: dict, by_name: dict, observed: dict,
                         excel_fields: list[dict]) -> list[str]:
    """Check copied routes using frozen expected pairings, never actual pairings."""
    errors = []
    for name, doc in by_name.items():
        if doc.get("role") != "Invoice" or doc["fields"].get("receipt_type") != "客运服务费":
            continue
        target = expected["matched_by_invoice"].get(name)
        # The JSON serializer omits route fields when no itinerary is matched.
        wanted = by_name[target]["fields"]["origin_dest"] if target else None
        actual = observed["fields_by_invoice"].get(name, {}).get("origin_dest")
        if actual != wanted:
            errors.append(f"route_json {name}: expected {wanted!r}, actual {actual!r}")
        number = doc["fields"]["invoice_number"]
        rows = [row for row in excel_fields if str(row.get("invoice_number")) == str(number)]
        wanted_excel = wanted if target else ""
        if len(rows) != 1 or (rows[0].get("origin_dest") or "") != wanted_excel:
            errors.append(f"route_excel {name}: expected one row with {wanted_excel!r}, actual {rows!r}")
    return errors


def execute_variant(case_id: str, code: str, *, real_pdf: bool = False) -> dict:
    folder = ROOT / "cases" / case_id / "variants" / code
    state = json.loads((folder / "imap_state.json").read_text(encoding="utf-8"))
    expected = json.loads((folder / "expected.json").read_text(encoding="utf-8"))["expected"]
    inventory = json.loads((folder / "input_inventory.json").read_text(encoding="utf-8"))
    extracted = json.loads((folder / "extract.json").read_text(encoding="utf-8"))["documents"]
    errors = integrity(folder, inventory, extracted)
    by_name = {item["original_name"]: item for item in extracted}
    calls: list[str] = []
    pre_read_snapshots: list[dict] = []
    ai_calls: list[dict] = []

    def observed_pre_read(path):
        # Forward to the production implementation and return its result unchanged.
        local_type, text = local_inspect_pdf(path)
        calls.append(Path(path).name)
        pre_read_snapshots.append({"file": Path(path).name,
                                   "local_type": local_type, "text": text})
        return local_type, text

    def offline_ai(*_args, **kwargs):
        ai_calls.append({"model": str(kwargs.get("model", ""))})
        return SimpleNamespace(status_code=401, code="OFFLINE_FIXTURE",
                               message="simulated AI unavailable")

    def controlled_pdf(file, *_args):
        name = file.meta.original_name
        calls.append(name)
        item = by_name.get(name)
        if item is None:
            raise AssertionError(f"unexpected PDF {name}")
        return DocumentResult(file_seq=file.meta.file_seq, role=item["role"],
                              fields=item["fields"],
                              source_file_seq=[file.meta.file_seq],
                              error_reason=item["error_reason"])

    if "mails" in state:
        raw_by_uid = {str(item["uid"]): (folder / item["raw_path"]).read_bytes()
                      for item in state["mails"]}
    else:
        raw_path = folder / "00_raw.eml"
        raw_by_uid = ({str(state["uid"]): raw_path.read_bytes()}
                      if raw_path.exists() else {})
    declared_mails = state.get("mails") or (
        [{"uid": state["uid"], "subject": state["subject"]}]
        if raw_by_uid and state.get("subject") is not None else []
    )
    for item in declared_mails:
        uid = str(item["uid"])
        if uid not in raw_by_uid:
            errors.append(f"missing EML for UID {uid}")
            continue
        parsed = email.message_from_bytes(raw_by_uid[uid], policy=email_policy.default)
        decoded_subject = str(parsed.get("Subject", ""))
        if decoded_subject != item["subject"]:
            errors.append(f"subject mismatch UID {uid}: {decoded_subject!r} != {item['subject']!r}")
    unseen_uids = [str(uid) for uid in state["unseen_uids"]]
    if any(uid not in raw_by_uid for uid in unseen_uids):
        raise ValueError(f"missing EML for one of {unseen_uids}")
    fake = FakeImap(unseen_uids, raw_by_uid, int(state["messages"]))
    temporary_root = ROOT / "result"
    temporary_root.mkdir(exist_ok=True)
    with tempfile.TemporaryDirectory(prefix=f"fixed-{case_id}-{code}-",
                                     dir=temporary_root) as tmp:
        cfg = RunConfig("imap.fixture.test", "fixture", "no-real-password",
                        output_root=tmp, dashscope_api_key="offline-fixture-only")
        with ExitStack() as stack:
            stack.enter_context(patch("pipeline.imaplib.IMAP4_SSL", return_value=fake))
            if real_pdf:
                # pipeline.process_pdf remains the actual product function.
                assert sys.modules["pipeline"].process_pdf is sys.modules["pdf_processor"].process_pdf
                stack.enter_context(patch("pdf_processor.local_inspect_pdf",
                                          side_effect=observed_pre_read))
                stack.enter_context(patch.object(OfflineGeneration, "call", side_effect=offline_ai))
                stack.enter_context(patch.object(ai_client, "__file__", str(Path(tmp) / "ai_client.py")))
            else:
                stack.enter_context(patch("pipeline.process_pdf", side_effect=controlled_pdf))
            _, stats = run_pipeline(cfg)
        debug = Path(cfg.debug_file)
        data = json.loads(debug.read_text(encoding="utf-8")) if debug.exists() else []
        workbook = Path(cfg.summary_excel_file)
        excel_fields = []
        if workbook.exists():
            book = load_workbook(workbook, read_only=True, data_only=True)
            sheet = book.active
            rows = max(sheet.max_row - 1, 0)
            excel_fields = [dict(zip([key for _, key in SUMMARY_COLUMNS], row))
                            for row in sheet.iter_rows(min_row=2, values_only=True)]
            book.close()
        else:
            rows = 0
        observed = inspect_json(data, expected)
        observed.update({"unseen_total": stats["unseen_total"],
                         "selected_mails": stats["invoice_keyword_mails"],
                         "excel_rows": rows, "pdf_calls": calls,
                         "excel_fields": excel_fields,
                         "alerts": [flag for alert in stats["not_found_alerts"]
                                    for flag in alert.get("flags", [])],
                         "imap_events": fake.events, "imap_error": stats.get("imap_error")})
        observed["mail_checks"] = {
            "header_fetched_uids": fake.header_fetched_uids,
            "full_fetched_uids": fake.full_fetched_uids,
            "records_by_uid": observed["records_by_uid"],
            "pdf_calls": calls,
        }
        route_errors = check_route_transfer(expected, by_name, observed, excel_fields)
        errors.extend(route_errors)
        observed["route_transfer_checked"] = any(
            doc.get("role") == "Invoice" and doc["fields"].get("receipt_type") == "客运服务费"
            for doc in by_name.values())
        observed["route_transfer_errors"] = route_errors
        if real_pdf:
            pdf_count = sum(file["file_type"] == "PDF"
                            for mail in data for file in mail.get("files", []))
            execution_checks = {
                "registered_pdfs_reached_real_pre_read": len(pre_read_snapshots) == pdf_count,
                "each_pdf_reached_offline_ai": len(ai_calls) == pdf_count,
                "mail_input_completed": not stats.get("imap_error"),
            }
            errors.extend(f"execution check failed: {key}" for key, ok in execution_checks.items() if not ok)
            observed["execution_checks"] = execution_checks
            observed["pre_read_snapshots"] = pre_read_snapshots
            observed["offline_ai_calls"] = ai_calls
            observed["final_json"] = data
            observed["excel_fields"] = excel_fields
            observed["pipeline_stats"] = stats
            artifacts = ROOT / "result" / "local_flow" / case_id / code
            artifacts.mkdir(parents=True, exist_ok=True)
            (artifacts / "final.json").write_text(
                json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
            (artifacts / "pre_read.json").write_text(
                json.dumps(pre_read_snapshots, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
            if workbook.exists():
                shutil.copyfile(workbook, artifacts / "final.xlsx")
            elif (artifacts / "final.xlsx").exists():
                (artifacts / "final.xlsx").unlink()
            observed["artifact_dir"] = str(artifacts.relative_to(ROOT)).replace("\\", "/")

    for key in ("unseen_total", "selected_mails", "file_count", "record_count",
                "itinerary_count", "excel_rows", "duplicate_skipped", "matched_by_invoice"):
        if observed[key] != expected[key]:
            errors.append(f"{key}: expected {expected[key]!r}, actual {observed[key]!r}")
    for name, checks in expected["field_checks"].items():
        actual_fields = observed["fields_by_invoice"].get(name, {})
        for key, value in checks.items():
            if str(actual_fields.get(key)) != str(value):
                errors.append(f"{name}.{key}: expected {value!r}, actual {actual_fields.get(key)!r}")
    for name, checks in expected.get("excel_field_checks", {}).items():
        source = by_name.get(name)
        number = source["fields"].get("invoice_number") if source else None
        matches = [row for row in excel_fields if str(row.get("invoice_number")) == str(number)]
        if len(matches) != 1:
            errors.append(f"excel {name}: expected one row for invoice {number}, actual {len(matches)}")
            continue
        for key, wanted in checks.items():
            actual = matches[0].get(key)
            if key == "total_amount":
                try:
                    equal = Decimal(str(actual)) == Decimal(str(wanted))
                except (InvalidOperation, TypeError):
                    equal = False
            else:
                equal = str(actual or "") == str(wanted or "")
            if not equal:
                errors.append(f"excel {name}.{key}: expected {wanted!r}, actual {actual!r}")
    for needle in expected["alerts_contain"]:
        if not any(needle in flag for flag in observed["alerts"]):
            errors.append(f"alert missing: {needle}")
    for key, wanted in expected.get("mail_checks", {}).items():
        actual = observed["mail_checks"].get(key)
        if actual != wanted:
            errors.append(f"mail_checks.{key}: expected {wanted!r}, actual {actual!r}")
    if stats.get("imap_error"):
        errors.append(f"imap_error: {stats['imap_error']}")
    return {"case_id": case_id, "variant": code,
            "status": "PASS" if not errors else "FAIL",
            "errors": errors, "expected": expected, "observed": observed,
            "proof_level": ("真实邮件解析/PDF/匹配/输出整链路；邮箱固定输入、AI模拟失败；与已审核标准答案比较，运行差异待分析"
                            if real_pdf else "受控 PDF 返回值的离线主流程和 T2-3 验证；非真实 PDF/AI 识别")}


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--pdf-mode", choices=("controlled", "local"), default="controlled",
                        help="local runs the actual PDF processor with simulated AI failure")
    args = parser.parse_args()
    real_pdf = args.pdf_mode == "local"
    index = json.loads((ROOT / "index.json").read_text(encoding="utf-8"))
    results = [execute_variant(case["id"], code, real_pdf=real_pdf)
               for case in index["cases"] for code in case["variants"]]
    result_dir = ROOT / "result"
    result_dir.mkdir(exist_ok=True)
    basename = "local_flow_latest" if real_pdf else "latest"
    (result_dir / f"{basename}.json").write_text(
        json.dumps(results, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    lines = [f"# {len(index['cases'])} 条必选固定用例离线验证", "",
             "本次 PDF 字段由各例 `extract.json` 受控提供。通过只证明输入、主流程、附件处理、数据组装、T2-3 匹配及 Excel/JSON 交接；不证明真实 PDF/AI 识别。",
             "", "| 用例 | 变体 | 结果 | 差异 |", "| --- | --- | --- | --- |"]
    for item in results:
        lines.append(f"| {item['case_id']} | {item['variant']} | {item['status']} | "
                     f"{'；'.join(item['errors']) if item['errors'] else '无'} |")
    if real_pdf:
        lines[0] = "# 固定邮件到最终输出：真实本地 PDF 整链路"
        lines[2] = ("固定 EML 模拟邮箱返回；项目实际解析附件、预读 PDF、分类、字段提取、匹配并输出。"
                    "只模拟 AI 失败，不用 extract.json 替代处理结果。表中 PASS/FAIL 是与原冻结草案的比较；"
                    "草案原用于受控提取，差异须审核，不能据此宣称所有差异都是程序缺陷。")
        complete = sum(all(item["observed"]["execution_checks"].values()) for item in results)
        lines[2] += f" 实际调用检查完成 {complete}/{len(results)} 个变体。"
        lines += ["", "## 实际输出索引", "",
                  "| 用例 / 变体 | 预读 | 最终 JSON | Excel |", "| --- | --- | --- | --- |"]
        for item in results:
            rel = f"local_flow/{item['case_id']}/{item['variant']}"
            excel = f"[打开]({rel}/final.xlsx)" if (result_dir / rel / "final.xlsx").exists() else "无输出行"
            lines.append(f"| {item['case_id']} / {item['variant']} | [打开]({rel}/pre_read.json) | "
                         f"[打开]({rel}/final.json) | {excel} |")
    lines += ["", f"通过 {sum(item['status'] == 'PASS' for item in results)}/{len(results)} 个变体。",
              "18 条必选用例已由用户审核通过；运行 PASS/FAIL 独立记录。", ""]
    (result_dir / f"{basename}.md").write_text("\n".join(lines), encoding="utf-8")
    print(f"passed={sum(item['status'] == 'PASS' for item in results)}/{len(results)}")
    for item in results:
        if item["errors"]:
            print(item["case_id"], item["variant"], "; ".join(item["errors"]))
    if any(item["errors"] for item in results):
        raise SystemExit(1)


if __name__ == "__main__":
    main()
