"""Formal independent T2-3 tests for user-approved 2026-10-09 rule cases.

Expected links, statuses and inputs are authored in approved_revision_cases.json.
No expected value is copied from a production output. IMAP and first PDF
extraction are controlled; all T2 stage ports and JSON/Excel I/O are real.
"""
from __future__ import annotations

import hashlib
import json
import sys
import tempfile
from datetime import datetime
from decimal import Decimal
from email.message import EmailMessage
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[3]
sys.path.insert(0, str(ROOT))
from config import RunConfig, SUMMARY_COLUMNS  # noqa: E402
from pipeline import run_pipeline  # noqa: E402
from processor_contracts import DocumentResult  # noqa: E402

BASE = ROOT / "03 tests" / "T2-3"
ORACLE = BASE / "fixtures" / "approved_revision_cases.json"
SPEC = json.loads(ORACLE.read_text(encoding="utf-8"))
RUN = BASE / "result" / f"approved_revision_{datetime.now().strftime('%Y%m%d_%H%M%S')}"
RUN.mkdir(parents=True, exist_ok=False)
tempfile.tempdir = str(RUN.resolve())


def sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def check(diffs: list, path: str, got, wanted) -> None:
    if got != wanted:
        diffs.append({"path": path, "expected": wanted, "actual": got})


class FakeImap:
    def __init__(self, uid: str, raw: bytes):
        self.uid_text, self.raw = uid, raw
        self.events: list[str] = []

    def login(self, *_):
        self.events.append("LOGIN")
        return "OK", []

    def select(self, *_):
        self.events.append("SELECT INBOX")
        return "OK", [b"1"]

    def status(self, *_):
        self.events.append("STATUS INBOX")
        return "OK", [b"INBOX (MESSAGES 1 UNSEEN 1)"]

    def uid(self, command, *args):
        if command == "SEARCH":
            self.events.append("UID SEARCH UNSEEN")
            return "OK", [self.uid_text.encode()]
        if command != "FETCH":
            raise AssertionError(f"Unexpected IMAP command: {command}")
        uid = args[0].decode() if isinstance(args[0], bytes) else str(args[0])
        if uid != self.uid_text:
            raise AssertionError(f"Unexpected UID: {uid}")
        self.events.append(f"UID FETCH {args[1]}")
        return "OK", [(b"body", self.raw)]

    def logout(self):
        self.events.append("LOGOUT")
        return "BYE", []


def make_fields(item: dict, invoice_index: int) -> dict:
    if item["role"] == "Itinerary":
        return {"match_amount": item["amount"],
                "travel_time": "2026-10-14 10:00", "origin_dest": item["route"]}
    return {**SPEC["common_invoice"],
            "invoice_number": f"263270000007752327{invoice_index:02d}",
            "total_amount": item["amount"],
            "receipt_type": item.get("receipt_type", SPEC["common_invoice"]["receipt_type"])}


def run_case(case: dict) -> dict:
    case_dir = RUN / case["id"]
    case_dir.mkdir()
    cfg = RunConfig("imap.approved.test", "fixture", "unused", output_folder_name="output",
                    output_root=str(case_dir), dashscope_api_key="controlled-test")
    message = EmailMessage()
    message["Subject"] = SPEC["mail_subject"]
    message["Date"] = SPEC["mail_date_header"]
    message.set_content("已审核的合成邮件输入")
    fields_by_name = {}
    input_hashes = {}
    invoice_index = 0
    for index, item in enumerate(case["files"], 1):
        if item["role"] == "Invoice":
            invoice_index += 1
        fields_by_name[item["name"]] = (item["role"], make_fields(item, invoice_index))
        pdf_bytes = f"T2-3 approved {case['id']} {index} {item['name']}".encode("utf-8")
        pdf = case_dir / f"input_{index:02d}.pdf"
        pdf.write_bytes(pdf_bytes)
        input_hashes[pdf.name] = sha(pdf)
        message.add_attachment(pdf_bytes, maintype="application", subtype="pdf",
                               filename=item["name"])
    eml = case_dir / "input.eml"
    eml.write_bytes(message.as_bytes())
    input_hashes[eml.name] = sha(eml)
    fake = FakeImap(case["uid"], eml.read_bytes())

    def controlled_extract(file, *_):
        role, fields = fields_by_name[file.meta.original_name]
        return DocumentResult(file.meta.file_seq, role, fields, [file.meta.file_seq])

    with patch("pipeline.imaplib.IMAP4_SSL", return_value=fake), \
         patch("pipeline.process_pdf", side_effect=controlled_extract), \
         patch("dedup_arbitrator.Generation.call", side_effect=AssertionError("unexpected AI call")):
        mismatches, stats = run_pipeline(cfg)

    diffs = []
    check(diffs, "mismatches", mismatches, [])
    check(diffs, "stats.imap_error", stats.get("imap_error"), None)
    check(diffs, "stats.run_error", stats.get("run_error"), None)
    check(diffs, "stats.unseen_total", stats.get("unseen_total"), 1)
    check(diffs, "stats.invoice_keyword_mails", stats.get("invoice_keyword_mails"), 1)
    check(diffs, "imap.READ_ONLY_FETCH", any("STORE" in event for event in fake.events), False)
    debug = Path(cfg.debug_file)
    data = json.loads(debug.read_text(encoding="utf-8")) if debug.is_file() else []
    check(diffs, "JSON.mail_count", len(data), 1)
    if len(data) == 1:
        mail = data[0]
        check(diffs, "JSON.keys", sorted(mail), sorted(("mail_id", "mail_uid", "subject",
              "mail_date", "files", "itineraries", "records", "dedup_skipped",
              "error_reason")))
        for key, wanted in (("mail_id", "Msg1"), ("mail_uid", case["uid"]),
                            ("subject", SPEC["mail_subject"]),
                            ("mail_date", "2026-10-15 09:30")):
            check(diffs, f"JSON.{key}", mail.get(key), wanted)
        check(diffs, "JSON.dedup_skipped", mail.get("dedup_skipped"), [])
        check(diffs, "JSON.files.length", len(mail.get("files", [])), len(case["files"]))
        by_seq = {f"{i:02d}": item for i, item in enumerate(case["files"], 1)}
        for index, item in enumerate(case["files"], 1):
            if index > len(mail.get("files", [])):
                break
            seq = f"{index:02d}"
            expected_file = {"file_seq": seq, "file_type": "PDF",
                             "original_name": item["name"], "pre_read_type": "Unknown",
                             "pre_read_invoice_no": "Not Found", "pre_read_amount": "Not Found",
                             "internal_name": f"Msg1_{seq}_{item['name']}", "exported_name": ""}
            check(diffs, f"JSON.files[{index - 1}]", mail["files"][index - 1], expected_file)
            physical = Path(cfg.extract_dir) / expected_file["internal_name"]
            check(diffs, f"file[{index}].bytes_sha256",
                  sha(physical) if physical.is_file() else "missing",
                  input_hashes[f"input_{index:02d}.pdf"])
        expected_itineraries = [
            {"record_id": f"Msg1-{index:02d}", "match_amount": item["amount"],
             "travel_time": "2026-10-14 10:00", "origin_dest": item["route"]}
            for index, item in enumerate(case["files"], 1) if item["role"] == "Itinerary"
        ]
        check(diffs, "JSON.itineraries", mail.get("itineraries"), expected_itineraries)
        expected_records = []
        invoice_index = 0
        expected_excel = []
        for index, item in enumerate(case["files"], 1):
            if item["role"] != "Invoice":
                continue
            invoice_index += 1
            seq = f"{index:02d}"
            link_seq = case["expected_links"].get(seq)
            linked = by_seq.get(link_seq) if link_seq else None
            fields = make_fields(item, invoice_index)
            if linked:
                fields.update({"travel_time": "2026-10-14 10:00",
                               "origin_dest": linked["route"]})
            record = {"record_id": f"Msg1-{seq}",
                      "kind": "b类" if item.get("receipt_type") == "办公用品" else "a类",
                      "fields": fields, "source": {"source_file_seq": [seq]}}
            if linked:
                record["matched_itinerary"] = {"file_seq": link_seq,
                                                "match_amount": linked["amount"]}
            record["status"] = case["expected_statuses"][seq]
            expected_records.append(record)
            expected_excel.append([fields["invoice_date"], fields["invoice_number"],
                                   float(Decimal(item["amount"])), None,
                                   fields["receipt_type"],
                                   "2026-10-14 10:00" if linked else None,
                                   fields["seller"], linked["route"] if linked else None])
        check(diffs, "JSON.records", mail.get("records"), expected_records)
        issue_seqs = case["expected_issue_seqs"]
        if issue_seqs and not mail.get("error_reason"):
            diffs.append({"path": "JSON.error_reason", "expected": "nonempty issue",
                          "actual": mail.get("error_reason")})
        for seq in issue_seqs:
            if f"Msg1-{seq}" not in mail.get("error_reason", ""):
                diffs.append({"path": "JSON.error_reason", "expected_contains": f"Msg1-{seq}",
                              "actual": mail.get("error_reason")})
        flags = [flag for alert in stats.get("not_found_alerts", [])
                 for flag in alert.get("flags", [])]
        if case["id"] == "丁" and not any("一对一关系冲突" in flag for flag in flags):
            diffs.append({"path": "stats.not_found_alerts", "expected_contains": "一对一关系冲突",
                          "actual": flags})
        for seq in case["expected_unmatched_itineraries"]:
            if not any(f"Msg1-{seq}" in flag and "未匹配" in flag for flag in flags):
                diffs.append({"path": "stats.not_found_alerts",
                              "expected_contains": f"Msg1-{seq} 未匹配", "actual": flags})
    else:
        expected_excel = []
    excel = Path(cfg.summary_excel_file)
    if not excel.is_file():
        diffs.append({"path": "Excel", "expected": "workbook exists", "actual": "missing"})
    else:
        book = load_workbook(excel, data_only=True)
        sheet = book.active
        check(diffs, "Excel.headers", [sheet.cell(1, col).value for col in range(1, 9)],
              [name for name, _ in SUMMARY_COLUMNS])
        check(diffs, "Excel.row_count", sheet.max_row - 1, len(expected_excel))
        for row_index, row in enumerate(expected_excel, 2):
            if row_index > sheet.max_row:
                break
            for col, wanted in enumerate(row, 1):
                cell = sheet.cell(row_index, col)
                if col == 3 and isinstance(cell.value, (int, float)):
                    check(diffs, f"Excel.row[{row_index}].amount", float(cell.value), wanted)
                    check(diffs, f"Excel.row[{row_index}].amount_format", cell.number_format, "0.00")
                    check(diffs, f"Excel.row[{row_index}].amount_type", cell.data_type, "n")
                else:
                    check(diffs, f"Excel.row[{row_index}].col[{col}]", cell.value, wanted)
        book.close()
    return {"case": case["id"], "passed": not diffs, "differences": diffs,
            "input_sha256": input_hashes,
            "output_sha256": {name: sha(path) for name, path in
                              (("task_debug.json", debug), ("summary.xlsx", excel))
                              if path.is_file()},
            "imap_events": fake.events}


def main() -> int:
    results = [run_case(case) for case in SPEC["cases"]]
    sources = ("pipeline.py", "itinerary_matcher.py", "stage_ports.py", "data_model.py",
               "field_validator.py", "dedup_arbitrator.py", "attachment_processor.py",
               "03 tests/T2-3/scripts/run_approved_revision_cases.py",
               "03 tests/T2-3/fixtures/approved_revision_cases.json")
    report = {"approval": SPEC["approval"],
              "scope": "formal approved T2-3 revision cases; controlled IMAP/first PDF extraction",
              "source_sha256": {path: sha(ROOT / path) for path in sources},
              "passed": sum(item["passed"] for item in results),
              "total": len(results), "results": results}
    target = RUN / "report.json"
    target.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"report": str(target), "passed": report["passed"],
                      "total": report["total"],
                      "failed": {item["case"]: item["differences"] for item in results
                                 if not item["passed"]}}, ensure_ascii=False))
    return 0 if report["passed"] == report["total"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
