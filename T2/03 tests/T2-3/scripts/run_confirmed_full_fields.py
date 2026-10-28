"""Independent full-field check for the already approved T2-3 fixed cases.

The existing runner supplies extract.json as controlled PDF results. Expected
business fields come from that frozen input and approved expected.json, never
from a pipeline output. The original runner and fixtures are not modified.
"""
from __future__ import annotations

import hashlib
import importlib.util
import json
import shutil
import sys
import tempfile
from datetime import datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[3]
LIB = ROOT / "03 tests" / "owner test" / "固定用例库"
HERE = ROOT / "03 tests" / "T2-3"
RUN = HERE / "result" / f"full_fields_{datetime.now().strftime('%Y%m%d_%H%M%S')}"
RUN.mkdir(parents=True, exist_ok=False)
tempfile.tempdir = str(RUN.resolve())
SPECS = [("C05", "v1"), ("C13", "v1"), ("C14", "v1"), ("C15", "v1"),
         ("C16", "d049"), ("C16", "d050"), ("C16", "d051_name"),
         ("C16", "d051_no_name"), ("C17", "v1"), ("C18", "v1"),
         ("C19", "no_evidence"), ("C19", "same_amount")]


def digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def check(diffs: list, path: str, got, wanted) -> None:
    if got != wanted:
        diffs.append({"path": path, "expected": wanted, "actual": got})


def amount_equal(got, wanted) -> bool:
    try:
        return Decimal(str(got)) == Decimal(str(wanted))
    except (InvalidOperation, TypeError, ValueError):
        return str(got) == str(wanted)


def expected_status(case_id: str, variant: str, name: str, matched: bool) -> str:
    if case_id == "C17" or not matched:
        return "人工待核"
    return "complete"


def main() -> int:
    runner_path = LIB / "run_fixed_cases.py"
    spec = importlib.util.spec_from_file_location("approved_fixed_runner_full_fields", runner_path)
    if spec is None or spec.loader is None:
        raise RuntimeError("fixed case runner unavailable")
    runner = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = runner
    spec.loader.exec_module(runner)
    original_pipeline = runner.run_pipeline
    capture_dir: Path | None = None

    def capture(cfg, *args, **kwargs):
        result = original_pipeline(cfg, *args, **kwargs)
        assert capture_dir is not None
        for src, name in ((Path(cfg.debug_file), "task_debug.json"),
                          (Path(cfg.summary_excel_file), "summary.xlsx")):
            if src.is_file():
                shutil.copy2(src, capture_dir / name)
        return result

    runner.run_pipeline = capture
    reports = []
    for case_id, variant in SPECS:
        folder = LIB / "cases" / case_id / "variants" / variant
        capture_dir = RUN / case_id / variant
        capture_dir.mkdir(parents=True)
        state = json.loads((folder / "imap_state.json").read_text(encoding="utf-8"))
        approved = json.loads((folder / "expected.json").read_text(encoding="utf-8"))["expected"]
        documents = json.loads((folder / "extract.json").read_text(encoding="utf-8"))["documents"]
        baseline = runner.execute_variant(case_id, variant, real_pdf=False)
        actual = json.loads((capture_dir / "task_debug.json").read_text(encoding="utf-8"))
        diffs: list[dict] = []
        check(diffs, "existing_assertions", baseline["errors"], [])
        check(diffs, "JSON.mail_count", len(actual), 1)
        if len(actual) != 1:
            reports.append({"case": f"{case_id}/{variant}", "passed": False,
                            "differences": diffs})
            continue
        mail = actual[0]
        check(diffs, "JSON.keys", sorted(mail), sorted(("mail_id", "mail_uid", "subject",
              "mail_date", "files", "itineraries", "records", "dedup_skipped",
              "error_reason")))
        for key, wanted in (("mail_id", "Msg1"), ("mail_uid", str(state["uid"])),
                            ("subject", state["subject"]),
                            ("mail_date", "2026-09-29 09:30")):
            check(diffs, f"JSON.{key}", mail[key], wanted)
        check(diffs, "JSON.dedup_skipped", mail["dedup_skipped"], [])
        check(diffs, "JSON.files.length", len(mail["files"]), len(documents))
        check(diffs, "JSON.itineraries.length", len(mail["itineraries"]),
              sum(doc["role"] == "Itinerary" for doc in documents))
        check(diffs, "JSON.records.length", len(mail["records"]),
              sum(doc["role"] == "Invoice" for doc in documents))
        seq_by_name = {doc["original_name"]: f"{index:02d}"
                       for index, doc in enumerate(documents, 1)}
        doc_by_name = {doc["original_name"]: doc for doc in documents}
        for index, doc in enumerate(documents, 1):
            if index > len(mail["files"]):
                break
            name = doc["original_name"]
            expected_file = {"file_seq": f"{index:02d}", "file_type": "PDF",
                             "original_name": name, "pre_read_type": "Unknown",
                             "pre_read_invoice_no": "Not Found",
                             "pre_read_amount": "Not Found",
                             "internal_name": f"Msg1_{index:02d}_{name}",
                             "exported_name": ""}
            check(diffs, f"JSON.files[{index - 1}]", mail["files"][index - 1],
                  expected_file)
        expected_itineraries = [
            {"record_id": f"Msg1-{seq_by_name[doc['original_name']]}",
             "match_amount": doc["fields"]["match_amount"],
             "travel_time": doc["fields"]["travel_time"],
             "origin_dest": doc["fields"]["origin_dest"]}
            for doc in documents if doc["role"] == "Itinerary"
        ]
        check(diffs, "JSON.itineraries", mail["itineraries"], expected_itineraries)
        expected_records = []
        expected_rows = []
        for doc in documents:
            if doc["role"] != "Invoice":
                continue
            name = doc["original_name"]
            match_name = approved["matched_by_invoice"].get(name)
            match_doc = doc_by_name.get(match_name) if match_name else None
            expected_fields = dict(doc["fields"])
            if match_doc:
                expected_fields.update({key: match_doc["fields"][key]
                                        for key in ("travel_time", "origin_dest")})
            record = {"record_id": f"Msg1-{seq_by_name[name]}", "kind": "a类",
                      "fields": expected_fields,
                      "source": {"source_file_seq": [seq_by_name[name]]}}
            if match_doc:
                record["matched_itinerary"] = {
                    "file_seq": seq_by_name[match_name],
                    "match_amount": match_doc["fields"]["match_amount"]}
            record["status"] = expected_status(case_id, variant, name, bool(match_doc))
            expected_records.append(record)
            expected_rows.append({
                "invoice_date": expected_fields["invoice_date"],
                "invoice_number": expected_fields["invoice_number"],
                "total_amount": expected_fields["total_amount"],
                "invoice_type": expected_fields["invoice_type"],
                "receipt_type": expected_fields["receipt_type"],
                "travel_time": match_doc["fields"]["travel_time"] if match_doc else None,
                "seller": expected_fields["seller"],
                "origin_dest": match_doc["fields"]["origin_dest"] if match_doc else None,
            })
        check(diffs, "JSON.records", mail["records"], expected_records)
        # Exact wording is only asserted where approved expected.json prescribes a phrase.
        expected_issue = any(record["status"] != "complete" for record in expected_records)
        if expected_issue and not mail["error_reason"]:
            diffs.append({"path": "JSON.error_reason", "expected": "nonempty issue",
                          "actual": mail["error_reason"]})
        workbook = capture_dir / "summary.xlsx"
        if not workbook.is_file():
            diffs.append({"path": "Excel", "expected": "workbook exists", "actual": "missing"})
        else:
            book = load_workbook(workbook, read_only=False, data_only=True)
            sheet = book.active
            headers = [cell.value for cell in sheet[1]]
            check(diffs, "Excel.headers", headers, ["发票日期", "发票号码", "报销金额",
                  "发票类型", "票面类型", "出行时间", "销售方", "出发地/目的地"])
            check(diffs, "Excel.row_count", sheet.max_row - 1, len(expected_rows))
            keys = ("invoice_date", "invoice_number", "total_amount", "invoice_type",
                    "receipt_type", "travel_time", "seller", "origin_dest")
            for row_index, expected_row in enumerate(expected_rows, 2):
                if row_index > sheet.max_row:
                    break
                for col_index, key in enumerate(keys, 1):
                    cell = sheet.cell(row_index, col_index)
                    got, wanted = cell.value, expected_row[key]
                    if key == "total_amount" and wanted != "Not Found":
                        if isinstance(got, bool) or not isinstance(got, (int, float)) or not amount_equal(got, wanted):
                            diffs.append({"path": f"Excel.row[{row_index}].{key}",
                                          "expected": wanted, "actual": got})
                        check(diffs, f"Excel.row[{row_index}].amount_format",
                              cell.number_format, "0.00")
                    else:
                        check(diffs, f"Excel.row[{row_index}].{key}", got, wanted)
            book.close()
        reports.append({"case": f"{case_id}/{variant}", "passed": not diffs,
                        "differences": diffs,
                        "input_sha256": {name: digest(folder / name) for name in
                                         ("expected.json", "extract.json", "00_raw.eml")},
                        "output_sha256": {name: digest(capture_dir / name) for name in
                                          ("task_debug.json", "summary.xlsx")
                                          if (capture_dir / name).is_file()}})
    report = {"scope": "approved fixed C05,C13-C19; controlled PDF outputs; full JSON/Excel business fields",
              "source_sha256": {name: digest(ROOT / name) for name in
                                ("pipeline.py", "itinerary_matcher.py", "data_model.py",
                                 "field_validator.py", "03 tests/owner test/固定用例库/run_fixed_cases.py",
                                 "03 tests/T2-3/scripts/run_confirmed_full_fields.py")},
              "passed": sum(row["passed"] for row in reports), "total": len(reports),
              "results": reports}
    path = RUN / "report.json"
    path.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"report": str(path), "passed": report["passed"],
                      "total": report["total"],
                      "failed": {row["case"]: row["differences"] for row in reports
                                 if not row["passed"]}}, ensure_ascii=False))
    return 0 if report["passed"] == report["total"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
