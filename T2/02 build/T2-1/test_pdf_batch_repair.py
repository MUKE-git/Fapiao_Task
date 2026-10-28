"""PDF 行程单字段与单封异常隔离的受控回归；不连接真实邮箱或模型。"""
from __future__ import annotations

import json
import tempfile
import unittest
from decimal import Decimal
from email.message import EmailMessage
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from openpyxl import load_workbook

from ai_client import call_ai_audit_and_rename
from attachment_processor import CollectedFile
from config import RunConfig, ensure_output_dirs
from data_model import FileRecord, NOT_FOUND
from fallback_extractor import build_fallback_audit_result
from pdf_processor import process_pdf
from pipeline import process_message, run_pipeline
from processor_contracts import DocumentResult
from stage_ports import StagePorts


ROOT = Path(__file__).resolve().parents[2]
ITINERARY_PDF = (ROOT / "03 tests/owner test/固定用例库/cases/C05/variants/v1/attachments/direct/行程单.pdf")


def message() -> EmailMessage:
    mail = EmailMessage()
    mail["Subject"] = "电子发票"
    mail["Date"] = "Thu, 15 Oct 2026 08:00:00 +0800"
    mail.set_content("受控测试")
    mail.add_attachment(b"%PDF-1.4", maintype="application", subtype="pdf", filename="发票.pdf")
    return mail


class FakeImap:
    def __init__(self, uids: bytes, body: bytes):
        self.uids = uids
        self.body = body

    def login(self, *_args): return "OK", []
    def select(self, *_args): return "OK", [b"3"]
    def uid(self, command, *_args):
        return ("OK", [self.uids]) if command == "SEARCH" else ("OK", [(b"1", self.body)])
    def status(self, *_args): return "OK", [b"INBOX (MESSAGES 3 UNSEEN 3)"]
    def logout(self): return "BYE", []


class PdfBatchRepairTests(unittest.TestCase):
    def test_itinerary_record_accepts_legacy_total_from_same_file(self):
        mail = EmailMessage()
        mail["Subject"] = "电子发票"
        mail["Date"] = "Thu, 15 Oct 2026 08:00:00 +0800"
        mail.set_content("受控测试")
        mail.add_attachment(b"%PDF-1.4 itinerary", maintype="application", subtype="pdf",
                            filename="行程单.pdf")
        with tempfile.TemporaryDirectory() as root:
            cfg = RunConfig("imap.test", "user", "unused", output_root=root,
                            dashscope_api_key="test")
            ensure_output_dirs(cfg)
            result = DocumentResult("01", "Itinerary", {
                "total_amount": "38.93", "travel_time": "2026-09-25 10:00",
                "origin_dest": "杭州→上海",
            })
            with patch("pipeline.process_pdf", return_value=result):
                outcome = process_message(mail, "Msg1", "701", cfg)
        self.assertEqual(outcome.mail.itineraries[0].match_amount, Decimal("38.93"))

    def test_itinerary_record_does_not_borrow_invoice_total_from_other_file(self):
        mail = message()
        mail.add_attachment(b"%PDF-1.4 itinerary", maintype="application", subtype="pdf",
                            filename="行程单.pdf")
        with tempfile.TemporaryDirectory() as root:
            cfg = RunConfig("imap.test", "user", "unused", output_root=root,
                            dashscope_api_key="test")
            ensure_output_dirs(cfg)

            def audit(file, *_args):
                if "行程单" in file.meta.original_name:
                    return DocumentResult(file.meta.file_seq, "Itinerary", {
                        "travel_time": "2026-09-25 10:00", "origin_dest": "杭州→上海"})
                return DocumentResult(file.meta.file_seq, "Invoice", {
                    "total_amount": "38.93", "invoice_date": "2026-09-25",
                    "invoice_number": "12345678901234567890", "receipt_type": "客运服务"})

            with patch("pipeline.process_pdf", side_effect=audit):
                outcome = process_message(mail, "Msg1", "701", cfg)
        self.assertEqual(len(outcome.mail.itineraries), 1,
                         f"files={outcome.mail.files}, failed={outcome.failed_files}")
        self.assertEqual(outcome.mail.itineraries[0].match_amount, NOT_FOUND)

    def test_fallback_keeps_itinerary_amount_separate_from_invoice_amount(self):
        invoice = build_fallback_audit_result([{"identified_type": "Invoice",
                                                "text_snapshot": "价税合计（小写）￥38.93"}])
        itinerary = build_fallback_audit_result([{"identified_type": "Itinerary",
                                                  "text_snapshot": "报销金额：53.9300\n出行时间：2026-09-25 10:00"}])
        self.assertEqual(invoice["total_amount"], "38.93")
        self.assertEqual(invoice["match_amount"], NOT_FOUND)
        self.assertEqual(itinerary["match_amount"], "53.9300")
        self.assertEqual(itinerary["total_amount"], NOT_FOUND)

    def test_real_text_pdf_pre_read_stores_itinerary_amount(self):
        self.assertTrue(ITINERARY_PDF.is_file())
        with tempfile.TemporaryDirectory() as root:
            cfg = RunConfig("imap.test", "user", "unused", output_root=root,
                            dashscope_api_key="test")
            file = CollectedFile(FileRecord("01", original_name="行程单.pdf"), str(ITINERARY_PDF))
            with patch("pdf_processor.call_ai_audit_and_rename",
                       return_value=(None, "AI_OFFLINE_TEST", [], [])):
                result = process_pdf(file, cfg, "电子发票", "2026-10-15")
        self.assertEqual(result.role, "Itinerary")
        self.assertEqual(result.fields["match_amount"], "38.93")
        self.assertEqual(file.meta.pre_read_amount, Decimal("38.93"))

    def test_model_itinerary_response_uses_its_own_fields(self):
        raw = ('{"match_amount":53.9300,"travel_time":"2026-09-25 10:00",'
               '"origin_dest":"杭州→上海","file_classifications":'
               '[{"file":"Msg1_01_行程单.pdf","role":"Itinerary"}]}')
        response = SimpleNamespace(status_code=200, output=SimpleNamespace(text=raw))
        task = {"files": [{"local_path": "Msg1_01_行程单.pdf",
                           "attachment_original_name": "行程单.pdf",
                           "identified_type": "Itinerary", "text_snapshot": "报销金额：53.9300"}]}
        with patch("ai_client.Generation.call", return_value=response), \
             patch("ai_client._rename_pdfs_with_audit", return_value=([], [])):
            data, error, _, _ = call_ai_audit_and_rename(task)
        self.assertEqual(data["match_amount"], Decimal("53.9300"))
        self.assertEqual(error, "")

    def test_missing_model_itinerary_amount_is_flagged_without_invoice_field_false_alarm(self):
        raw = ('{"travel_time":"2026-09-25 10:00","origin_dest":"杭州→上海",'
               '"file_classifications":[{"file":"Msg1_01_行程单.pdf","role":"Itinerary"}]}')
        response = SimpleNamespace(status_code=200, output=SimpleNamespace(text=raw))
        task = {"files": [{"local_path": "Msg1_01_行程单.pdf",
                           "attachment_original_name": "行程单.pdf",
                           "identified_type": "Itinerary", "text_snapshot": "出行时间：2026-09-25 10:00"}]}
        with patch("ai_client.Generation.call", return_value=response), \
             patch("ai_client._rename_pdfs_with_audit", return_value=([], [])):
            data, error, _, _ = call_ai_audit_and_rename(task)
        self.assertEqual(data["match_amount"], NOT_FOUND)
        self.assertIn("AI_MISSING_FIELDS:match_amount", error)
        self.assertNotIn("invoice_number", error)

    def test_text_pdf_pair_saves_itinerary_amount_only_in_internal_json(self):
        invoice_pdf = ITINERARY_PDF.with_name("发票.pdf")
        mail = EmailMessage()
        mail["Subject"] = "电子发票"
        mail["Date"] = "Thu, 15 Oct 2026 08:00:00 +0800"
        mail.set_content("受控测试")
        for path in (invoice_pdf, ITINERARY_PDF):
            mail.add_attachment(path.read_bytes(), maintype="application", subtype="pdf",
                                filename=path.name)
        offline = SimpleNamespace(status_code=400, code="OFFLINE_TEST", message="mocked")
        with tempfile.TemporaryDirectory() as root:
            cfg = RunConfig("imap.test", "user", "unused", output_root=root,
                            dashscope_api_key="test")
            with patch("pipeline.imaplib.IMAP4_SSL",
                       return_value=FakeImap(b"701", mail.as_bytes())), \
                 patch("ai_client.Generation.call", return_value=offline), \
                 patch("ai_client._rename_pdfs_with_audit", return_value=([], [])):
                _, stats = run_pipeline(cfg)
            saved = json.loads(Path(cfg.debug_file).read_text(encoding="utf-8"))
            book = load_workbook(cfg.summary_excel_file, read_only=True)
            try:
                self.assertEqual(book.active.max_row - 1, 1)
                self.assertEqual(book.active.max_column, 8)
            finally:
                book.close()
        self.assertIsNone(stats["imap_error"])
        self.assertEqual(saved[0]["itineraries"][0]["match_amount"], "38.93")
        self.assertEqual(saved[0]["records"][0]["matched_itinerary"]["match_amount"], "38.93")
        self.assertEqual(saved[0]["files"][1]["pre_read_amount"], "38.93")

    def test_one_failed_mail_does_not_discard_other_json_and_excel_records(self):
        calls = 0

        def extract(file, *_args):
            nonlocal calls
            calls += 1
            if calls == 2:
                raise RuntimeError("synthetic PDF failure")
            return DocumentResult(file.meta.file_seq, "Invoice", {
                "invoice_date": "2026-10-15", "invoice_number": str(calls) * 20,
                "total_amount": "20.00", "invoice_type": "", "receipt_type": "办公用品",
                "seller": "测试公司"}, [file.meta.file_seq])

        with tempfile.TemporaryDirectory() as root:
            cfg = RunConfig("imap.test", "user", "unused", output_root=root,
                            dashscope_api_key="test")
            with patch("pipeline.imaplib.IMAP4_SSL",
                       return_value=FakeImap(b"701 702 703", message().as_bytes())), \
                 patch("pipeline.process_pdf", side_effect=extract):
                _, stats = run_pipeline(cfg)
            saved = json.loads(Path(cfg.debug_file).read_text(encoding="utf-8"))
            book = load_workbook(cfg.summary_excel_file, read_only=True)
            try:
                self.assertEqual(book.active.max_row - 1, 2)
                self.assertEqual(book.active.max_column, 8)
            finally:
                book.close()
        self.assertIsNone(stats["imap_error"])
        self.assertEqual(len(stats["failed_mails"]), 1)
        self.assertEqual([len(mail["records"]) for mail in saved], [1, 0, 1])
        self.assertEqual(saved[1]["mail_uid"], "702")
        self.assertIn("RuntimeError", saved[1]["error_reason"])

    def test_failed_second_pdf_preserves_first_pdf_in_same_mail(self):
        mail = message()
        mail.add_attachment(b"%PDF-1.4 second", maintype="application", subtype="pdf",
                            filename="第二张发票.pdf")

        def extract(file, *_args):
            if file.meta.file_seq == "02":
                raise RuntimeError("synthetic second file failure")
            return DocumentResult(file.meta.file_seq, "Invoice", {
                "invoice_date": "2026-10-15", "invoice_number": "1" * 20,
                "total_amount": "20.00", "receipt_type": "办公用品",
                "seller": "测试公司"}, [file.meta.file_seq])

        with tempfile.TemporaryDirectory() as root:
            cfg = RunConfig("imap.test", "user", "unused", output_root=root,
                            dashscope_api_key="test")
            with patch("pipeline.imaplib.IMAP4_SSL",
                       return_value=FakeImap(b"701", mail.as_bytes())), \
                 patch("pipeline.process_pdf", side_effect=extract):
                _, stats = run_pipeline(cfg)
            saved = json.loads(Path(cfg.debug_file).read_text(encoding="utf-8"))
            book = load_workbook(cfg.summary_excel_file, read_only=True)
            try:
                self.assertEqual(book.active.max_row - 1, 1)
            finally:
                book.close()
        self.assertEqual(len(saved[0]["files"]), 2)
        self.assertEqual(saved[0]["records"][0]["record_id"], "Msg1-01")
        self.assertEqual(stats["failed_mails"][0]["failed_files"][0]["file_seq"], "02")

    def test_late_unexpected_error_keeps_partial_json_as_manual_review(self):
        def extract(file, *_args):
            return DocumentResult(file.meta.file_seq, "Invoice", {
                "invoice_date": "2026-10-15", "invoice_number": "1" * 20,
                "total_amount": "20.00", "receipt_type": "办公用品",
                "seller": "测试公司"}, [file.meta.file_seq])

        def fail_arbitration(_mails):
            raise RuntimeError("synthetic arbitration failure")

        with tempfile.TemporaryDirectory() as root:
            cfg = RunConfig("imap.test", "user", "unused", output_root=root,
                            dashscope_api_key="test")
            with patch("pipeline.imaplib.IMAP4_SSL",
                       return_value=FakeImap(b"701", message().as_bytes())), \
                 patch("pipeline.process_pdf", side_effect=extract), \
                 patch("pipeline.load_stage_ports",
                       return_value=StagePorts(arbitrate_records=fail_arbitration)):
                _, stats = run_pipeline(cfg)
            saved = json.loads(Path(cfg.debug_file).read_text(encoding="utf-8"))
            self.assertFalse(Path(cfg.summary_excel_file).exists())
        self.assertIn("RuntimeError", stats["run_error"])
        self.assertEqual(saved[0]["records"][0]["status"], "人工待核")


if __name__ == "__main__":
    unittest.main()
