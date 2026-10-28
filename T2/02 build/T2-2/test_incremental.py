"""2026-10-04 T2-2 增量研发自测；不代表测试工程师独立验收。"""
from __future__ import annotations

import json
import os
import tempfile
import unittest
from decimal import Decimal
from email.message import EmailMessage
from unittest.mock import patch
from types import SimpleNamespace

from openpyxl import load_workbook

from config import RunConfig, ensure_output_dirs
from ai_client import call_ai_audit_and_rename
from attachment_processor import CollectedFile
from data_model import (
    FileRecord, InvoiceFields, InvoiceRecord, ItineraryRecord, MailRecord,
    MatchedItinerary, NOT_FOUND, amount_requires_review, amount_to_json,
    mail_to_dict, normalize_amount,
)
from excel_writer import append_to_summary_excel
from pipeline import _check_source_chain, process_message
from pipeline import run_pipeline
from pdf_processor import process_pdf
from processor_contracts import DocumentResult
from stage_ports import StagePorts


def message(*names: str) -> EmailMessage:
    msg = EmailMessage()
    msg["Subject"] = "电子发票"
    msg["Date"] = "Sun, 04 Oct 2026 09:30:00 +0800"
    msg.set_content("研发自测")
    for index, name in enumerate(names):
        msg.add_attachment(f"file-{index}".encode(), maintype="application",
                           subtype="octet-stream", filename=name)
    return msg


class T22IncrementalTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.cfg = RunConfig("imap.test", "user", "password",
                             output_root=self.temp.name, dashscope_api_key="test")
        ensure_output_dirs(self.cfg)

    def test_all_itineraries_persist_with_only_four_fields(self):
        def pdf(file, *_args):
            if "行程单" in file.meta.original_name:
                number = file.meta.file_seq
                return DocumentResult(file.meta.file_seq, "Itinerary", {
                    "total_amount": "53.9300" if number == "02" else "48.31",
                    "travel_time": "2026-10-01 10:00",
                    "origin_dest": "杭州→上海",
                    "seller": "不应保存的行程单企业",
                })
            return DocumentResult(file.meta.file_seq, "Invoice", {
                "invoice_date": "2026-10-01", "invoice_number": "1" * 20,
                "total_amount": "53.9300", "receipt_type": "客运服务",
                "seller": "发票销售方",
            })

        with patch("pipeline.process_pdf", side_effect=pdf):
            outcome = process_message(
                message("甲发票.pdf", "甲行程单.pdf", "乙行程单.pdf"),
                "Msg1", "101", self.cfg, ports=StagePorts(),
            )
        self.assertEqual(len(outcome.itineraries), 2)
        self.assertEqual(len(outcome.mail.itineraries), 2)
        self.assertEqual(outcome.mail.records[0].fields.seller, "发票销售方")
        self.assertEqual(outcome.mail.records[0].status, "complete")
        outcome.mail.records[0].matched_itinerary = MatchedItinerary("02", "53.9300")
        _check_source_chain([outcome.mail])
        saved = json.loads(json.dumps(mail_to_dict(outcome.mail), ensure_ascii=False))
        self.assertEqual([i["record_id"] for i in saved["itineraries"]],
                         ["Msg1-02", "Msg1-03"])
        self.assertEqual(set(saved["itineraries"][0]),
                         {"record_id", "match_amount", "travel_time", "origin_dest"})
        self.assertEqual(saved["itineraries"][0]["match_amount"], "53.9300")
        self.assertEqual(saved["records"][0]["matched_itinerary"]["match_amount"], "53.9300")
        self.assertEqual(saved["records"][0]["fields"]["invoice_type"], "")
        self.assertTrue(all(f["internal_name"].startswith("Msg1_") for f in saved["files"]))
        self.assertTrue(all(f["exported_name"] == "" for f in saved["files"]))
        outcome.mail.records[0].matched_itinerary.match_amount = "53.94"
        with self.assertRaisesRegex(ValueError, "匹配金额与行程单条目不一致"):
            _check_source_chain([outcome.mail])

    def test_missing_receipt_type_does_not_default_to_b(self):
        def pdf(file, *_args):
            return DocumentResult(file.meta.file_seq, "Invoice", {
                "invoice_date": "2026-10-01", "invoice_number": "2" * 20,
                "total_amount": "10.00", "invoice_type": NOT_FOUND,
                "receipt_type": NOT_FOUND, "seller": "发票销售方",
            })

        with patch("pipeline.process_pdf", side_effect=pdf):
            outcome = process_message(message("发票.pdf"), "Msg2", "102", self.cfg,
                                      ports=StagePorts())
        row = outcome.mail.records[0]
        self.assertEqual(row.kind, "")
        self.assertEqual(row.status, "人工待核")
        self.assertEqual(row.fields.invoice_type, "")
        self.assertTrue(any("票面类型缺失" in str(a["flags"]) for a in outcome.alerts))

    def test_amount_precision_and_final_excel_display(self):
        self.assertEqual(normalize_amount("¥1,234.5600"), Decimal("1234.5600"))
        self.assertEqual(amount_to_json(Decimal("1234.5600")), "1234.5600")
        self.assertFalse(amount_requires_review("0.1234"))
        self.assertTrue(amount_requires_review("0.12345"))
        self.assertEqual(amount_to_json("0.12345"), "0.12345")
        self.assertTrue(amount_requires_review("not-a-number"))
        self.assertFalse(amount_requires_review(NOT_FOUND))
        path = os.path.join(self.temp.name, "amount.xlsx")
        error, written = append_to_summary_excel([{
            "invoice_date": "2026-10-01", "invoice_number": "3" * 20,
            "total_amount": Decimal("53.9300"), "invoice_type": "",
            "receipt_type": "办公用品", "seller": "发票销售方",
            "travel_time": "", "origin_dest": "",
        }], path)
        self.assertIsNone(error)
        self.assertEqual(written, path)
        book = load_workbook(path, read_only=True)
        try:
            self.assertEqual(book.active.cell(2, 3).value, 53.93)
            self.assertEqual(book.active.cell(2, 3).number_format, "0.00")
        finally:
            book.close()

        rounded_path = os.path.join(self.temp.name, "rounded.xlsx")
        append_to_summary_excel([{
            "invoice_date": "2026-10-01", "invoice_number": "7" * 20,
            "total_amount": "53.935", "invoice_type": "", "receipt_type": "办公用品",
            "seller": "发票销售方", "travel_time": "", "origin_dest": "",
        }], rounded_path)
        rounded = load_workbook(rounded_path, read_only=True)
        try:
            self.assertEqual(rounded.active.cell(2, 3).value, 53.94)
        finally:
            rounded.close()

    def test_excess_precision_keeps_value_and_flags_record(self):
        def pdf(file, *_args):
            return DocumentResult(file.meta.file_seq, "Invoice", {
                "invoice_date": "2026-10-01", "invoice_number": "4" * 20,
                "total_amount": "53.93001", "receipt_type": "办公用品",
                "seller": "发票销售方",
            })

        with patch("pipeline.process_pdf", side_effect=pdf):
            outcome = process_message(message("高精度发票.pdf"), "Msg3", "103", self.cfg,
                                      ports=StagePorts())
        self.assertEqual(outcome.mail.records[0].fields.total_amount, Decimal("53.93001"))
        self.assertEqual(outcome.mail.records[0].status, "人工待核")
        self.assertEqual(mail_to_dict(outcome.mail)["records"][0]["fields"]["total_amount"],
                         "53.93001")
        self.assertTrue(any("金额格式或精度" in str(a["flags"]) for a in outcome.alerts))

    def test_ai_json_number_keeps_four_decimal_places(self):
        response = SimpleNamespace(status_code=200, output=SimpleNamespace(text=json.dumps({
            "total_amount": 53.9300, "invoice_date": "2026-10-01",
            "invoice_number": "5" * 20, "receipt_type": "办公用品",
            "seller": "发票销售方", "file_classifications": [],
        })))
        # Python 的 json.dumps(float) 已丢掉尾零；直接用 AI 原始文本检验解析边界。
        response.output.text = response.output.text.replace("53.93", "53.9300")
        task = {"files": [{"local_path": "Msg1_01_发票.pdf",
                           "attachment_original_name": "发票.pdf",
                           "identified_type": "Invoice", "text_snapshot": ""}]}
        with patch("ai_client.Generation.call", return_value=response), \
             patch("ai_client._rename_pdfs_with_audit", return_value=([], [])):
            data, error, _, _ = call_ai_audit_and_rename(task)
        self.assertEqual(data["total_amount"], Decimal("53.9300"))
        self.assertEqual(data["invoice_type"], "")
        self.assertNotIn("invoice_type", error)

    def test_pdf_pre_read_keeps_amount_precision(self):
        path = os.path.join(self.temp.name, "Msg1_01_发票.pdf")
        with open(path, "wb") as stream:
            stream.write(b"test-pdf-copy")
        item = CollectedFile(FileRecord(file_seq="01", original_name="发票.pdf"), path)
        fallback = {"invoice_number": "6" * 20, "total_amount": "53.9300"}
        with patch("pdf_processor.local_inspect_pdf", return_value=("Invoice", "发票")), \
             patch("pdf_processor.build_fallback_audit_result", return_value=fallback), \
             patch("pdf_processor.call_ai_audit_and_rename", return_value=(fallback, "", [], [])):
            process_pdf(item, self.cfg, "电子发票", "2026-10-01")
        self.assertEqual(item.meta.pre_read_amount, Decimal("53.9300"))
        self.assertEqual(mail_to_dict(MailRecord(files=[item.meta]))["files"][0]["pre_read_amount"],
                         "53.9300")

    def test_run_pipeline_writes_unmatched_itinerary_to_debug_json(self):
        msg = message("行程单.pdf")

        class FakeImap:
            def login(self, *_args): return "OK", []
            def select(self, *_args): return "OK", [b"1"]
            def uid(self, command, *_args):
                if command == "SEARCH":
                    return "OK", [b"101"]
                return "OK", [(b"1", msg.as_bytes())]
            def status(self, *_args):
                return "OK", [b"INBOX (MESSAGES 1 UNSEEN 1)"]
            def logout(self): return "BYE", []

        def pdf(file, *_args):
            return DocumentResult(file.meta.file_seq, "Itinerary", {
                "total_amount": "53.9300", "travel_time": "2026-10-01 10:00",
                "origin_dest": "杭州→上海",
            })

        with patch("pipeline.imaplib.IMAP4_SSL", return_value=FakeImap()), \
             patch("pipeline.process_pdf", side_effect=pdf), \
             patch("pipeline.load_stage_ports", return_value=StagePorts()):
            _, stats = run_pipeline(self.cfg)
        self.assertIsNone(stats["imap_error"])
        with open(self.cfg.debug_file, encoding="utf-8") as stream:
            saved = json.load(stream)
        self.assertEqual(saved[0]["records"], [])
        self.assertEqual(saved[0]["itineraries"], [{
            "record_id": "Msg1-01", "match_amount": "53.9300",
            "travel_time": "2026-10-01 10:00", "origin_dest": "杭州→上海",
        }])


if __name__ == "__main__":
    unittest.main()
