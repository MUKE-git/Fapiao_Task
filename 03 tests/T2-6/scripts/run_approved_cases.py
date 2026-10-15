"""Formal T2-6 controlled main-flow cases approved by the user.

Inputs and expected outcomes are in fixtures/approved_cases.json. No expected
business value is copied from the production output. External IMAP and first
PDF extraction are controlled; all T2 stage ports and JSON/Excel I/O are real.
"""
from __future__ import annotations

import hashlib
import json
import sys
from datetime import datetime
from decimal import Decimal, ROUND_HALF_UP
from email.message import EmailMessage
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[3]
sys.path.insert(0, str(ROOT))
from config import RunConfig, SUMMARY_COLUMNS  # noqa: E402
from pipeline import run_pipeline  # noqa: E402
from processor_contracts import DocumentResult  # noqa: E402

BASE = ROOT / "03 tests" / "T2-6"
SPEC_PATH = BASE / "fixtures" / "approved_cases.json"
SPEC = json.loads(SPEC_PATH.read_text(encoding="utf-8"))
STAMP = datetime.now().strftime("%Y%m%d_%H%M%S")
RUNS = BASE / "result" / f"approved_{STAMP}"


def sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def expected_input(invoice: dict) -> dict:
    return {**SPEC["common"], **invoice}


def expected_amount(value) -> str:
    if value in (None, ""):
        return "Not Found"
    return str(value)


def expected_receipt(value) -> str:
    if value is None or value == "":
        return "Not Found"
    if isinstance(value, str) and value.startswith("*"):
        chunks = value.split("*")
        return chunks[2] if len(chunks) == 3 else ""
    return str(value)


def expected_fields(invoice: dict, itinerary: dict | None) -> dict:
    values = expected_input(invoice)
    fields = {
        "invoice_date": values["invoice_date"] or "Not Found",
        "invoice_number": values["invoice_number"] or "Not Found",
        "total_amount": expected_amount(values["total_amount"]),
        "invoice_type": values["invoice_type"],
        "receipt_type": expected_receipt(values["receipt_type"]),
        "seller": str(values["seller"]),
    }
    if itinerary:
        fields["travel_time"] = itinerary["travel_time"]
        fields["origin_dest"] = itinerary["origin_dest"]
    return fields


def expected_excel_row(fields: dict, kind: str) -> list:
    amount = fields["total_amount"]
    if amount == "1,2,3":
        excel_amount = "1,2,3"  # Invalid value must not silently become 123.
    elif amount == "Not Found":
        excel_amount = "Not Found"
    else:
        excel_amount = float(Decimal(amount).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP))
    return [fields["invoice_date"], fields["invoice_number"], excel_amount,
            fields["invoice_type"] or None, fields["receipt_type"],
            fields.get("travel_time") if kind == "a类" else None,
            fields["seller"], fields.get("origin_dest") if kind == "a类" else None]


def compare(path: str, actual, expected, diffs: list) -> None:
    if actual != expected:
        diffs.append({"path": path, "expected": expected, "actual": actual})


class FakeImap:
    def __init__(self, raw: dict[str, bytes]): self.raw = raw
    def login(self, *_): return "OK", []
    def select(self, *_): return "OK", []
    def logout(self): return "BYE", []
    def status(self, *_):
        n = len(self.raw)
        return "OK", [f"INBOX (MESSAGES {n} UNSEEN {n})".encode()]
    def uid(self, command, *args):
        if command == "SEARCH": return "OK", [" ".join(self.raw).encode()]
        uid = args[0].decode() if isinstance(args[0], bytes) else str(args[0])
        return "OK", [(b"body", self.raw[uid])]


def run(case: dict) -> dict:
    case_dir = RUNS / case["id"]
    case_dir.mkdir(parents=True, exist_ok=True)
    cfg = RunConfig("imap.test", "user", "unused", output_folder_name="output",
                    output_root=str(case_dir), dashscope_api_key="controlled-test")
    raw_messages: dict[str, bytes] = {}
    fields_by_name: dict[str, tuple[str, dict]] = {}
    input_hashes: dict[str, str] = {}
    for i, item in enumerate(case["mails"], 1):
        msg = EmailMessage()
        msg["Subject"] = "电子发票"
        msg["Date"] = "Thu, 15 Oct 2026 09:00:00 +0800"
        msg.set_content("approved synthetic input")
        invoice_name = f"发票{i}.pdf"
        invoice_bytes = f"controlled PDF invoice {case['id']} {i}".encode()
        msg.add_attachment(invoice_bytes, maintype="application", subtype="pdf", filename=invoice_name)
        fields_by_name[invoice_name] = ("Invoice", expected_input(item["invoice"]))
        if "itinerary" in item:
            trip_name = f"行程单{i}.pdf"
            trip_bytes = f"controlled PDF itinerary {case['id']} {i}".encode()
            msg.add_attachment(trip_bytes, maintype="application", subtype="pdf", filename=trip_name)
            fields_by_name[trip_name] = ("Itinerary", item["itinerary"])
            trip_path = case_dir / f"input_{i:02d}_trip.pdf"
            trip_path.write_bytes(trip_bytes)
            input_hashes[str(trip_path.relative_to(ROOT))] = sha(trip_path)
        invoice_path = case_dir / f"input_{i:02d}_invoice.pdf"
        invoice_path.write_bytes(invoice_bytes)
        input_hashes[str(invoice_path.relative_to(ROOT))] = sha(invoice_path)
        raw = msg.as_bytes()
        eml_path = case_dir / f"input_{i:02d}.eml"
        eml_path.write_bytes(raw)
        input_hashes[str(eml_path.relative_to(ROOT))] = sha(eml_path)
        raw_messages[item["uid"]] = raw

    def extract(file, *_):
        role, values = fields_by_name[file.meta.original_name]
        return DocumentResult(file.meta.file_seq, role, values, [file.meta.file_seq])

    with patch("pipeline.imaplib.IMAP4_SSL", return_value=FakeImap(raw_messages)), \
         patch("pipeline.process_pdf", side_effect=extract), \
         patch("dedup_arbitrator.Generation.call", side_effect=AssertionError("unexpected AI call")):
        mismatches, stats = run_pipeline(cfg)

    diffs: list[dict] = []
    compare("mismatches", mismatches, [], diffs)
    compare("stats.imap_error", stats["imap_error"], None, diffs)
    compare("stats.unseen_total", stats["unseen_total"], len(case["mails"]), diffs)
    debug_path = Path(cfg.debug_file)
    actual = json.loads(debug_path.read_text(encoding="utf-8")) if debug_path.is_file() else []
    compare("JSON.mail_count", len(actual), len(case["mails"]), diffs)
    expected_records = []
    for i, (item, mail) in enumerate(zip(case["mails"], actual), 1):
        prefix = f"JSON[{i - 1}]"
        compare(prefix + ".keys", sorted(mail), sorted(("mail_id", "mail_uid", "subject",
                "mail_date", "files", "itineraries", "records", "dedup_skipped",
                "error_reason")), diffs)
        compare(prefix + ".mail_id", mail.get("mail_id"), f"Msg{i}", diffs)
        compare(prefix + ".mail_uid", mail.get("mail_uid"), item["uid"], diffs)
        compare(prefix + ".subject", mail.get("subject"), "电子发票", diffs)
        compare(prefix + ".mail_date", mail.get("mail_date"), "2026-10-15 09:00", diffs)
        expected_error = []
        if case.get("duplicate") and i == 2:
            expected_error.append(f"Msg2-01 与 Msg1-01 的发票号及五个必填字段一致，已跳过重复 Excel 行；原记录保留在 JSON")
        for issue in case["issues"]:
            expected_error.append(f"T2-6 Msg{i}-01 {issue}")
        compare(prefix + ".error_reason", mail.get("error_reason"), ";".join(expected_error), diffs)
        compare(prefix + ".dedup_skipped", mail.get("dedup_skipped"), [], diffs)
        compare(prefix + ".record_count", len(mail.get("records", [])), 1, diffs)
        itinerary = item.get("itinerary")
        compare(prefix + ".file_count", len(mail.get("files", [])), 2 if itinerary else 1, diffs)
        for j, f in enumerate(mail.get("files", []), 1):
            expected_name = (f"发票{i}.pdf" if j == 1 else f"行程单{i}.pdf")
            expected_file = {"file_seq": f"{j:02d}", "file_type": "PDF",
                             "original_name": expected_name, "pre_read_type": "Unknown",
                             "pre_read_invoice_no": "Not Found", "pre_read_amount": "Not Found",
                             "internal_name": f"Msg{i}_{j:02d}_{expected_name}", "exported_name": ""}
            compare(f"{prefix}.files[{j - 1}]", f, expected_file, diffs)
            actual_path = Path(cfg.extract_dir) / expected_file["internal_name"]
            original_path = case_dir / (f"input_{i:02d}_invoice.pdf" if j == 1 else f"input_{i:02d}_trip.pdf")
            compare(f"{prefix}.files[{j - 1}].sha256",
                    sha(actual_path) if actual_path.is_file() else "missing", sha(original_path), diffs)
        expected_itineraries = ([{"record_id": f"Msg{i}-02",
                                  "match_amount": itinerary["match_amount"],
                                  "travel_time": itinerary["travel_time"],
                                  "origin_dest": itinerary["origin_dest"]}]
                                if itinerary else [])
        compare(prefix + ".itineraries", mail.get("itineraries"), expected_itineraries, diffs)
        fields = expected_fields(item["invoice"], itinerary)
        record = {"record_id": f"Msg{i}-01", "kind": case["kinds"][i - 1],
                  "fields": fields, "source": {"source_file_seq": ["01"]},
                  "status": case["statuses"][i - 1]}
        if itinerary:
            record["matched_itinerary"] = {"file_seq": "02", "match_amount": itinerary["match_amount"]}
        expected_records.append(record)
        if mail.get("records"):
            compare(prefix + ".records[0]", mail["records"][0], record, diffs)

    excel_path = Path(cfg.summary_excel_file)
    if not case["excel_from"]:
        compare("Excel.absent", excel_path.exists(), False, diffs)
    elif not excel_path.is_file():
        diffs.append({"path": "Excel.file", "expected": "exists", "actual": "missing"})
    else:
        workbook = load_workbook(excel_path)
        sheet = workbook.active
        compare("Excel.headers", [sheet.cell(1, c).value for c in range(1, 9)],
                [label for label, _ in SUMMARY_COLUMNS], diffs)
        compare("Excel.row_count", sheet.max_row - 1, len(case["excel_from"]), diffs)
        for rindex, source_index in enumerate(case["excel_from"], 2):
            record = expected_records[source_index]
            expected_row = expected_excel_row(record["fields"], record["kind"])
            for column, expected in enumerate(expected_row, 1):
                cell = sheet.cell(rindex, column)
                actual_value = cell.value
                if expected in (None, "") and actual_value is None:
                    continue
                if column == 3 and isinstance(expected, float) and isinstance(actual_value, (int, float)):
                    compare(f"Excel.row[{rindex - 1}].amount", float(actual_value), expected, diffs)
                    compare(f"Excel.row[{rindex - 1}].amount_format", cell.number_format, "0.00", diffs)
                    compare(f"Excel.row[{rindex - 1}].amount_type", cell.data_type, "n", diffs)
                else:
                    compare(f"Excel.row[{rindex - 1}].col[{column}]", actual_value, expected, diffs)
        workbook.close()

    all_flags = [flag for alert in stats["not_found_alerts"] for flag in alert.get("flags", [])]
    for i in range(1, len(case["mails"]) + 1):
        for issue in case["issues"]:
            flag = f"T2-6 Msg{i}-01 {issue}"
            if flag not in all_flags:
                diffs.append({"path": "stats.not_found_alerts", "expected_contains": flag,
                              "actual": all_flags})
    if not case["issues"] and any(flag.startswith("T2-6 ") for flag in all_flags):
        diffs.append({"path": "stats.not_found_alerts", "expected": "no T2-6 issues",
                      "actual": all_flags})
    return {"case": case["id"], "passed": not diffs, "differences": diffs,
            "input_sha256": input_hashes,
            "output_sha256": {str(path.relative_to(ROOT)): sha(path)
                              for path in (debug_path, excel_path) if path.is_file()},
            "actual_statuses": [mail["records"][0]["status"] for mail in actual if mail.get("records")]}


def main() -> int:
    RUNS.mkdir(parents=True, exist_ok=True)
    results = [run(case) for case in SPEC["cases"]]
    source_names = ("pipeline.py", "field_validator.py", "data_model.py", "stage_ports.py",
                    "dedup_arbitrator.py", "excel_writer.py", "attachment_processor.py",
                    "03 tests/T2-6/scripts/run_approved_cases.py",
                    "03 tests/T2-6/fixtures/approved_cases.json")
    report = {"approval": SPEC["approval"], "scope": "controlled main flow, formal approved cases",
              "source_sha256": {name: sha(ROOT / name) for name in source_names},
              "passed": sum(result["passed"] for result in results), "total": len(results),
              "results": results}
    path = BASE / "result" / f"approved_report_{STAMP}.json"
    path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"passed": report["passed"], "total": report["total"],
                      "failed": {r["case"]: r["differences"] for r in results if not r["passed"]},
                      "report": str(path)}, ensure_ascii=False, indent=2))
    return 0 if report["passed"] == report["total"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
