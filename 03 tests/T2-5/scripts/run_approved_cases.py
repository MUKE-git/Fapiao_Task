"""Formal T2-5 independent cases approved by the user.

Case A and D run the main pipeline with controlled IMAP/PDF extraction; B/C
exercise actual MIME attachment storage and production JSON serialization.
Expected values originate in fixtures/approved_cases.json and approved rules.
"""
from __future__ import annotations

import hashlib
import importlib.util
import io
import json
import re
import sys
import zipfile
from datetime import datetime
from email.message import EmailMessage
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[3]
sys.path.insert(0, str(ROOT))
from attachment_processor import collect_mail_files  # noqa: E402
from config import RunConfig, SUMMARY_COLUMNS, ensure_output_dirs  # noqa: E402
from data_model import MailRecord, mail_to_dict  # noqa: E402
from pipeline import _check_source_chain, run_pipeline  # noqa: E402
from processor_contracts import DocumentResult  # noqa: E402

BASE = ROOT / "03 tests" / "T2-5"
SPEC_PATH = BASE / "fixtures" / "approved_cases.json"
SPEC = json.loads(SPEC_PATH.read_text(encoding="utf-8"))
STAMP = datetime.now().strftime("%Y%m%d_%H%M%S")
RUNS = BASE / "result" / f"approved_{STAMP}"


def sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def compare(path: str, actual, expected, diffs: list) -> None:
    if actual != expected:
        diffs.append({"path": path, "expected": expected, "actual": actual})


def common_file_expected(seq: str, name: str, internal: str, role="Unknown") -> dict:
    return {"file_seq": seq, "file_type": Path(internal).suffix.lower().lstrip(".").upper()
            if Path(internal).suffix.lower() in (".pdf", ".ofd", ".xml") else "其他",
            "original_name": name, "pre_read_type": role,
            "pre_read_invoice_no": "Not Found", "pre_read_amount": "Not Found",
            "internal_name": internal, "exported_name": ""}


def case_a() -> dict:
    path = BASE / "scripts" / "explore_trace.py"
    spec = importlib.util.spec_from_file_location("t25_trace_explorer", path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    module.BASE = RUNS / "甲"
    module.main()
    cfg = RunConfig("imap.test", "user", "unused", output_folder_name="output",
                    output_root=str(module.BASE), dashscope_api_key="controlled-key")
    mails = json.loads(Path(cfg.debug_file).read_text(encoding="utf-8"))
    expected = SPEC["case_A"]
    diffs = []
    compare("JSON.mail_count", len(mails), 2, diffs)
    expected_bytes = {"invoice_one": b"synthetic-invoice-701",
                      "itinerary": b"synthetic-itinerary-701",
                      "invoice_two": b"synthetic-invoice-702"}
    for i, mail in enumerate(mails):
        prefix = f"JSON[{i}]"
        compare(prefix + ".keys", sorted(mail), sorted(("mail_id", "mail_uid", "subject",
                "mail_date", "files", "itineraries", "records", "dedup_skipped",
                "error_reason")), diffs)
        compare(prefix + ".mail_id", mail["mail_id"], f"Msg{i + 1}", diffs)
        compare(prefix + ".mail_uid", mail["mail_uid"], expected["uids"][i], diffs)
        compare(prefix + ".subject", mail["subject"], "电子发票", diffs)
        compare(prefix + ".mail_date", mail["mail_date"], "2026-10-15 09:00", diffs)
        compare(prefix + ".error_reason", mail["error_reason"], "", diffs)
        compare(prefix + ".dedup_skipped", mail["dedup_skipped"],
                expected["dedup_skipped"] if i == 0 else [], diffs)
        compare(prefix + ".file_count", len(mail["files"]), len(expected["files"][i]), diffs)
        for j, (f, file_spec) in enumerate(zip(mail["files"], expected["files"][i])):
            role = "Itinerary" if i == 0 and j == 1 else "Invoice"
            internal = f"Msg{i + 1}_{file_spec['file_seq']}_{Path(file_spec['original_name']).name}"
            expected_file = common_file_expected(file_spec["file_seq"], file_spec["original_name"], internal, role)
            compare(f"{prefix}.files[{j}]", f, expected_file, diffs)
            physical = Path(cfg.extract_dir) / internal
            compare(f"{prefix}.files[{j}].sha256", sha(physical) if physical.is_file() else "missing",
                    hashlib.sha256(expected_bytes[file_spec["bytes_label"]]).hexdigest(), diffs)
        expected_itineraries = [expected["itinerary"]] if i == 0 else []
        compare(prefix + ".itineraries", mail["itineraries"], expected_itineraries, diffs)
        record = {"record_id": f"Msg{i + 1}-01", "kind": "a类" if i == 0 else "b类",
                  "fields": expected["invoice_fields"][i],
                  "source": {"source_file_seq": ["01"]}, "status": "complete"}
        if i == 0:
            record["matched_itinerary"] = {"file_seq": "02", "match_amount": "900.00"}
        compare(prefix + ".records", mail["records"], [record], diffs)
    workbook = load_workbook(cfg.summary_excel_file)
    sheet = workbook.active
    compare("Excel.headers", [sheet.cell(1, c).value for c in range(1, 9)],
            [title for title, _ in SUMMARY_COLUMNS], diffs)
    compare("Excel.row_count", sheet.max_row - 1, 2, diffs)
    for i, row in enumerate(expected["excel_rows"], 2):
        for col, value in enumerate(row, 1):
            cell = sheet.cell(i, col)
            if value is None and cell.value is None: continue
            if col == 3 and isinstance(cell.value, (int, float)):
                compare(f"Excel.row[{i - 1}].amount", float(cell.value), value, diffs)
                compare(f"Excel.row[{i - 1}].amount_format", cell.number_format, "0.00", diffs)
                compare(f"Excel.row[{i - 1}].amount_type", cell.data_type, "n", diffs)
            else:
                compare(f"Excel.row[{i - 1}].col[{col}]", cell.value, value, diffs)
    workbook.close()
    return {"case": "甲", "passed": not diffs, "differences": diffs,
            "input_sha256": {uid: sha(module.BASE / f"input_{uid}.eml") for uid in expected["uids"]},
            "output_sha256": {"json": sha(Path(cfg.debug_file)), "excel": sha(Path(cfg.summary_excel_file))}}


def make_mail(attachments: list[tuple[str, bytes, str]]) -> EmailMessage:
    msg = EmailMessage()
    msg["Subject"] = "电子发票"
    msg["Date"] = "Thu, 15 Oct 2026 09:00:00 +0800"
    msg.set_content("approved synthetic input")
    for name, data, subtype in attachments:
        msg.add_attachment(data, maintype="application", subtype=subtype, filename=name)
    return msg


def case_b() -> dict:
    spec = SPEC["case_B"]
    root = RUNS / "乙"
    cfg = RunConfig("imap.test", "user", "unused", output_folder_name="output",
                    output_root=str(root), dashscope_api_key="controlled-key")
    ensure_output_dirs(cfg)
    name1 = spec["exact_stem_char"] * spec["exact_stem_length"] + spec["extension"]
    name2 = spec["over_stem_char"] * spec["over_stem_length"] + spec["extension"]
    sentinel = Path(cfg.extract_dir) / ("Msg1_01_" + name1)
    sentinel.write_bytes(spec["sentinel_bytes"].encode())
    msg = make_mail([(name1, spec["attachment_bytes"][0].encode(), "pdf"),
                     (name2, spec["attachment_bytes"][1].encode(), "pdf")])
    eml_path = root / "input.eml"
    eml_path.write_bytes(msg.as_bytes())
    mail = MailRecord(mail_id="Msg1", mail_uid=spec["uid"], subject="电子发票",
                      mail_date="2026-10-15 09:00")
    collect_mail_files(msg, mail, cfg)
    _check_source_chain([mail], cfg)
    saved = mail_to_dict(mail)
    debug = root / "task_debug.json"
    debug.write_text(json.dumps([saved], ensure_ascii=False, indent=2), encoding="utf-8")
    diffs = []
    check = lambda label, actual, expected: compare(label, actual, expected, diffs)
    check("JSON.mail_id", saved["mail_id"], "Msg1")
    check("JSON.mail_uid", saved["mail_uid"], spec["uid"])
    check("JSON.subject", saved["subject"], "电子发票")
    check("JSON.mail_date", saved["mail_date"], "2026-10-15 09:00")
    check("JSON.records", saved["records"], [])
    check("JSON.itineraries", saved["itineraries"], [])
    check("JSON.dedup_skipped", saved["dedup_skipped"], [])
    check("JSON.error_reason", saved["error_reason"], "")
    check("JSON.file_count", len(saved["files"]), 2)
    hash1 = hashlib.sha256(name1.encode()).hexdigest()[:8]
    hash2 = hashlib.sha256(name2.encode()).hexdigest()[:8]
    expected_names = ["Msg1_01_" + spec["exact_stem_char"] * 96 + "_" + hash1 + "(1).pdf",
                      "Msg1_02_" + spec["over_stem_char"] * 99 + "_" + hash2 + ".pdf"]
    for i, (f, name, internal, content) in enumerate(zip(saved["files"], (name1, name2),
                                                          expected_names, spec["attachment_bytes"]), 1):
        check(f"JSON.files[{i}]", f, common_file_expected(f"{i:02d}", name, internal))
        path = Path(cfg.extract_dir) / internal
        check(f"file[{i}].sha256", sha(path) if path.is_file() else "missing",
              hashlib.sha256(content.encode()).hexdigest())
        check(f"file[{i}].length", len(internal), 120)
    check("sentinel.sha256", sha(sentinel), hashlib.sha256(spec["sentinel_bytes"].encode()).hexdigest())
    return {"case": "乙", "passed": not diffs, "differences": diffs,
            "input_sha256": sha(eml_path), "output_sha256": sha(debug)}


def case_c() -> dict:
    spec = SPEC["case_C"]
    root = RUNS / "丙"
    cfg = RunConfig("imap.test", "user", "unused", output_folder_name="output",
                    output_root=str(root), dashscope_api_key="controlled-key")
    ensure_output_dirs(cfg)
    names = ["", "CON.ofd", "a<bad>?.xml", "=?utf-8?b?5Y+R56WoLnBkZg==?="]
    msg = make_mail([(name, content.encode(), ext.lstrip(".")) for name, content, ext in
                     zip(names, spec["contents"], spec["extensions"])])
    eml_path = root / "input.eml"
    eml_path.write_bytes(msg.as_bytes())
    mail = MailRecord(mail_id="Msg1", mail_uid=spec["uid"], subject="电子发票",
                      mail_date="2026-10-15 09:00")
    collect_mail_files(msg, mail, cfg)
    _check_source_chain([mail], cfg)
    saved = mail_to_dict(mail)
    debug = root / "task_debug.json"
    debug.write_text(json.dumps([saved], ensure_ascii=False, indent=2), encoding="utf-8")
    diffs = []
    compare("JSON.mail_id", saved["mail_id"], "Msg1", diffs)
    compare("JSON.mail_uid", saved["mail_uid"], spec["uid"], diffs)
    compare("JSON.subject", saved["subject"], "电子发票", diffs)
    compare("JSON.mail_date", saved["mail_date"], "2026-10-15 09:00", diffs)
    compare("JSON.file_count", len(saved["files"]), 4, diffs)
    compare("JSON.original_names", [f["original_name"] for f in saved["files"]], spec["names"], diffs)
    compare("JSON.file_types", [f["file_type"] for f in saved["files"]], spec["types"], diffs)
    compare("JSON.records", saved["records"], [], diffs)
    compare("JSON.itineraries", saved["itineraries"], [], diffs)
    compare("JSON.dedup_skipped", saved["dedup_skipped"], [], diffs)
    compare("JSON.error_reason", saved["error_reason"], "", diffs)
    for i, (f, ext, content) in enumerate(zip(saved["files"], spec["extensions"], spec["contents"]), 1):
        internal = f["internal_name"]
        if len(internal) > 120 or not internal.startswith(f"Msg1_{i:02d}_") or not internal.endswith(ext):
            diffs.append({"path": f"JSON.files[{i}].internal_name", "expected": "safe prefix/extension/length", "actual": internal})
        if re.search(r'[<>:"/\\|?*]', internal):
            diffs.append({"path": f"JSON.files[{i}].unsafe_name", "expected": "safe", "actual": internal})
        expected_file = common_file_expected(f"{i:02d}", spec["names"][i - 1], internal)
        compare(f"JSON.files[{i}]", f, expected_file, diffs)
        physical = Path(cfg.extract_dir) / internal
        compare(f"file[{i}].sha256", sha(physical) if physical.is_file() else "missing",
                hashlib.sha256(content.encode()).hexdigest(), diffs)
    return {"case": "丙", "passed": not diffs, "differences": diffs,
            "input_sha256": sha(eml_path), "output_sha256": sha(debug)}


class FakeImap:
    def __init__(self, raw): self.raw = raw
    def login(self, *_): return "OK", []
    def select(self, *_): return "OK", []
    def logout(self): return "BYE", []
    def status(self, *_): return "OK", [b"INBOX (MESSAGES 1 UNSEEN 1)"]
    def uid(self, command, *_):
        return ("OK", [b"705"]) if command == "SEARCH" else ("OK", [(b"body", self.raw)])


def case_d(variant: dict) -> dict:
    root = RUNS / variant["id"]
    root.mkdir(parents=True, exist_ok=True)
    msg = make_mail([("发票.pdf", b"source-test-pdf", "pdf")])
    raw = msg.as_bytes()
    eml_path = root / "input.eml"
    eml_path.write_bytes(raw)
    cfg = RunConfig("imap.test", "user", "unused", output_folder_name="output",
                    output_root=str(root), dashscope_api_key="controlled-key")

    def extract(file, *_):
        if variant.get("delete_file"):
            Path(file.path).unlink()
        return DocumentResult("01", "Invoice", {
            "invoice_date": "2026-10-15", "invoice_number": "26327000000775232739",
            "total_amount": "100.00", "invoice_type": "", "receipt_type": "办公用品",
            "seller": "甲公司"}, [variant.get("bad_source", "01")])

    with patch("pipeline.imaplib.IMAP4_SSL", return_value=FakeImap(raw)), \
         patch("pipeline.process_pdf", side_effect=extract):
        mismatches, stats = run_pipeline(cfg)
    diffs = []
    compare("mismatches", mismatches, [], diffs)
    if variant["expected_error"] not in (stats.get("imap_error") or ""):
        diffs.append({"path": "stats.imap_error", "expected_contains": variant["expected_error"],
                      "actual": stats.get("imap_error")})
    compare("Excel.absent", Path(cfg.summary_excel_file).exists(), False, diffs)
    debug = Path(cfg.debug_file)
    if debug.is_file():
        saved = json.loads(debug.read_text(encoding="utf-8"))
        compare("JSON.mail_count", len(saved), 1, diffs)
        compare("JSON[0].mail_uid", saved[0]["mail_uid"], SPEC["case_D"]["uid"], diffs)
        compare("JSON[0].record_count", len(saved[0]["records"]), 1, diffs)
    return {"case": variant["id"], "passed": not diffs, "differences": diffs,
            "input_sha256": sha(eml_path), "output_sha256": sha(debug) if debug.is_file() else None}


def main() -> int:
    RUNS.mkdir(parents=True, exist_ok=True)
    results = [case_a(), case_b(), case_c()] + [case_d(v) for v in SPEC["case_D"]["variants"]]
    source_names = ("attachment_processor.py", "pipeline.py", "data_model.py", "field_validator.py",
                    "03 tests/T2-5/scripts/explore_trace.py", "03 tests/T2-5/scripts/run_approved_cases.py",
                    "03 tests/T2-5/fixtures/approved_cases.json")
    report = {"approval": SPEC["approval"], "scope": "formal controlled T2-5 cases",
              "source_sha256": {name: sha(ROOT / name) for name in source_names},
              "passed": sum(r["passed"] for r in results), "total": len(results), "results": results}
    path = BASE / "result" / f"approved_report_{STAMP}.json"
    path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"passed": report["passed"], "total": report["total"],
                      "failed": {r["case"]: r["differences"] for r in results if not r["passed"]},
                      "report": str(path)}, ensure_ascii=False, indent=2))
    return 0 if report["passed"] == report["total"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
