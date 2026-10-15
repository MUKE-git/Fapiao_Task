"""T2-4 independent, approved, controlled main-flow acceptance cases.

Input and business expectations are authored in fixtures/approved_cases.json.
The first-pass PDF extractor, IMAP and model response are controlled; PDF bytes
are real and the T2-4 arbitrator reads them with its production pdfplumber path.
"""
from __future__ import annotations

import hashlib
import json
import re
import sys
from datetime import datetime
from email.message import EmailMessage
from pathlib import Path
from unittest.mock import patch

import fitz
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[3]
sys.path.insert(0, str(ROOT))

from config import RunConfig, SUMMARY_COLUMNS  # noqa: E402
from data_model import NOT_FOUND  # noqa: E402
from pipeline import run_pipeline  # noqa: E402
from processor_contracts import DocumentResult  # noqa: E402

BASE = ROOT / "03 tests" / "T2-4"
FIXTURES = BASE / "fixtures"
RUN_TAG = datetime.now().strftime("%Y%m%d_%H%M%S")
RUNS = BASE / "result" / f"approved_cases_{RUN_TAG}"
MANIFEST = FIXTURES / "approved_cases.json"
SPEC = json.loads(MANIFEST.read_text(encoding="utf-8"))
DEFAULTS = SPEC["defaults"]
NUMBER = SPEC["invoice_number"]
DATE_HEADER = "Tue, 13 Oct 2026 10:00:00 +0800"


def sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def pdf_bytes(text: str) -> bytes:
    doc = fitz.open()
    page = doc.new_page()
    page.insert_font(fontname="SimHei", fontfile="C:/Windows/Fonts/simhei.ttf")
    page.insert_text((54, 72), text, fontname="SimHei", fontsize=12)
    data = doc.tobytes(garbage=4, deflate=True)
    doc.close()
    return data


def field_values(invoice: dict) -> dict:
    values = dict(DEFAULTS)
    values.update({k: invoice[k] for k in DEFAULTS if k in invoice})
    values["invoice_number"] = invoice.get("invoice_number", NUMBER) or NOT_FOUND
    return values


def expected_error(case: dict, index: int) -> str:
    mode = case["mode"]
    rid = f"Msg{index + 1}-01"
    if mode == "duplicate" and index == 1 or mode == "mixed" and index == 1:
        return f"{rid} 与 Msg1-01 的发票号及五个必填字段一致，已跳过重复 Excel 行；原记录保留在 JSON"
    if mode in ("conflict_unresolved", "mixed") and not (mode == "mixed" and index == 1):
        return f"{rid} 同票号 {NUMBER} 的必填字段冲突，AI 无法确认其一；本条保留 Excel 并需人工核对"
    if mode == "conflict_selected":
        if index == 0:
            notice = f"{rid} 与 Msg2-01 同票号但必填字段冲突；AI 采信后者，本条未写入 Excel，原记录保留在 JSON"
            return (f"T2-3 Msg1-02 行程单未匹配汇入，保留候选供人工核对;{notice}"
                    if case["id"] == "庚" else notice)
        evidence = "；".join(("报销金额" if item["field"] == "total_amount" else "销售方") + "：" + item["quote"] for item in case["evidence"])
        return f"同票号 {NUMBER} 存在字段冲突；AI 二次仲裁采信 Msg2-01；原文证据 {evidence}"
    return ""


def compare(path: str, actual, expected, differences: list[dict]) -> None:
    if type(actual) is not type(expected) or actual != expected:
        differences.append({"path": path, "expected": expected, "actual": actual})


class FakeImap:
    def __init__(self, messages: dict[str, bytes]):
        self.messages = messages

    def login(self, *_): return "OK", []
    def select(self, *_): return "OK", []
    def logout(self): return "BYE", []
    def status(self, *_): return "OK", [f"INBOX (MESSAGES {len(self.messages)} UNSEEN {len(self.messages)})".encode()]

    def uid(self, command, *args):
        if command.upper() == "SEARCH":
            return "OK", [" ".join(self.messages).encode()]
        uid = args[0].decode() if isinstance(args[0], bytes) else str(args[0])
        return "OK", [(b"body", self.messages[uid])]


def run_case(case: dict) -> dict:
    case_id = case["id"]
    case_dir = RUNS / case_id
    case_dir.mkdir(parents=True, exist_ok=True)
    cfg = RunConfig("imap.test", "user", "password", output_folder_name="output",
                    output_root=str(case_dir), dashscope_api_key="controlled-test-only")
    messages: dict[str, bytes] = {}
    source_hashes: dict[str, str] = {}
    input_fields: dict[str, dict] = {}
    for index, invoice in enumerate(case["invoices"]):
        msg = EmailMessage()
        msg["Subject"] = "电子发票"
        msg["Date"] = DATE_HEADER
        msg.set_content("T2-4 approved synthetic input")
        name = invoice["name"]
        pdf = pdf_bytes(invoice["pdf_text"])
        source_path = case_dir / f"input_{index + 1:02d}.pdf"
        source_path.write_bytes(pdf)
        source_hashes[str(source_path.relative_to(ROOT))] = sha(source_path)
        msg.add_attachment(pdf, maintype="application", subtype="pdf", filename=name)
        input_fields[name] = field_values(invoice)
        if "extra_pdf_name" in invoice:
            extra = pdf_bytes(invoice["extra_pdf_text"])
            extra_path = case_dir / "input_extra.pdf"
            extra_path.write_bytes(extra)
            source_hashes[str(extra_path.relative_to(ROOT))] = sha(extra_path)
            msg.add_attachment(extra, maintype="application", subtype="pdf", filename=invoice["extra_pdf_name"])
        raw = msg.as_bytes()
        eml_path = case_dir / f"input_{index + 1:02d}.eml"
        eml_path.write_bytes(raw)
        source_hashes[str(eml_path.relative_to(ROOT))] = sha(eml_path)
        messages[invoice["uid"]] = raw

    def controlled_extract(file, *_):
        if file.meta.original_name.endswith("行程单.pdf"):
            return DocumentResult(file.meta.file_seq, "Itinerary", {
                "match_amount": "50.00", "travel_time": "2026-10-13 09:00",
                "origin_dest": "A→B",
            }, [file.meta.file_seq])
        return DocumentResult(file.meta.file_seq, "Invoice", input_fields[file.meta.original_name],
                              [file.meta.file_seq])

    calls: list[dict] = []

    def controlled_ai(**kwargs):
        calls.append(kwargs)
        choice = case.get("ai_selection")
        return {"selected_record_id": None if choice is None else f"Msg{choice + 1}-01",
                "evidence": case.get("evidence", []), "reason": "controlled response"}

    with patch("pipeline.imaplib.IMAP4_SSL", return_value=FakeImap(messages)), \
         patch("pipeline.process_pdf", side_effect=controlled_extract), \
         patch("dedup_arbitrator.Generation.call", side_effect=controlled_ai):
        mismatches, stats = run_pipeline(cfg)

    differences: list[dict] = []
    compare("mismatches", mismatches, [], differences)
    compare("stats.imap_error", stats.get("imap_error"), None, differences)
    compare("stats.unseen_total", stats.get("unseen_total"), len(messages), differences)
    compare("AI.calls", len(calls), case["ai_calls"], differences)
    debug_path = Path(cfg.debug_file)
    actual = json.loads(debug_path.read_text(encoding="utf-8")) if debug_path.exists() else []
    compare("JSON.mail_count", len(actual), len(case["invoices"]), differences)
    for index, (invoice, mail) in enumerate(zip(case["invoices"], actual)):
        prefix = f"JSON[{index}]"
        compare(prefix + ".keys", sorted(mail), sorted(("mail_id", "mail_uid", "subject",
                "mail_date", "files", "itineraries", "records", "dedup_skipped",
                "error_reason")), differences)
        compare(prefix + ".mail_id", mail.get("mail_id"), f"Msg{index + 1}", differences)
        compare(prefix + ".mail_uid", mail.get("mail_uid"), invoice["uid"], differences)
        compare(prefix + ".subject", mail.get("subject"), "电子发票", differences)
        compare(prefix + ".mail_date", mail.get("mail_date"), "2026-10-13 10:00", differences)
        compare(prefix + ".record_count", len(mail.get("records", [])), 1, differences)
        compare(prefix + ".dedup_skipped", mail.get("dedup_skipped"), [], differences)
        compare(prefix + ".error_reason", mail.get("error_reason"), expected_error(case, index), differences)
        files = mail.get("files", [])
        expected_file_count = 2 if "extra_pdf_name" in invoice else 1
        compare(prefix + ".file_count", len(files), expected_file_count, differences)
        for fidx, file in enumerate(files):
            seq = f"{fidx + 1:02d}"
            name = invoice["name"] if fidx == 0 else invoice["extra_pdf_name"]
            compare(f"{prefix}.files[{fidx}].keys", sorted(file), sorted(("file_seq", "file_type",
                    "original_name", "pre_read_type", "pre_read_invoice_no", "pre_read_amount",
                    "internal_name", "exported_name")), differences)
            expected_file = {"file_seq": seq, "file_type": "PDF", "original_name": name,
                             "pre_read_type": "Unknown", "pre_read_invoice_no": NOT_FOUND,
                             "pre_read_amount": NOT_FOUND, "exported_name": ""}
            for key, value in expected_file.items():
                compare(f"{prefix}.files[{fidx}].{key}", file.get(key), value, differences)
            internal = file.get("internal_name", "")
            if not re.fullmatch(rf"Msg{index + 1}_{seq}_.+\.pdf", internal):
                differences.append({"path": f"{prefix}.files[{fidx}].internal_name", "expected": "Msg序号_文件序号_原名.pdf", "actual": internal})
            saved_path = Path(cfg.extract_dir) / internal
            if not saved_path.is_file():
                differences.append({"path": f"{prefix}.files[{fidx}].stored_file", "expected": "exists", "actual": "missing"})
            else:
                original_path = case_dir / (f"input_{index + 1:02d}.pdf" if fidx == 0 else "input_extra.pdf")
                compare(f"{prefix}.files[{fidx}].bytes_sha256", sha(saved_path), sha(original_path), differences)
        if mail.get("records"):
            record = mail["records"][0]
            values = field_values(invoice)
            expected_fields = {"invoice_date": values["invoice_date"],
                               "invoice_number": values["invoice_number"],
                               "total_amount": values["total_amount"],
                               "invoice_type": values["invoice_type"],
                               "receipt_type": values["receipt_type"],
                               "seller": values["seller"]}
            compare(prefix + ".records[0]", record, {
                "record_id": f"Msg{index + 1}-01", "kind": "b类", "fields": expected_fields,
                "source": {"source_file_seq": ["01"]}, "status": "人工待核",
            }, differences)
        expected_itineraries = ([{"record_id": f"Msg{index + 1}-02", "match_amount": "50.00",
                                  "travel_time": "2026-10-13 09:00", "origin_dest": "A→B"}]
                                if "extra_pdf_name" in invoice else [])
        compare(prefix + ".itineraries", mail.get("itineraries"), expected_itineraries, differences)

    excel_path = Path(cfg.summary_excel_file)
    expected_headers = [h for h, _ in SUMMARY_COLUMNS]
    expected_rows = []
    for selected in case["excel_from"]:
        fields = field_values(case["invoices"][selected])
        expected_rows.append([fields["invoice_date"], fields["invoice_number"],
                              float(fields["total_amount"]), fields["invoice_type"],
                              fields["receipt_type"], None, fields["seller"], None])
    if not excel_path.is_file():
        differences.append({"path": "Excel.file", "expected": "exists", "actual": "missing"})
    else:
        book = load_workbook(excel_path)
        sheet = book.active
        compare("Excel.headers", [sheet.cell(1, c).value for c in range(1, 9)], expected_headers, differences)
        compare("Excel.row_count", sheet.max_row - 1, len(expected_rows), differences)
        for row_idx, expected_row in enumerate(expected_rows, 2):
            for col_idx, expected_value in enumerate(expected_row, 1):
                cell = sheet.cell(row_idx, col_idx)
                if col_idx == 3 and isinstance(cell.value, (float, int)):
                    compare(f"Excel.row[{row_idx - 1}].col[{col_idx}].numeric_value",
                            float(cell.value), expected_value, differences)
                elif expected_value == "" and cell.value is None:
                    pass  # Empty spreadsheet cells reopen as None in openpyxl.
                else:
                    compare(f"Excel.row[{row_idx - 1}].col[{col_idx}]", cell.value, expected_value, differences)
                if col_idx == 3:
                    compare(f"Excel.row[{row_idx - 1}].amount_format", cell.number_format, "0.00", differences)
                    if cell.data_type != "n":
                        differences.append({"path": f"Excel.row[{row_idx - 1}].amount_type", "expected": "number", "actual": cell.data_type})
        book.close()

    for call_index, call in enumerate(calls):
        messages = call.get("messages", [])
        if len(messages) != 2:
            differences.append({"path": f"AI[{call_index}].messages", "expected": 2, "actual": len(messages)})
            continue
        payload = json.loads(messages[1]["content"])
        compare("AI.different_fields", payload.get("different_fields"), case["different_fields"], differences)
        compare("AI.candidate_count", len(payload.get("candidates", [])),
                len(case["invoices"]) - (1 if case_id == "戊" else 0), differences)
        groups = payload.get("attachment_groups", [])
        expected_group_uids = [case["invoices"][i]["uid"] for i in range(len(case["invoices"])) if not (case_id == "戊" and i == 1)]
        compare("AI.attachment_group_uids", [group.get("mail_uid") for group in groups], expected_group_uids, differences)
        for i, group in enumerate(groups):
            original_index = next(j for j, inv in enumerate(case["invoices"]) if inv["uid"] == group["mail_uid"])
            invoice = case["invoices"][original_index]
            expected_attachments = [{"file_seq": "01", "file_type": "PDF", "original_name": invoice["name"], "pdf_text": invoice["pdf_text"]}]
            if "extra_pdf_name" in invoice:
                expected_attachments.append({"file_seq": "02", "file_type": "PDF", "original_name": invoice["extra_pdf_name"], "pdf_text": invoice["extra_pdf_text"]})
            compare(f"AI.groups[{i}].attachments", group.get("attachments"), expected_attachments, differences)
            candidate = payload["candidates"][i]
            compare(f"AI.candidates[{i}].record_id", candidate.get("record_id"), f"Msg{original_index + 1}-01", differences)
            compare(f"AI.candidates[{i}].pdf_text", candidate.get("pdf_text"), invoice["pdf_text"], differences)
            compare(f"AI.candidates[{i}].source_files", candidate.get("source_files"),
                    [{"file_seq": "01", "original_name": invoice["name"]}], differences)
            compare(f"AI.candidates[{i}].mail_uid", candidate.get("mail_uid"), invoice["uid"], differences)
            for field in ("invoice_date", "invoice_number", "total_amount", "receipt_type", "seller"):
                compare(f"AI.candidates[{i}].fields.{field}", candidate.get("fields", {}).get(field), field_values(invoice)[field], differences)

    output_hashes = {str(path.relative_to(ROOT)): sha(path) for path in (debug_path, excel_path) if path.is_file()}
    return {"case": case_id, "passed": not differences, "differences": differences,
            "source_hashes": source_hashes, "output_hashes": output_hashes,
            "ai_calls": len(calls), "stats_alerts": stats.get("not_found_alerts", [])}


def main() -> int:
    RUNS.mkdir(parents=True, exist_ok=True)
    results = [run_case(case) for case in SPEC["cases"]]
    source_files = ["dedup_arbitrator.py", "pipeline.py", "data_model.py", "stage_ports.py",
                    "attachment_processor.py", "excel_writer.py", "processor_contracts.py",
                    "03 tests/T2-4/scripts/run_approved_cases.py", "03 tests/T2-4/fixtures/approved_cases.json"]
    report = {"approval": SPEC["approval"], "execution": "controlled main-flow; real PDF byte storage/read",
              "source_versions": {name: sha(ROOT / name) for name in source_files},
              "passed": sum(result["passed"] for result in results), "total": len(results), "cases": results}
    path = BASE / "result" / f"approved_cases_report_{RUN_TAG}.json"
    path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"passed": report["passed"], "total": report["total"],
                      "failed": {r["case"]: r["differences"] for r in results if not r["passed"]},
                      "report": str(path)}, ensure_ascii=False, indent=2))
    return 0 if report["passed"] == report["total"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
