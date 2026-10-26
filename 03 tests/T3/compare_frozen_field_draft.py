"""Compare frozen cases with the separate, review-pending full-field draft.

The original frozen runner supplies controlled PDF extraction and performs its
unchanged approved checks. This script only captures those same run artifacts
and compares an additional prewritten answer; it never edits the frozen bank.
"""
from __future__ import annotations

import importlib.util
import json
import os
import shutil
from pathlib import Path

from openpyxl import load_workbook

from compare_full_fields import _valid_format
from prepare_frozen_blob_copy import prepare


HERE = Path(__file__).resolve().parent
PROJECT = HERE.parents[1]
DRAFT = HERE / "frozen_18_full_fields_draft.json"
RESULT = HERE / "result" / "frozen_field_draft"
COLUMNS = (
    ("发票日期", "invoice_date"), ("发票号码", "invoice_number"),
    ("报销金额", "total_amount"), ("发票类型", "invoice_type"),
    ("票面类型", "receipt_type"), ("出行时间", "travel_time"),
    ("销售方", "seller"), ("出发地/目的地", "origin_dest"),
)
RULES = {
    "invoice_date": "date_or_missing",
    "invoice_number": "20_digits_or_missing",
    "total_amount": "decimal_string_4dp_or_missing",
    "invoice_type": "text_or_empty",
    "receipt_type": "text_or_missing",
    "seller": "text_or_missing",
    "match_amount": "decimal_string_4dp_or_missing",
    "travel_time": "datetime_minute_or_empty",
    "origin_dest": "text_or_empty",
}


def validate_draft(draft: dict) -> None:
    """Reject an incomplete or silently altered review artifact before testing."""
    if draft.get("review_status") != "测试工程师制作、待用户审核":
        raise ValueError("Only the separate pending-review draft may run here")
    schema = draft["schema"]
    if schema["excel_columns"] != [[title, key] for title, key in COLUMNS]:
        raise ValueError("Draft eight-column mapping differs from the agreed mapping")
    if schema["json_invoice_fields"] != [
            "invoice_date", "invoice_number", "total_amount", "invoice_type",
            "receipt_type", "seller", "travel_time", "origin_dest"]:
        raise ValueError("Draft JSON invoice field declaration incomplete")
    if schema["json_itinerary_fields"] != ["match_amount", "travel_time", "origin_dest"]:
        raise ValueError("Draft JSON itinerary field declaration incomplete")
    if set(schema["formats"]) != {
            "invoice_date", "invoice_number", "total_amount_json",
            "total_amount_excel", "invoice_type", "receipt_type", "seller",
            "match_amount", "travel_time", "origin_dest"}:
        raise ValueError("Draft field format declarations incomplete")
    if schema["empty_values"] != {
            "json_unmatched_travel_fields": "absent",
            "json_unmatched_itinerary": "absent",
            "excel_unmatched_travel_fields": "",
            "required_missing": "Not Found"}:
        raise ValueError("Draft empty-value rules differ from the agreed rules")
    if schema["excel_amount_number_format"] != "0.00":
        raise ValueError("Draft Excel amount format must be 0.00")
    original_index = json.loads((HERE.parents[0] / "owner test" / "固定用例库" /
                                 "index.json").read_text(encoding="utf-8"))
    original_order = [(case["id"], variant) for case in original_index["cases"]
                      for variant in case["variants"]]
    draft_order = [(item["case_id"], item["variant"]) for item in draft["variants"]]
    if len(original_index["cases"]) != 18 or draft_order != original_order:
        raise ValueError("Draft must preserve the approved 18-case, 26-variant order")
    invoice_keys = {"invoice_date", "invoice_number", "total_amount",
                    "invoice_type", "receipt_type", "seller"}
    trip_keys = {"match_amount", "travel_time", "origin_dest"}
    excel_keys = {key for _, key in COLUMNS}
    for item in draft["variants"]:
        if (len(item["json_invoice_records"]) != item["counts"]["record_count"]
                or len(item["json_itinerary_records"]) != item["counts"]["itinerary_count"]
                or len(item["excel_rows"]) != item["counts"]["excel_rows"]):
            raise ValueError(f"Draft counts incomplete: {item['case_id']}/{item['variant']}")
        for record in item["json_invoice_records"]:
            target = record["matched_itinerary"]
            expected_keys = invoice_keys | ({"travel_time", "origin_dest"} if target else set())
            if set(record["fields"]) != expected_keys:
                raise ValueError(f"Draft invoice fields incomplete: {item['case_id']}/{item['variant']}")
            if target is not None and set(target) != {"original_name", "match_amount"}:
                raise ValueError("Draft match relation incomplete")
        for trip in item["json_itinerary_records"]:
            if set(trip["fields"]) != trip_keys:
                raise ValueError("Draft itinerary fields incomplete")
        for row in item["excel_rows"]:
            if set(row["fields"]) != excel_keys:
                raise ValueError("Draft Excel row must contain all eight fields")


def _field_compare(wanted: dict, actual: dict, path: str, errors: list[str]) -> None:
    if not isinstance(actual, dict):
        errors.append(f"{path}: expected object, actual {actual!r}")
        return
    if set(actual) != set(wanted):
        errors.append(f"{path}.keys: expected {sorted(wanted)}, actual {sorted(actual)}")
    for key, expected_value in wanted.items():
        observed_value = actual.get(key, "__ABSENT__")
        if observed_value != expected_value or type(observed_value) is not type(expected_value):
            errors.append(f"{path}.{key}: expected {expected_value!r} ({type(expected_value).__name__}), actual {observed_value!r} ({type(observed_value).__name__})")
        rule = RULES[key]
        if not _valid_format(rule, expected_value):
            errors.append(f"{path}.{key}: draft expectation violates {rule}: {expected_value!r}")
        if observed_value != "__ABSENT__" and not _valid_format(rule, observed_value):
            errors.append(f"{path}.{key}: actual violates {rule}: {observed_value!r}")


def compare_variant(wanted: dict, actual_json: Path, actual_excel: Path) -> list[str]:
    errors: list[str] = []
    mails = json.loads(actual_json.read_text(encoding="utf-8"))
    if len(mails) != wanted["counts"]["selected_mails"]:
        errors.append(f"json.mails: expected {wanted['counts']['selected_mails']}, actual {len(mails)}")
    invoice_records = []
    itinerary_records = []
    for mail in mails:
        names = {file["file_seq"]: file["original_name"] for file in mail.get("files", [])}
        for record in mail.get("records", []):
            source = (record.get("source") or {}).get("source_file_seq") or []
            name = names.get(source[0], f"unresolved:{source[0]}") if source else "no source"
            target = record.get("matched_itinerary")
            invoice_records.append({
                "original_name": name, "fields": record.get("fields", {}),
                "matched_itinerary": ({
                    "original_name": names.get(target["file_seq"], f"unresolved:{target['file_seq']}"),
                    "match_amount": target["match_amount"],
                } if target else None),
            })
        for trip in mail.get("itineraries", []):
            seq = trip.get("record_id", "").rsplit("-", 1)[-1]
            itinerary_records.append({
                "original_name": names.get(seq, f"unresolved:{seq}"),
                "fields": {key: trip.get(key, "__ABSENT__") for key in
                           ("match_amount", "travel_time", "origin_dest")},
            })
    for label, expected_items, observed_items in (
        ("json_invoice_records", wanted["json_invoice_records"], invoice_records),
        ("json_itinerary_records", wanted["json_itinerary_records"], itinerary_records),
    ):
        if len(expected_items) != len(observed_items):
            errors.append(f"{label}.count: expected {len(expected_items)}, actual {len(observed_items)}")
        for i, (left, right) in enumerate(zip(expected_items, observed_items)):
            path = f"{label}[{i}]"
            if right["original_name"] != left["original_name"]:
                errors.append(f"{path}.original_name: expected {left['original_name']!r}, actual {right['original_name']!r}")
            _field_compare(left["fields"], right["fields"], path + ".fields", errors)
            if label == "json_invoice_records":
                if right["matched_itinerary"] != left["matched_itinerary"]:
                    errors.append(f"{path}.matched_itinerary: expected {left['matched_itinerary']!r}, actual {right['matched_itinerary']!r}")

    expected_rows = wanted["excel_rows"]
    if not actual_excel.is_file():
        errors.append(f"excel.file: missing {actual_excel}")
        return errors
    book = load_workbook(actual_excel, read_only=False, data_only=True)
    try:
        sheet = book.active
        values = list(sheet.iter_rows(values_only=True))
        expected_headers = [title for title, _ in COLUMNS]
        headers = list(values[0]) if values else []
        if headers != expected_headers:
            errors.append(f"excel.headers: expected {expected_headers!r}, actual {headers!r}")
        rows = values[1:]
        if len(rows) != len(expected_rows):
            errors.append(f"excel.rows: expected {len(expected_rows)}, actual {len(rows)}")
        for ri, (left, actual_row) in enumerate(zip(expected_rows, rows)):
            expected_fields = left["fields"]
            if set(expected_fields) != {key for _, key in COLUMNS}:
                errors.append(f"excel.rows[{ri}]: draft must specify all eight columns")
            for ci, (_, key) in enumerate(COLUMNS):
                want = expected_fields[key]
                got = actual_row[ci] if ci < len(actual_row) else "__ABSENT__"
                if got is None and want == "":
                    got = ""
                path = f"excel.rows[{ri}].{key}"
                if key == "total_amount" and want != "Not Found":
                    from decimal import Decimal, ROUND_HALF_UP
                    expected_number = Decimal(want).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
                    if isinstance(got, bool) or not isinstance(got, (int, float)) or Decimal(str(got)) != expected_number:
                        errors.append(f"{path}: expected numeric {expected_number}, actual {got!r}")
                    if sheet.cell(ri + 2, ci + 1).number_format != "0.00":
                        errors.append(f"{path}.number_format: expected '0.00', actual {sheet.cell(ri + 2, ci + 1).number_format!r}")
                elif got != want or type(got) is not type(want):
                    errors.append(f"{path}: expected {want!r}, actual {got!r}")
                if key != "total_amount":
                    rule = RULES[key]
                    if not _valid_format(rule, want):
                        errors.append(f"{path}: draft expectation violates {rule}: {want!r}")
                    if got != "__ABSENT__" and not _valid_format(rule, got):
                        errors.append(f"{path}: actual violates {rule}: {got!r}")
    finally:
        book.close()
    return errors


def run() -> dict:
    draft = json.loads(DRAFT.read_text(encoding="utf-8"))
    validate_draft(draft)
    prepared = prepare()
    os.environ["FIXED_CASE_PROJECT_ROOT"] = str(PROJECT)
    runner_path = Path(prepared["temporary_case_root"]) / "run_fixed_cases.py"
    spec = importlib.util.spec_from_file_location("unchanged_frozen_runner_draft", runner_path)
    runner = importlib.util.module_from_spec(spec)
    assert spec and spec.loader
    spec.loader.exec_module(runner)
    original_run_pipeline = runner.run_pipeline
    report = []
    for item in draft["variants"]:
        case_id, variant = item["case_id"], item["variant"]
        folder = RESULT / case_id / variant
        folder.mkdir(parents=True, exist_ok=True)
        # Never let an earlier run supply a missing artifact for this run.
        (folder / "actual.json").unlink(missing_ok=True)
        (folder / "actual.xlsx").unlink(missing_ok=True)
        def captured_pipeline(cfg):
            output = original_run_pipeline(cfg)
            debug = Path(cfg.debug_file)
            excel = Path(cfg.summary_excel_file)
            if debug.is_file():
                shutil.copyfile(debug, folder / "actual.json")
            if excel.is_file():
                shutil.copyfile(excel, folder / "actual.xlsx")
            return output
        runner.run_pipeline = captured_pipeline
        frozen = runner.execute_variant(case_id, variant, real_pdf=False)
        actual_json = folder / "actual.json"
        actual_excel = folder / "actual.xlsx"
        errors = ([] if actual_json.is_file() else ["json.file: missing captured output"])
        if not errors:
            errors.extend(compare_variant(item, actual_json, actual_excel))
        report.append({"case_id": case_id, "variant": variant,
                       "original_frozen_status": frozen["status"],
                       "original_frozen_errors": frozen["errors"],
                       "draft_status": "PASS" if not errors else "FAIL",
                       "draft_errors": errors})
    result = {"mode": "草案比较；非正式验收",
              "review_status": draft["review_status"],
              "called": "未修改的 run_fixed_cases.execute_variant 和生产 run_pipeline",
              "replaced": "IMAP 固定 EML 与每份 PDF 的 extract.json 文件级返回值",
              "checked": "原已审核检查 + 扩展 JSON 发票/行程单业务字段、匹配关系、Excel 八列表头/行序/单元格/类型/金额格式",
              "variants": report,
              "frozen_pass": sum(x["original_frozen_status"] == "PASS" for x in report),
              "draft_pass": sum(x["draft_status"] == "PASS" for x in report)}
    RESULT.mkdir(parents=True, exist_ok=True)
    (RESULT / "comparison.json").write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n",
                                            encoding="utf-8")
    return result


if __name__ == "__main__":
    output = run()
    print(f"frozen={output['frozen_pass']}/26 draft={output['draft_pass']}/26")
    for row in output["variants"]:
        if row["original_frozen_status"] != "PASS" or row["draft_errors"]:
            print(f"{row['case_id']}/{row['variant']}: frozen={row['original_frozen_status']} draft={row['draft_status']}")
            for difference in row["original_frozen_errors"] + row["draft_errors"]:
                print("  " + difference)
