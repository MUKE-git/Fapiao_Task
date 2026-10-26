"""Independent JSON/Excel comparison against a prewritten T3 expectation.

The expectation must be authored from approved rules and input evidence before
the program is run. This module never copies observed values into expectations.
"""
from __future__ import annotations

import argparse
import json
import re
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook


COLUMNS = [
    ("发票日期", "invoice_date"), ("发票号码", "invoice_number"),
    ("报销金额", "total_amount"), ("发票类型", "invoice_type"),
    ("票面类型", "receipt_type"), ("出行时间", "travel_time"),
    ("销售方", "seller"), ("出发地/目的地", "origin_dest"),
]
INVOICE_FIELDS = {"invoice_date", "invoice_number", "total_amount",
                  "invoice_type", "receipt_type", "seller"}
TRIP_FIELDS = {"match_amount", "travel_time", "origin_dest"}
SCHEMA_FORMATS = {
    "invoice_date": "date_or_missing",
    "invoice_number": "20_digits_or_missing",
    "total_amount": "decimal_string_4dp_or_missing",
    "invoice_type": "text_or_empty",
    "receipt_type": "text_or_missing",
    "travel_time": "datetime_minute_or_empty",
    "seller": "text_or_missing",
    "origin_dest": "text_or_empty",
    "match_amount": "decimal_string_4dp_or_missing",
}


def _valid_format(rule: str, value: object) -> bool:
    if rule == "raw":
        return True
    if rule.endswith("_or_missing") and value == "Not Found":
        return True
    if rule.endswith("_or_empty") and value == "":
        return True
    if rule in {"text_or_empty", "text_or_missing"}:
        return isinstance(value, str) and bool(value.strip())
    if rule == "20_digits_or_missing":
        return isinstance(value, str) and re.fullmatch(r"[0-9]{20}", value) is not None
    if rule == "date_or_missing":
        if not isinstance(value, str) or not re.fullmatch(r"[0-9]{4}-[0-9]{2}-[0-9]{2}", value):
            return False
        try:
            return date.fromisoformat(value).isoformat() == value
        except ValueError:
            return False
    if rule == "datetime_minute_or_empty":
        if not isinstance(value, str) or not re.fullmatch(r"[0-9]{4}-[0-9]{2}-[0-9]{2} [0-9]{2}:[0-9]{2}", value):
            return False
        try:
            return datetime.strptime(value, "%Y-%m-%d %H:%M").strftime("%Y-%m-%d %H:%M") == value
        except ValueError:
            return False
    if rule == "decimal_string_4dp_or_missing":
        if not isinstance(value, str) or not re.fullmatch(r"[0-9]+(?:\.[0-9]{1,4})?", value):
            return False
        try:
            return Decimal(value) > 0
        except InvalidOperation:
            return False
    return False


def _schema_errors(schema: dict) -> list[str]:
    errors = []
    if schema.get("excel_columns") != [[title, key] for title, key in COLUMNS]:
        errors.append("expectation.schema.excel_columns: must declare the exact eight-column mapping")
    formats = schema.get("formats")
    if not isinstance(formats, dict) or set(formats) != set(SCHEMA_FORMATS):
        errors.append("expectation.schema.formats: must declare all nine business fields")
    elif any(rule not in set(SCHEMA_FORMATS.values()) | {"raw"} for rule in formats.values()):
        errors.append("expectation.schema.formats: unknown format rule")
    if schema.get("empty_values") != {
            "invoice_type": "", "travel_time": "", "origin_dest": "",
            "required_missing": "Not Found"}:
        errors.append("expectation.schema.empty_values: explicit empty-value rules required")
    if schema.get("excel_amount_number_format") != "0.00":
        errors.append("expectation.schema.excel_amount_number_format: expected 0.00")
    return errors


def _money(value: object, path: str, errors: list[str]) -> Decimal | None:
    try:
        result = Decimal(str(value))
    except (InvalidOperation, ValueError, TypeError):
        errors.append(f"{path}: expected decimal-compatible amount, actual {value!r}")
        return None
    if not result.is_finite():
        errors.append(f"{path}: non-finite amount {value!r}")
        return None
    return result


def _compare_value(expected: object, actual: object, path: str,
                   errors: list[str], *, excel: bool = False) -> None:
    if expected == "__ABSENT__":
        if actual is not _ABSENT:
            errors.append(f"{path}: expected absent, actual {actual!r}")
        return
    if actual is _ABSENT:
        errors.append(f"{path}: missing; expected {expected!r}")
        return
    if path.endswith(("total_amount", "match_amount")) and expected != "Not Found":
        wanted = _money(expected, path + ".expected", errors)
        found = _money(actual, path + ".actual", errors)
        if wanted is not None and found is not None:
            if excel and found != wanted.quantize(Decimal("0.01")):
                errors.append(f"{path}: expected Excel {wanted:.2f}, actual {actual!r}")
            elif not excel and (not isinstance(actual, str) or actual != expected):
                errors.append(f"{path}: expected JSON decimal string {expected!r}, actual {actual!r}")
        return
    if expected != actual:
        errors.append(f"{path}: expected {expected!r}, actual {actual!r}")


_ABSENT = object()


def _compare_fields(expected: dict, actual: dict, required: set[str], path: str,
                    errors: list[str], formats: dict[str, str]) -> None:
    if not isinstance(actual, dict):
        errors.append(f"{path}: expected object, actual {actual!r}")
        return
    if not required <= set(expected):
        errors.append(f"{path}: expectation incomplete, missing {sorted(required - set(expected))}")
    for key, wanted in expected.items():
        _compare_value(wanted, actual.get(key, _ABSENT), f"{path}.{key}", errors)
        rule = formats.get(key, "raw")
        if wanted != "__ABSENT__" and not _valid_format(rule, wanted):
            errors.append(f"{path}.{key}: expectation violates declared format {rule}")
        found = actual.get(key, _ABSENT)
        if found is not _ABSENT and not _valid_format(rule, found):
            errors.append(f"{path}.{key}: actual violates declared format {rule}: {found!r}")
    extra = set(actual) - set(expected)
    if extra:
        errors.append(f"{path}: unexpected business fields {sorted(extra)}")


def compare_json(expected: dict, observed: list[dict], formats: dict[str, str] | None = None) -> list[str]:
    """Compare every declared business field, counts and source identities."""
    errors: list[str] = []
    formats = formats or {}
    wanted_mails = expected.get("mails")
    if not isinstance(wanted_mails, list):
        return ["expectation.mails must be a list"]
    if not isinstance(observed, list) or len(observed) != len(wanted_mails):
        return [f"json.mails: expected {len(wanted_mails)}, actual {len(observed) if isinstance(observed, list) else type(observed).__name__}"]
    for mi, (wanted_mail, mail) in enumerate(zip(wanted_mails, observed)):
        path = f"json.mails[{mi}]"
        for key in ("mail_id", "mail_uid", "subject", "mail_date"):
            _compare_value(wanted_mail[key], mail.get(key, _ABSENT), f"{path}.{key}", errors)
        for collection in ("files", "records", "itineraries"):
            wanted = wanted_mail.get(collection, [])
            actual = mail.get(collection, [])
            if len(actual) != len(wanted):
                errors.append(f"{path}.{collection}: expected {len(wanted)}, actual {len(actual)}")
            for ri, (left, right) in enumerate(zip(wanted, actual)):
                item_path = f"{path}.{collection}[{ri}]"
                if collection == "records":
                    for key in ("record_id", "kind", "status"):
                        _compare_value(left[key], right.get(key, _ABSENT), f"{item_path}.{key}", errors)
                    _compare_fields(left["fields"], right.get("fields"), INVOICE_FIELDS,
                                    item_path + ".fields", errors, formats)
                    _compare_value(left["source_file_seq"],
                                   (right.get("source") or {}).get("source_file_seq", _ABSENT),
                                   item_path + ".source_file_seq", errors)
                    _compare_value(left.get("matched_itinerary", "__ABSENT__"),
                                   right.get("matched_itinerary", _ABSENT),
                                   item_path + ".matched_itinerary", errors)
                elif collection == "itineraries":
                    _compare_value(left["record_id"], right.get("record_id", _ABSENT),
                                   item_path + ".record_id", errors)
                    _compare_fields(left["fields"],
                                    {key: right.get(key, _ABSENT) for key in TRIP_FIELDS},
                                    TRIP_FIELDS,
                                    item_path, errors, formats)
                    if "source_file_seq" in left:
                        _compare_value(left["source_file_seq"],
                                       right.get("source_file_seq", _ABSENT),
                                       item_path + ".source_file_seq", errors)
                else:
                    for key in ("file_seq", "file_type", "original_name"):
                        _compare_value(left[key], right.get(key, _ABSENT), f"{item_path}.{key}", errors)
    return errors


def compare_excel(expected: dict, path: Path, formats: dict[str, str] | None = None) -> list[str]:
    errors: list[str] = []
    formats = formats or {}
    wanted_rows = expected.get("rows")
    if not isinstance(wanted_rows, list):
        return ["expectation.excel.rows must be a list"]
    if not path.is_file():
        return [] if not wanted_rows and expected.get("allow_missing_file") else [f"excel: missing file {path}"]
    book = load_workbook(path, read_only=False, data_only=True)
    try:
        sheet = book.active
        values = list(sheet.iter_rows(values_only=True))
        headers = list(values[0]) if values else []
        expected_headers = [title for title, _ in COLUMNS]
        if headers != expected_headers:
            errors.append(f"excel.headers: expected {expected_headers!r}, actual {headers!r}")
        actual_rows = values[1:]
        if len(actual_rows) != len(wanted_rows):
            errors.append(f"excel.rows: expected {len(wanted_rows)}, actual {len(actual_rows)}")
        for ri, (wanted, actual) in enumerate(zip(wanted_rows, actual_rows)):
            if set(wanted) != {key for _, key in COLUMNS}:
                errors.append(f"expectation.excel.rows[{ri}]: must declare all eight fields")
            for ci, (_, key) in enumerate(COLUMNS):
                value = actual[ci] if ci < len(actual) else _ABSENT
                if value is None and wanted.get(key) == "":
                    value = ""
                _compare_value(wanted.get(key, _ABSENT), value,
                               f"excel.rows[{ri}].{key}", errors, excel=True)
                rule = formats.get(key, "raw")
                if key != "total_amount":
                    if not _valid_format(rule, wanted.get(key)):
                        errors.append(f"excel.rows[{ri}].{key}: expectation violates declared format {rule}")
                    if value is not _ABSENT and not _valid_format(rule, value):
                        errors.append(f"excel.rows[{ri}].{key}: actual violates declared format {rule}: {value!r}")
                if key == "total_amount" and wanted.get(key) != "Not Found" and value is not _ABSENT:
                    if not isinstance(value, (int, float)) or isinstance(value, bool):
                        errors.append(f"excel.rows[{ri}].total_amount: expected numeric cell, actual {type(value).__name__}")
                    elif sheet.cell(row=ri + 2, column=ci + 1).number_format != "0.00":
                        errors.append(f"excel.rows[{ri}].total_amount: expected 0.00 display format")
    finally:
        book.close()
    return errors


def compare_files(expectation: Path, actual_json: Path, actual_excel: Path,
                  mode: str = "draft") -> dict:
    manifest = json.loads(expectation.read_text(encoding="utf-8"))
    status = manifest.get("review_status")
    if mode == "formal" and status != "已由用户审核":
        return {"mode": mode, "status": "REFUSED", "errors": ["预期尚未由用户审核，不能作为正式验收"]}
    schema = manifest.get("schema") or {}
    schema_errors = _schema_errors(schema)
    observed = json.loads(actual_json.read_text(encoding="utf-8"))
    errors = list(schema_errors)
    errors.extend(compare_json(manifest["json"], observed, schema.get("formats") or {}))
    errors.extend(compare_excel(manifest["excel"], actual_excel, schema.get("formats") or {}))
    return {"mode": mode, "review_status": status,
            "status": "PASS" if not errors else "FAIL", "errors": errors}


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("expected", type=Path)
    parser.add_argument("actual_json", type=Path)
    parser.add_argument("actual_excel", type=Path)
    parser.add_argument("--mode", choices=("draft", "formal"), default="draft")
    args = parser.parse_args()
    print(json.dumps(compare_files(args.expected, args.actual_json, args.actual_excel,
                                   args.mode), ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
