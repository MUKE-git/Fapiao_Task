"""测试工程师制作、待用户审核：真实文件 + 固定外部响应的整链草案比较。

样本为用户本机原件；不复制到仓库。PDF 的外部 AI 答复、图片的
多模态候选和邮件服务器分别由预先写定的响应替换。处理器、匹配、
Excel/JSON 和 T3-8 ZIP 在集成分支上实际执行。
"""
from __future__ import annotations

import io
import json
import sys
import tempfile
import unittest
import zipfile
from email.message import EmailMessage
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

import pdfplumber
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))
from ai_client import ModelFileOutput
from config import RunConfig
from mail_link_downloader import DownloadedFile
from pipeline import run_pipeline

SAMPLES = Path(r"C:\Users\lenovo\Downloads")
XML = SAMPLES / "T3出行电子发票(XML)A.xml"
PDF = SAMPLES / "T3出行电子发票A (1).pdf"
IMAGE = SAMPLES / "发票文件1.png"
OFD_A = SAMPLES / "T3出行电子发票(OFD)A.ofd"
OFD_COMPLETE = SAMPLES / "【呼我出行-48.31元】百度地图打车电子发票.ofd"
NUMBER_A = "26447000000342589935"
NUMBER_B = "26522000000868668241"
HEADERS = ["发票日期", "发票号码", "报销金额", "发票类型", "票面类型", "出行时间", "销售方", "出发地/目的地"]


def mail(*files: Path, body: str = "固定测试邮件") -> EmailMessage:
    msg = EmailMessage()
    msg["Subject"] = "电子发票：测试工程师制作，待用户审核"
    msg["Date"] = "Mon, 27 Apr 2026 09:00:00 +0800"
    msg.set_content(body)
    for file in files:
        msg.add_attachment(file.read_bytes(), maintype="application",
                           subtype="octet-stream", filename=file.name)
    return msg


class FakeImap:
    def __init__(self, message: EmailMessage):
        self.payload = message.as_bytes()
        self.fetch_queries = []

    def login(self, *_args): return "OK", []
    def select(self, *_args): return "OK", [b"1"]
    def status(self, *_args): return "OK", [b"INBOX (MESSAGES 1 UNSEEN 1)"]
    def logout(self): return "BYE", []
    def uid(self, command, *args):
        if command == "SEARCH":
            return "OK", [b"101"]
        self.fetch_queries.append(args[-1])
        return "OK", [(b"1 (UID 101)", self.payload)]


def controlled_pdf_ai(task_info, **_kwargs):
    source_text = task_info["files"][0]["text_snapshot"]
    assert NUMBER_A in source_text and "121.74" in source_text
    internal = Path(task_info["files"][0]["local_path"]).name
    return ({
        "file_classifications": [{"file": internal, "role": "Invoice"}],
        "invoice_date": "2026-04-27", "invoice_number": NUMBER_A,
        "total_amount": "121.74", "invoice_type": "电子发票（普通发票）",
        "receipt_type": "客运服务费",
        "seller": "南京领行科技股份有限公司广州分公司",
    }, "", [], [])


def controlled_image_model(request, _cfg):
    assert request.file_type == "JPEG"
    assert request.original_name == IMAGE.name
    fields = {
        "invoice_date": "2026-08-24", "invoice_number": NUMBER_B,
        "total_amount": "21.60", "invoice_type": "电子发票（普通发票）",
        "receipt_type": "驱蚊花露水",
        "seller": "贵阳皓康医疗器械有限公司",
    }
    evidence = {key: [{"file_seq": request.file_seq, "raw": value,
                       "normalized": value, "location": "图像可见区域"}]
                for key, value in fields.items() if value}
    return ModelFileOutput(role="Invoice", fields=fields, field_evidence=evidence)


class RealLocalChainDraft(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.cfg = RunConfig("imap.test", "user", "password",
                             output_root=self.temp.name, dashscope_api_key="fixed-no-network")

    def execute(self, msg: EmailMessage, *, pdf_ai=None, image_model=None, links=None):
        fake = FakeImap(msg)
        with patch("pipeline.imaplib.IMAP4_SSL", return_value=fake):
            from contextlib import ExitStack
            with ExitStack() as stack:
                if pdf_ai:
                    stack.enter_context(patch("pdf_processor.call_ai_audit_and_rename", side_effect=pdf_ai))
                if image_model:
                    stack.enter_context(patch("image_processor.extract_single_file", side_effect=image_model))
                if links is not None:
                    stack.enter_context(patch("pipeline.download_links", return_value=links))
                _mismatches, stats = run_pipeline(self.cfg)
        self.assertEqual(fake.fetch_queries, ["(BODY.PEEK[HEADER.FIELDS (SUBJECT DATE)])", "(BODY.PEEK[])"])
        self.assertIsNone(stats["imap_error"])
        self.assertIsNone(stats["run_error"])
        self.assertEqual(stats["delivery_status"], "ready", stats.get("delivery_error"))
        raw_zip = stats["delivery_zip_bytes"]
        self.assertIsInstance(raw_zip, bytes)
        with zipfile.ZipFile(io.BytesIO(raw_zip)) as archive:
            self.assertIsNone(archive.testzip())
            names = archive.namelist()
            self.assertTrue(all(name.startswith("Invoice_Task/") for name in names))
            self.assertTrue(all(name.endswith((".xlsx", ".json", ".pdf")) for name in names))
            self.assertFalse(any("2_Extracted" in name or "1_Downloaded" in name for name in names))
            zipped_json = json.loads(archive.read("Invoice_Task/task_debug.json"))
            with Path(self.cfg.debug_file).open(encoding="utf-8") as source:
                self.assertEqual(zipped_json, json.load(source))
            self.assertEqual(archive.read("Invoice_Task/发票信息汇总表.xlsx"),
                             Path(self.cfg.summary_excel_file).read_bytes())
            pdf_entries = [name for name in names if name.endswith(".pdf")]
            return stats, zipped_json, pdf_entries, raw_zip

    @unittest.skipUnless(XML.is_file(), "缺用户 XML A 原件")
    def test_real_xml_only_from_eml_to_excel_json_zip(self):
        stats, mails, pdf_entries, _ = self.execute(mail(XML))
        self.assertEqual(stats["summary_row_count"], 1)
        self.assertEqual(pdf_entries, [])
        record = mails[0]["records"][0]
        self.assertEqual(record["fields"]["invoice_number"], NUMBER_A)
        self.assertEqual(record["fields"]["invoice_date"], "2026-04-27")
        self.assertEqual(record["fields"]["total_amount"], "121.74")
        self.assertEqual(record["status"], "缺凭证")
        self.assertEqual(record["source"]["source_file_seq"], ["01"])
        self.assertEqual(mails[0]["files"][0]["exported_name"], "")
        book = load_workbook(self.cfg.summary_excel_file, data_only=True)
        sheet = book.active
        self.assertEqual([sheet.cell(1, col).value for col in range(1, 9)], HEADERS)
        self.assertEqual(sheet.max_row, 2)
        self.assertEqual(sheet.cell(2, 2).value, NUMBER_A)
        self.assertEqual(sheet.cell(2, 3).value, 121.74)

    @unittest.skipUnless(XML.is_file() and PDF.is_file(), "缺用户 XML/PDF A 原件")
    def test_real_xml_pdf_with_controlled_ai_then_exact_pdf_delivery(self):
        with patch("pdf_processor.call_ai_audit_and_rename", side_effect=controlled_pdf_ai):
            stats, mails, pdf_entries, raw_zip = self.execute(mail(XML, PDF))
        self.assertEqual(stats["summary_row_count"], 1)
        self.assertEqual(len(mails[0]["records"]), 1)
        record = mails[0]["records"][0]
        self.assertEqual(record["fields"]["invoice_number"], NUMBER_A)
        self.assertEqual(set(record["source"]["source_file_seq"]), {"01", "02"})
        self.assertEqual(len(pdf_entries), 1)
        self.assertEqual(mails[0]["files"][1]["exported_name"], Path(pdf_entries[0]).name)
        with zipfile.ZipFile(io.BytesIO(raw_zip)) as archive:
            self.assertEqual(archive.read(pdf_entries[0]), PDF.read_bytes())

    @unittest.skipUnless(IMAGE.is_file(), "缺用户图片 B 原件")
    def test_real_jpeg_named_png_with_controlled_model_and_converted_pdf(self):
        with patch("image_processor.extract_single_file", side_effect=controlled_image_model):
            stats, mails, pdf_entries, raw_zip = self.execute(mail(IMAGE))
        self.assertEqual(stats["summary_row_count"], 1)
        self.assertEqual(mails[0]["files"][0]["file_type"], "图片")
        self.assertEqual(mails[0]["files"][0]["original_name"], IMAGE.name)
        self.assertEqual(mails[0]["records"][0]["fields"]["invoice_number"], NUMBER_B)
        self.assertEqual(len(pdf_entries), 1)
        with zipfile.ZipFile(io.BytesIO(raw_zip)) as archive:
            with pdfplumber.open(io.BytesIO(archive.read(pdf_entries[0]))) as pdf:
                self.assertEqual(len(pdf.pages), 1)
                self.assertAlmostEqual(pdf.pages[0].width / pdf.pages[0].height,
                                       1080 / 698, delta=0.02)

    @unittest.skipUnless(XML.is_file() and PDF.is_file(), "缺用户 XML/PDF A 原件")
    def test_attachment_and_controlled_link_share_md5_and_pdf_delivery(self):
        linked = DownloadedFile(PDF.name, PDF.read_bytes(),
                                "正文链接 example.test（PDF）", "PDF")
        with patch("pdf_processor.call_ai_audit_and_rename", side_effect=controlled_pdf_ai):
            stats, mails, pdf_entries, _ = self.execute(
                mail(XML), links=([DownloadedFile("duplicate.xml", XML.read_bytes(),
                    "正文链接 example.test（XML）", "XML"), linked], []))
        self.assertEqual(stats["summary_row_count"], 1)
        self.assertEqual(len(mails[0]["files"]), 2)
        self.assertEqual(len(mails[0]["dedup_skipped"]), 1)
        self.assertEqual(mails[0]["files"][1]["acquisition"], "正文链接")
        self.assertEqual(len(pdf_entries), 1)

    @unittest.skipUnless(OFD_A.is_file(), "缺用户 OFD A 原件")
    def test_real_ofd_a_only_fields_preserved_conversion_rejected(self):
        stats, mails, pdf_entries, _ = self.execute(mail(OFD_A))
        self.assertEqual(stats["summary_row_count"], 1)
        self.assertEqual(pdf_entries, [])
        record = mails[0]["records"][0]
        self.assertEqual(record["fields"]["invoice_number"], NUMBER_A)
        self.assertEqual(record["fields"]["total_amount"], "121.74")
        self.assertEqual(record["source"]["source_file_seq"], ["01"])
        self.assertEqual(mails[0]["files"][0]["exported_name"], "")
        self.assertIn("OFD_CONVERSION_FAILED", json.dumps(mails[0], ensure_ascii=False))
        sheet = load_workbook(self.cfg.summary_excel_file, data_only=True).active
        self.assertEqual(sheet.cell(2, 2).value, NUMBER_A)
        self.assertEqual(sheet.cell(2, 3).value, 121.74)

    @unittest.skipUnless(XML.is_file() and PDF.is_file() and OFD_A.is_file(),
                         "缺用户 XML/PDF/OFD A 原件")
    def test_real_xml_pdf_ofd_a_merges_and_keeps_original_pdf(self):
        with patch("pdf_processor.call_ai_audit_and_rename", side_effect=controlled_pdf_ai):
            stats, mails, pdf_entries, raw_zip = self.execute(mail(XML, PDF, OFD_A))
        self.assertEqual(stats["summary_row_count"], 1)
        self.assertEqual(len(mails[0]["records"]), 1)
        self.assertEqual(mails[0]["records"][0]["fields"]["invoice_number"], NUMBER_A)
        self.assertEqual(set(mails[0]["records"][0]["source"]["source_file_seq"]),
                         {"01", "02", "03"})
        self.assertEqual(len(pdf_entries), 1)
        self.assertIn("OFD_CONVERSION_FAILED", json.dumps(mails[0], ensure_ascii=False))
        with zipfile.ZipFile(io.BytesIO(raw_zip)) as archive:
            self.assertEqual(archive.read(pdf_entries[0]), PDF.read_bytes())

    @unittest.skipUnless(OFD_COMPLETE.is_file(), "缺用户完整呼我 OFD 原件")
    def test_real_complete_ofd_only_converts_pdf_in_zip(self):
        stats, mails, pdf_entries, raw_zip = self.execute(mail(OFD_COMPLETE))
        self.assertEqual(stats["summary_row_count"], 1)
        record = mails[0]["records"][0]
        self.assertEqual(record["fields"]["invoice_number"], "26442000004589632231")
        self.assertEqual(record["fields"]["total_amount"], "48.31")
        self.assertEqual(len(pdf_entries), 1)
        self.assertEqual(mails[0]["files"][0]["exported_name"],
                         Path(pdf_entries[0]).name)
        with zipfile.ZipFile(io.BytesIO(raw_zip)) as archive:
            pdf_bytes = archive.read(pdf_entries[0])
            self.assertTrue(pdf_bytes.startswith(b"%PDF-"))
            with pdfplumber.open(io.BytesIO(pdf_bytes)) as pdf:
                self.assertEqual(len(pdf.pages), 1)
                page = pdf.pages[0]
                text = page.extract_text() or ""
                self.assertIn("26442000004589632231", text)
                self.assertIn("48.31", text)
                self.assertGreaterEqual(len(page.images), 1)


if __name__ == "__main__":
    unittest.main()
