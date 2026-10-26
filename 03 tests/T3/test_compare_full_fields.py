"""Comparator tests use invented values, not application output as an oracle."""
from __future__ import annotations

import json
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook

sys.path.insert(0, str(Path(__file__).resolve().parent))
from compare_full_fields import COLUMNS, SCHEMA_FORMATS, compare_files  # noqa: E402


class FullFieldComparisonTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.expected = self.root / "expected.json"
        self.actual_json = self.root / "task_debug.json"
        self.actual_excel = self.root / "summary.xlsx"
        fields = {
            "invoice_date": "2026-10-20", "invoice_number": "26330000000000000001",
            "total_amount": "121.74", "invoice_type": "电子发票（普通发票）",
            "receipt_type": "办公用品", "seller": "测试销售方",
        }
        self.manifest = {
            "review_status": "测试工程师制作、待用户审核",
            "schema": {
                "excel_columns": [[title, key] for title, key in COLUMNS],
                "formats": dict(SCHEMA_FORMATS),
                "empty_values": {"invoice_type": "", "travel_time": "",
                                 "origin_dest": "", "required_missing": "Not Found"},
                "excel_amount_number_format": "0.00",
            },
            "json": {"mails": [{
                "mail_id": "Msg1", "mail_uid": "9101", "subject": "测试电子发票",
                "mail_date": "2026-10-26 09:00",
                "files": [{"file_seq": "01", "file_type": "XML", "original_name": "invoice.xml"}],
                "itineraries": [],
                "records": [{"record_id": "Msg1-01", "kind": "b类", "status": "缺凭证",
                             "fields": fields, "source_file_seq": ["01"]}],
            }]},
            "excel": {"rows": [{**fields, "travel_time": "", "origin_dest": ""}]},
        }
        self.observed = [{
            "mail_id": "Msg1", "mail_uid": "9101", "subject": "测试电子发票",
            "mail_date": "2026-10-26 09:00",
            "files": [{"file_seq": "01", "file_type": "XML", "original_name": "invoice.xml"}],
            "itineraries": [],
            "records": [{"record_id": "Msg1-01", "kind": "b类", "status": "缺凭证",
                         "fields": fields.copy(), "source": {"source_file_seq": ["01"]}}],
        }]
        self._write()

    def tearDown(self):
        self.temp.cleanup()

    def _write(self):
        self.expected.write_text(json.dumps(self.manifest, ensure_ascii=False), encoding="utf-8")
        self.actual_json.write_text(json.dumps(self.observed, ensure_ascii=False), encoding="utf-8")
        book = Workbook()
        sheet = book.active
        sheet.append([title for title, _ in COLUMNS])
        row = self.manifest["excel"]["rows"][0]
        sheet.append([float(row[key]) if key == "total_amount" else row[key]
                      for _, key in COLUMNS])
        sheet.cell(row=2, column=3).number_format = "0.00"
        book.save(self.actual_excel)

    def test_complete_draft_can_be_compared_but_formal_is_refused(self):
        self.assertEqual(compare_files(self.expected, self.actual_json,
                                       self.actual_excel)["status"], "PASS")
        self.assertEqual(compare_files(self.expected, self.actual_json,
                                       self.actual_excel, "formal")["status"], "REFUSED")

    def test_json_and_excel_each_report_precise_independent_differences(self):
        self.observed[0]["records"][0]["fields"]["seller"] = "错误销售方"
        self._write()
        book = Workbook()
        sheet = book.active
        sheet.append([title for title, _ in COLUMNS])
        row = self.manifest["excel"]["rows"][0]
        sheet.append([float(row[key]) if key == "total_amount" else
                      "错误票面类型" if key == "receipt_type" else row[key]
                      for _, key in COLUMNS])
        sheet.cell(row=2, column=3).number_format = "0.00"
        book.save(self.actual_excel)
        errors = compare_files(self.expected, self.actual_json,
                               self.actual_excel)["errors"]
        self.assertTrue(any("json.mails[0].records[0].fields.seller" in item for item in errors))
        self.assertTrue(any("excel.rows[0].receipt_type" in item for item in errors))

    def test_incomplete_expectation_is_not_silently_accepted(self):
        del self.manifest["json"]["mails"][0]["records"][0]["fields"]["seller"]
        self._write()
        errors = compare_files(self.expected, self.actual_json,
                               self.actual_excel)["errors"]
        self.assertTrue(any("expectation incomplete" in item for item in errors))


if __name__ == "__main__":
    unittest.main()
